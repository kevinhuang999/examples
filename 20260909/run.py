# -*- coding: utf-8 -*-
"""一键运行：读原始数据，按模板批量生成报告，并自动回读验证。

跨平台：Windows / macOS / Linux 都是一条命令，缺依赖会自动装。
    python run.py
    python run.py --no-pause   # 不暂停（CI 或脚本里用）

Windows 上直接双击 run.py 也行，跑完会停住窗口等回车。

打包给别人时，把整个目录拷走即可，路径全部相对定位，无硬编码。
"""
import csv
import socket
import subprocess
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path

# 控制台编码随环境变，只加容错，不强行改编码。
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(errors="replace")
    except Exception:
        pass

MIRROR = "https://pypi.tuna.tsinghua.edu.cn/simple"
NO_PAUSE = "--no-pause" in sys.argv


def _pause(msg=""):
    """Windows 双击 .py 时留住窗口，避免一闪而过。"""
    if NO_PAUSE:
        if msg:
            print(msg)
        return
    if msg:
        print(msg)
    try:
        if sys.platform == "win32" and sys.stdin and sys.stdin.isatty():
            input("按回车退出")
    except Exception:
        pass


def _ensure_openpyxl():
    """没有就装。先直接装，失败再退到 --user（无管理员权限的机器）。"""
    try:
        from openpyxl import load_workbook
        return load_workbook
    except ImportError:
        pass

    print(f"缺 openpyxl，正在用清华源安装（解释器: {sys.executable}）...")
    base = [sys.executable, "-m", "pip", "install", "-i", MIRROR, "openpyxl"]
    attempts = [base]
    if sys.prefix == sys.base_prefix:   # 不在虚拟环境里，普通安装可能没权限
        attempts.append(base[:4] + ["--user"] + base[4:])

    for cmd in attempts:
        subprocess.run(cmd, check=False)
        try:
            from openpyxl import load_workbook
            return load_workbook
        except ImportError:
            continue

    _pause(f"自动安装失败。手动执行：\n  {sys.executable} -m pip install -i {MIRROR} openpyxl")
    sys.exit(1)


load_workbook = _ensure_openpyxl()

HERE = Path(__file__).resolve().parent
RAW_DIR = HERE / "01_raw_data"                           # 原始数据（CSV）
TEMPLATE = HERE / "source" / "templates" / "检验报告模板.xlsx"
OUT_DIR = HERE / "02_output"                             # 生成的报告
LOG_DIR = HERE / "source"                                # 运行日志（run_log.txt 就落在 source/ 下）
DETAIL_START = 12      # 模板里明细区第一行
DETAIL_RESERVED = 60   # 模板里预留的明细行数
FOOTER_LAST = 76       # 模板最后一行（结论 + 签字栏），打印区域必须固定到这里

# 左边是模板里的命名区域，右边是 CSV 列名
FIELDS = {
    "样品编号": "SampleID",
    "样品名称": "SampleName",
    "检品批号": "BatchNo",
    "检验日期": "TestDate",
}


def cell_by_name(ws, name):
    """按命名区域取单元格。openpyxl 不支持 ws["命名区域"]，得自己解析。"""
    dn = ws.parent.defined_names[name]
    _, ref = next(iter(dn.destinations))
    return ws[ref.replace("$", "")]


def fill_one(header: dict, items: list[dict]) -> Path:
    """按模板生成一份报告。header 是抬头字段，items 是检项明细。"""
    wb = load_workbook(TEMPLATE)
    wb.calculation.fullCalcOnLoad = True  # 打开 Excel 时强制重算
    ws = wb["报告"]

    # 1) 填抬头
    for name, col in FIELDS.items():
        cell_by_name(ws, name).value = header[col]

    # 2) 填明细行
    if len(items) > DETAIL_RESERVED:
        raise ValueError(f"明细 {len(items)} 行，超过模板预留的 {DETAIL_RESERVED} 行")

    for i, it in enumerate(items):
        r = DETAIL_START + i
        ws.cell(r, 1, it["检验项目"])
        ws.cell(r, 2, it["标准规定"])
        v = (it["检验结果"] or "").strip()
        ws.cell(r, 3, float(v) if v else None)  # 写数值，写字符串数字格式不生效
        ws.cell(r, 4, it["单项结论"])

    # 3) 多余预留行整行隐藏
    last = DETAIL_START + len(items) - 1
    for r in range(last + 1, DETAIL_START + DETAIL_RESERVED):
        ws.row_dimensions[r].hidden = True

    # 4) 打印区域固定到模板最后一行，别跟着明细走
    ws.print_area = f"A1:G{FOOTER_LAST}"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"检验报告_{header['SampleID']}.xlsx"
    wb.save(out)
    wb.close()
    return out


def main() -> bool:
    with open(RAW_DIR / "results.csv", encoding="utf-8-sig") as f:   # 别省 utf-8-sig
        rows = list(csv.DictReader(f))

    grouped = defaultdict(list)
    for row in rows:
        grouped[row["SampleID"]].append(row)

    generated = []
    for sid, items in grouped.items():
        path = fill_one(items[0], items)
        generated.append((sid, len(items), path))
        print(f"已生成 {path.name}，明细 {len(items)} 行")

    print("\n---- 回读验证 ----")
    ok = True
    for sid, n, path in generated:
        wb = load_workbook(path)
        ws = wb["报告"]

        checks = {
            "抬头写入": cell_by_name(ws, "样品编号").value == sid,
            "数字格式保留": ws.cell(12, 3).number_format == "0.000",
            "边框保留": ws.cell(12, 3).border.left.style == "thin",
            "合并区保留": any(str(r) == "A1:G1" for r in ws.merged_cells.ranges),
            "结论公式在": str(ws["B72"].value).startswith("=IF("),
            "结论行在打印区内": ws.print_area == f"'报告'!$A$1:$G${FOOTER_LAST}",
            "明细末行未隐藏": ws.row_dimensions[12 + n - 1].hidden is False,
            "多余行已隐藏": ws.row_dimensions[12 + n].hidden is True,
            "强制重算": wb.calculation.fullCalcOnLoad is True,
        }
        bad = [k for k, v in checks.items() if not v]
        ok = ok and not bad
        print(f"{sid}: {'全部通过' if not bad else '不通过 -> ' + ', '.join(bad)}")
        wb.close()

    print("\n结论行文本（公式，打开 Excel 才显示结果）:")
    print(f"  {load_workbook(generated[0][2])['报告']['B72'].value}")
    print(f"\n{'PASS' if ok else 'FAIL'}  输出目录: {OUT_DIR.name}/")
    return ok


class _Tee:
    """同时写控制台和日志文件，万一控制台中文乱码还有日志可看。"""

    def __init__(self, *streams):
        self.streams = streams

    def write(self, s):
        for st in self.streams:
            try:
                st.write(s)
            except Exception:
                pass

    def flush(self):
        for st in self.streams:
            try:
                st.flush()
            except Exception:
                pass


def _note_unfinished():
    """上一条执行若被手动关窗口打断，日志里会缺完成行。开跑前补一句说明。"""
    LOG_FILE = LOG_DIR / "run_log.txt"
    if not LOG_FILE.exists():
        return
    tail = ""
    with open(LOG_FILE, encoding="utf-8", errors="replace") as f:
        for ln in f:
            if ln.strip():
                tail = ln.strip()
    if tail and not tail.startswith("执行完成"):
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write("（上一条没有留下完成时间，多半是窗口被手动关闭或中途强制中断）\n")


def _host_info() -> str:
    """主机名 + 本机 IP，用来分辨这段日志是哪台机器跑出来的。"""
    try:
        host = socket.gethostname()
    except Exception:
        host = "unknown"
    try:
        # UDP 不会真的发包，只按路由表取本机对外网卡的地址
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:
            s.connect(("8.8.8.8", 80))
            ip = s.getsockname()[0]
    except Exception:
        try:
            ip = socket.gethostbyname(host)
        except Exception:
            ip = "未知"
    return f"{host}（{ip}）"


def _run():
    """一次执行 = 追加一段：开始时间 + 过程 + 完成时间。历史日志不覆盖。"""
    LINE = "=" * 60
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    _note_unfinished()
    _log = open(LOG_DIR / "run_log.txt", "a", encoding="utf-8")   # 追加，别用 w
    _log.write(f"\n{LINE}\n")
    _log.write(f"运行主机：{_host_info()}\n")
    _log.write(f"开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}\n")
    _log.write(f"Python   ：{sys.version.split()[0]}  （{Path(sys.executable).name}）\n")
    _log.flush()

    _orig = sys.stdout
    sys.stdout = _Tee(sys.stdout, _log)
    t0, ok, tail = time.time(), False, ""
    try:
        try:
            ok = main()
        except KeyboardInterrupt:
            tail = "\n被手动中断（Ctrl+C）。"
        except Exception as e:
            import traceback
            traceback.print_exc()
            tail = f"\n跑失败了：{type(e).__name__}: {e}"
        else:
            tail = "\n跑完了。结果在 02_output/，运行记录已写入 source/run_log.txt。"
        if tail:
            print(tail)
        # 完成行必须写在暂停之前：双击后关窗口就丢不了
        print(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}"
              f"  用时 {time.time() - t0:.1f}s  判定 {'PASS' if ok else 'FAIL'}")
    finally:
        sys.stdout.flush()
        _log.close()
        sys.stdout = _orig

    _pause()          # 等回车放在日志收尾之后
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    _run()

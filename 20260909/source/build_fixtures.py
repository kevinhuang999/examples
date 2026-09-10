# -*- coding: utf-8 -*-
"""生成示例用的模板与数据。改动模板结构后重跑一次即可。"""
import csv
import subprocess
import sys
from pathlib import Path

MIRROR = "https://pypi.tuna.tsinghua.edu.cn/simple"

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.workbook.defined_name import DefinedName
except ImportError:
    print(f"缺 openpyxl，正在用清华源安装（解释器: {sys.executable}）...")
    subprocess.run([sys.executable, "-m", "pip", "install", "-i", MIRROR, "openpyxl"],
                   check=False)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.workbook.defined_name import DefinedName

BASE = Path(__file__).resolve().parent.parent   # 示例根目录（本脚本在 source/ 里）
TPL_DIR = BASE / "source" / "templates"
TPL_DIR.mkdir(parents=True, exist_ok=True)

DETAIL_START = 12      # 明细区第一行
DETAIL_RESERVED = 60   # 预留明细行数
FOOTER_LAST = 76       # 模板最后一行（结论 + 签字栏）

thin = Side(style="thin")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "报告"

# --- 抬头 ---
ws.merge_cells("A1:G1")
ws["A1"] = "XX 检测技术有限公司  检验报告"
ws["A1"].font = Font(size=16, bold=True)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30

# 报告编号（合并单元格，演示"只有左上角能写"）
ws.merge_cells("A2:B2")
ws["A2"] = "报告编号"
ws["A2"].font = Font(bold=True)
ws.merge_cells("C2:G2")
ws["C2"] = "BG-2026-0909"
ws["C2"].alignment = Alignment(horizontal="left")

# --- 样品信息区：A 列标签，B 列值 ---
INFO = [("样品编号", "B3"), ("样品名称", "B4"), ("检品批号", "B5"), ("检验日期", "B6")]
for label, value_cell in INFO:
    r = int(value_cell[1:])
    ws.cell(r, 1, label).font = Font(bold=True)
    ws.cell(r, 1).border = box
    ws.cell(r, 2).border = box
    ws.merge_cells(f"C{r}:G{r}")  # 留白区，合并后只读

# 命名区域（对应 Excel 里 公式 -> 定义名称）
for name, ref in INFO:
    wb.defined_names.add(DefinedName(name, attr_text=f"报告!${ref}"))

# --- 检验依据（横向合并，演示合并区） ---
ws.merge_cells("A8:B8")
ws["A8"] = "检验依据"
ws["A8"].font = Font(bold=True)
ws.merge_cells("C8:G8")
ws["C8"] = "《中国药典》2025 年版 四部通则"

# --- 明细表头 ---
headers = ["检验项目", "标准规定", "检验结果", "单项结论"]
for c, h in enumerate(headers, start=1):
    cell = ws.cell(11, c, h)
    cell.font = Font(bold=True)
    cell.border = box
    cell.alignment = Alignment(horizontal="center")

# --- 预留明细行：60 行带样式空行，C 列数字格式 0.000 ---
for r in range(DETAIL_START, DETAIL_START + DETAIL_RESERVED):
    for c in range(1, 5):
        cell = ws.cell(r, c)
        cell.border = box
        if c == 3:
            cell.number_format = "0.000"
        if c == 4:
            cell.alignment = Alignment(horizontal="center")
    ws.row_dimensions[r].height = 18

# --- 结论行（在预留区之后，行号固定不变） ---
ws.cell(72, 1, "结论").font = Font(bold=True)
ws.cell(72, 1).border = box
ws.merge_cells("B72:G72")
ws["B72"] = (
    '=IF(COUNTIF(D12:D71,"不符合规定")=0,'
    '"本品按标准检验，结果符合规定",'
    '"本品按标准检验，结果不符合规定")'
)
ws["B72"].border = box
ws["B72"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[72].height = 26

# --- 签字栏 ---
for i, who in enumerate(["检验人", "复核人", "批准人"]):
    r = 74 + i
    ws.cell(r, 1, who).font = Font(bold=True)
    ws.cell(r, 1).border = box
    ws.merge_cells(f"B{r}:C{r}")
    ws.cell(r, 2).border = box
    ws.merge_cells(f"E{r}:F{r}")
    ws.cell(r, 5, "日期").font = Font(bold=True)
    ws.cell(r, 5).border = box
    ws.cell(r, 7).border = box

# --- 打印设置 ---
ws.print_area = f"A1:G{FOOTER_LAST}"
ws.print_title_rows = "11:11"
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True

for col, width in zip("ABCDEFG", [18, 22, 14, 16, 12, 12, 12]):
    ws.column_dimensions[col].width = width

tpl = TPL_DIR / "检验报告模板.xlsx"
wb.save(tpl)
print(f"模板已生成: {tpl}")

# ---------- 造 CSV（utf-8-sig，模拟仪器导出带 BOM） ----------
rows = []


def add(sid, name, batch, date, items):
    for it in items:
        rows.append({
            "SampleID": sid, "SampleName": name, "BatchNo": batch, "TestDate": date,
            "检验项目": it[0], "标准规定": it[1], "检验结果": it[2], "单项结论": it[3],
        })


add("S-2026-0001", "注射用头孢曲松钠", "B260901", "2026-09-09", [
    ("性状", "白色或类白色结晶性粉末", "", "符合规定"),
    ("含量", "90.0%~110.0%", "99.512", "符合规定"),
    ("干燥失重", "≤0.5%", "0.050", "符合规定"),
])

add("S-2026-0002", "阿托伐他汀钙片", "B260812", "2026-09-09", [
    ("性状", "薄膜衣片，除去包衣后显白色", "", "符合规定"),
    ("鉴别（HPLC）", "供试品主峰保留时间应与对照品一致", "", "符合规定"),
    ("含量均匀度", "应符合规定", "", "符合规定"),
    ("溶出度", "限度为标示量的 80%", "95.400", "符合规定"),
    ("有关物质-单个杂质", "≤0.5%", "0.185", "符合规定"),
    ("有关物质-总杂质", "≤2.0%", "0.720", "符合规定"),
    ("含量测定", "95.0%~105.0%", "100.230", "符合规定"),
    ("干燥失重", "≤0.5%", "0.310", "符合规定"),
    ("重金属", "≤20ppm", "", "符合规定"),
    ("微生物限度", "应符合规定", "", "符合规定"),
    ("装量差异", "应符合规定", "", "符合规定"),
])

add("S-2026-0003", "氯化钠注射液", "0260903", "2026-09-09", [
    ("含量", "0.85%~0.95%(g/ml)", "0.900", "符合规定"),
])

csv_path = BASE / "01_raw_data" / "results.csv"
with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
    w.writeheader()
    w.writerows(rows)
print(f"数据已生成: {csv_path}（{len(rows)} 行明细，3 个样品）")

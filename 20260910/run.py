# -*- coding: utf-8 -*-
"""引流长文《用 Python 按 Word 模板批量生成检验报告书，怎么保留原模板样式》配套脚本。

核心代码与正文逐字一致，末尾多了自动回读校验，方便确认生成结果对不对。
跨平台：Windows / macOS / Linux 都是一条命令，缺依赖会自动装。
    python run.py
    python run.py --no-pause   # 不暂停（CI 或脚本里用）

打包给别人时，把整个目录拷走即可，路径全部相对定位，无硬编码。
"""
import csv
import subprocess
import sys
import time
from collections import defaultdict
from datetime import datetime
from pathlib import Path

# 控制台编码随环境变，只加容错，不强行改编码。
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(errors="replace")
    except Exception:
        pass

MIRROR = "https://pypi.tuna.tsinghua.edu.cn/simple"
NO_PAUSE = "--no-pause" in sys.argv


def _pause(msg=""):
    """Windows 双击 .py 时留住窗口，避免一闪而过。"""
    if msg:
        print(msg)
    if NO_PAUSE:
        return
    try:
        if sys.platform == "win32" and sys.stdin and sys.stdin.isatty():
            input("按回车退出")
    except Exception:
        pass


def _ensure(*modules_pkgs):
    """缺哪个装哪个。先直接装，失败再退到 --user（无管理员权限的机器）。"""
    for mod, pkg in modules_pkgs:
        try:
            __import__(mod)
            continue
        except ImportError:
            pass
        print(f"缺 {pkg}，正在用清华源安装（解释器: {sys.executable}）...")
        base = [sys.executable, "-m", "pip", "install", "-i", MIRROR, pkg]
        attempts = [base]
        if sys.prefix == sys.base_prefix:      # 不在虚拟环境里，普通安装可能没权限
            attempts.append(base[:4] + ["--user"] + base[4:])
        for cmd in attempts:
            subprocess.run(cmd, check=False)
            try:
                __import__(mod)
                break
            except ImportError:
                continue
        else:
            _pause(f"自动安装失败。手动执行：\n  {sys.executable} -m pip install -i {MIRROR} {pkg}")
            sys.exit(1)


_ensure(("docx", "python-docx"), ("docxtpl", "docxtpl"), ("jinja2", "jinja2"))

from docx import Document
from docxtpl import DocxTemplate
from jinja2 import Environment

HERE = Path(__file__).resolve().parent
RAW_DIR = HERE / "01_raw_data"                           # 原始数据（CSV）
TEMPLATE = HERE / "source" / "templates" / "检验报告书模板.docx"
OUT_DIR = HERE / "02_output"                             # 生成的报告
LOG_DIR = HERE / "source"                                # 运行日志（run_log.txt 就落在 source/ 下）
CSV_FILE = RAW_DIR / "检验结果.csv"

FIELDS = {          # 左边是模板里的占位符名，右边是 CSV 列名
    "样品编号": "SampleID",
    "样品名称": "SampleName",
    "批号": "BatchNo",
}


def read_source(csv_path: Path) -> dict:
    """一个样品编号一份报告，明细按编号归堆。"""
    with open(csv_path, encoding="utf-8-sig") as f:   # 别省 utf-8-sig
        rows = list(csv.DictReader(f))
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["SampleID"]].append(row)
    return grouped


def make_context(sid: str, items: list) -> dict:
    head = items[0]
    ctx = {name: head[col] for name, col in FIELDS.items()}
    ctx["报告编号"] = f"COA-{sid}"
    ctx["结论"] = "本品按企业内控标准检验，结果符合规定。"
    ctx["items"] = [
        {"proj": r["检验项目"], "spec": r["标准规定"],
         "res": r["检验结果"], "concl": r["单项结论"]}
        for r in items
    ]
    return ctx


def render_one(ctx: dict) -> Path:
    doc = DocxTemplate(TEMPLATE)
    # autoescape=True：数据里的 < 和 & 才不会被 XML 吞掉
    doc.render(ctx, Environment(autoescape=True))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / f"检验报告书_{ctx['样品编号']}.docx"
    doc.save(out)
    return out


def main() -> bool:
    grouped = read_source(CSV_FILE)
    for sid, items in grouped.items():
        path = render_one(make_context(sid, items))
        print(f"已生成 {path.name}，明细 {len(items)} 行")

    # ---------- 以下为校验，正文里没有 ----------
    print("\n===== 回读校验 =====")
    expect = {
        "S-2026-0910-001": (4, True),    # 表头 1 行 + 明细 4 行；数据里含 < 和 &
        "S-2026-0910-002": (3, False),   # 明细 3 行；数据里无特殊字符
    }
    ok = True
    for sid, (n, has_special) in expect.items():
        f = OUT_DIR / f"检验报告书_{sid}.docx"
        d = Document(f)
        rows = d.tables[0].rows
        head = d.paragraphs[1].text
        hdr = d.sections[0].header.paragraphs[0].text
        specs = [c.text for row in rows for c in row.cells if "<" in c.text or "&" in c.text]
        sizes = {r.font.size.pt for row in rows for c in row.cells
                 for p in c.paragraphs for r in p.runs if r.font.size}
        check = [
            (f"行数 = 表头1+明细{n} = {n+1}", len(rows) == n + 1, len(rows)),
            ("页眉变量已替换", hdr == f"报告编号：COA-{sid}", hdr),
            ("抬头含样品编号", sid in head, head[:24]),
            ("特殊字符 < & 保留", (not has_special)
             or any("<0.5%" in s and "&" in s for s in specs), specs or "本样品无特殊字符"),
            ("明细行字号 9.0pt", sizes == {9.0}, sizes),
        ]
        print(f"\n{f.name}")
        for name, passed, val in check:
            print(f"  [{'OK ' if passed else 'FAIL'}] {name}  ->  {val}")
            ok &= passed

    print("\n总判定:", "PASS" if ok else "FAIL")
    return bool(ok)


class _Tee:
    """同时写控制台和日志文件，万一控制台中文乱码还有日志可看。"""

    def __init__(self, *streams):
        self.streams = streams

    def write(self, s):
        for st in self.streams:
            try:
                st.write(s)
            except Exception:
                pass

    def flush(self):
        for st in self.streams:
            try:
                st.flush()
            except Exception:
                pass


def _note_unfinished():
    """上一条执行若被手动关窗口打断，日志里会缺完成行。开跑前补一句说明。"""
    LOG_FILE = LOG_DIR / "run_log.txt"
    if not LOG_FILE.exists():
        return
    tail = ""
    with open(LOG_FILE, encoding="utf-8", errors="replace") as f:
        for ln in f:
            if ln.strip():
                tail = ln.strip()
    if tail and not tail.startswith("执行完成"):
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write("（上一条没有留下完成时间，多半是窗口被手动关闭或中途强制中断）\n")


def _run():
    """一次执行 = 追加一段：开始时间 + 过程 + 完成时间。历史日志不覆盖。"""
    LINE = "=" * 60
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    _note_unfinished()
    _log = open(LOG_DIR / "run_log.txt", "a", encoding="utf-8")   # 追加，别用 w
    _log.write(f"\n{LINE}\n")
    _log.write(f"开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}\n")
    _log.write(f"Python   ：{sys.version.split()[0]}  （{Path(sys.executable).name}）\n")
    _log.flush()

    _orig = sys.stdout
    sys.stdout = _Tee(sys.stdout, _log)
    t0, ok, tail = time.time(), False, ""
    try:
        try:
            ok = main()
        except KeyboardInterrupt:
            tail = "\n被手动中断（Ctrl+C）。"
        except Exception as e:
            import traceback
            traceback.print_exc()
            tail = f"\n跑失败了：{type(e).__name__}: {e}"
        else:
            tail = "\n跑完了。结果在 02_output/，运行记录已写入 source/run_log.txt。"
        if tail:
            print(tail)
        # 完成行必须写在暂停之前：双击后关窗口就丢不了
        print(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}"
              f"  用时 {time.time() - t0:.1f}s  判定 {'PASS' if ok else 'FAIL'}")
    finally:
        sys.stdout.flush()
        _log.close()
        sys.stdout = _orig

    _pause()          # 等回车放在日志收尾之后
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    _run()

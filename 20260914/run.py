# -*- coding: utf-8 -*-
"""引流长文《Word 邮件合并批量生成检验报告，多条检项却只出来一行》配套脚本。

一个样品一份 Word 检验报告：读样品信息 + 仪器导出的检项明细，
把模板里的 {{字段}} 换成值、把明细表的样例行按检项数复制成 N 行。
核心代码与正文逐字一致，末尾多了自动回读校验，方便确认报告里该有的都在。
跨平台：Windows / macOS / Linux 都是一条命令，缺依赖会自动装。
    python run.py
    python run.py --no-pause   # 不暂停（CI 或脚本里用）

打包给别人时，把整个目录拷走即可，路径全部相对定位，无硬编码。
"""
import copy
import csv
import re
import socket
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

# 控制台编码随环境变，只加容错，不强行改编码。
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(errors="replace")
    except Exception:
        pass

MIRROR = "https://pypi.tuna.tsinghua.edu.cn/simple"
NO_PAUSE = "--no-pause" in sys.argv


def _pause(msg=""):
    """Windows 双击 .py 时留住窗口，避免一闪而过。"""
    if msg:
        print(msg)
    if NO_PAUSE:
        return
    try:
        if sys.platform == "win32" and sys.stdin and sys.stdin.isatty():
            input("按回车退出")
    except Exception:
        pass


def _ensure(*modules_pkgs):
    """缺哪个装哪个。先直接装，失败再退到 --user（无管理员权限的机器）。"""
    for mod, pkg in modules_pkgs:
        try:
            __import__(mod)
            continue
        except ImportError:
            pass
        print(f"缺 {pkg}，正在用清华源安装（解释器: {sys.executable}）...")
        base = [sys.executable, "-m", "pip", "install", "-i", MIRROR, pkg]
        attempts = [base]
        if sys.prefix == sys.base_prefix:      # 不在虚拟环境里，普通安装可能没权限
            attempts.append(base[:4] + ["--user"] + base[4:])
        for cmd in attempts:
            subprocess.run(cmd, check=False)
            try:
                __import__(mod)
                break
            except ImportError:
                continue
        else:
            _pause(f"自动安装失败。手动执行：\n  {sys.executable} -m pip install -i {MIRROR} {pkg}")
            sys.exit(1)


_ensure(("docx", "python-docx"))

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
SAMPLE_FILE = HERE / "01_raw_data" / "样品信息.csv"            # 样品级信息，一个样品一行
ITEM_FILE = HERE / "01_raw_data" / "检项明细_仪器导出.csv"      # 仪器导出的检项明细（长表）
TPL_FILE = HERE / "source" / "templates" / "检验报告_模板.docx"  # 报告模板
OUT_DIR = HERE / "02_output"                                  # 一个样品一份 Word
LOG_DIR = HERE / "source"                                     # 运行日志（run_log.txt 就落在 source/ 下）

KEY_COL = "样品编号"
SAMPLE_COLS = ["样品编号", "样品名称", "批号", "规格", "送检部门", "收样日期", "检测日期", "检验人", "复核人"]
# 仪器导出的明细里比报告上多一列「判定」，报告表上不用摆（结论行拿它算）
ITEM_COLS = ["检项", "标准规定", "检验结果", "单位", "判定"]
TABLE_COLS = ["检项", "标准规定", "检验结果", "单位"]

# 检项在报告里的固定顺序（药典体例：性状、鉴别、检查项、含量）。
# 不写死这个顺序，明细表按拼音排，pH 值会跑到最前面，报告没法看。
ITEM_ORDER = ["性状", "鉴别", "溶解度", "溶液颜色", "澄清度", "pH 值",
              "有关物质", "干燥失重", "炽灼残渣", "重金属", "微生物限度", "含量测定"]

# 文件名里不能出现的字符，换成横线才敢拿样品名去命名文件
BAD_CHARS = r'[\\/:*?"<>|\s]+'


def read_samples(path: Path) -> dict:
    """读样品信息，一个样品一行，返回 {样品编号: {字段: 值}}。"""
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            sid = (row.get(KEY_COL) or "").strip()
            if not sid:
                continue                       # 空行、表尾的合计行都不是数据
            out[sid] = {c: (row.get(c) or "").strip() for c in SAMPLE_COLS}
    return out


def read_items(path: Path) -> dict:
    """读仪器导出的检项明细，按样品编号归堆：一个样品一个检项列表。

    值一律当文本读，不转数字：`<0.5`、`符合` 这些和数字混在同一列，
    批号这种带前导零的字段也只有按文本读才保得住 00903。
    """
    out = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            sid = (row.get(KEY_COL) or "").strip()
            item = (row.get("检项") or "").strip()
            if not (sid and item):
                continue                       # 表尾的空行、合计行、说明行都落在这儿
            out.setdefault(sid, []).append({c: (row.get(c) or "").strip() for c in ITEM_COLS})
    return out


def sort_items(rows: list) -> list:
    """按报告体例排行——检项顺序照药典来，不按拼音排。"""
    return sorted(rows, key=lambda r: ITEM_ORDER.index(r["检项"])
                  if r["检项"] in ITEM_ORDER else len(ITEM_ORDER))


def fill_paragraph(para, mapping: dict) -> int:
    """把段落里的 {{字段}} 换成值，返回换掉几处。

    坑：Word 存盘或拼写检查会把 {{样品编号}} 拆进多个 run（模板里那行就是拆成三个 run 存的），
    逐个 run 做 replace 一处都换不到——得先把整段文本拼起来，再写回第一个 run。
    """
    full = para.text
    if "{{" not in full:
        return 0
    hits = 0
    for key, val in mapping.items():
        token = "{{" + key + "}}"
        hits += full.count(token)
        full = full.replace(token, str(val))
    if not hits:
        return 0
    runs = para.runs
    runs[0].text = full                       # 文本写回第一个 run，后面几个清空
    for r in runs[1:]:
        r.text = ""
    return hits


def write_cell(cell, text: str) -> None:
    """往单元格写字：写进第一个 run，模板排好的字号才守得住。

    直接用 cell.text = "..." 会把单元格清空重建，模板里的字号字体一起没。
    """
    para = cell.paragraphs[0]
    if not para.runs:
        para.add_run("")
    para.runs[0].text = str(text)
    for r in para.runs[1:]:
        r.text = ""


def fill_detail_table(table, items: list) -> int:
    """按检项数把明细表的样例行复制成 N 行，一行一个检项。

    python-docx 没有"插入行"的 API，只能把模板那行的 XML 深拷一份挂回表格；
    复制出来的行自带模板行的字号和边框，比自己画一行省事。
    """
    tpl_row = table.rows[-1]
    if not items:                             # 一个检项都没有的时候，别留一行空占位符
        table._tbl.remove(tpl_row._tr)
        return 0
    for _ in range(len(items) - 1):
        table._tbl.append(copy.deepcopy(tpl_row._tr))
    for row, item in zip(list(table.rows)[-len(items):], items):
        for cell, col in zip(row.cells, TABLE_COLS):
            write_cell(cell, item[col])
    return len(items)


def make_conclusion(items: list) -> str:
    """结论按明细的判定列算：有不符合就是不符合规定；有还没判的先不下结论；全符合才算符合规定。"""
    verdicts = [r["判定"] for r in items]
    if any("不符合" in v for v in verdicts):
        return "不符合规定"
    if not items or any(not v for v in verdicts):
        return "待结果出齐后判定"
    return "符合规定"


def safe_name(text: str) -> str:
    """把样品名里不能进文件名的字符换成横线——「阿莫西林/克拉维酸钾片」的斜杠会直接建出子目录。"""
    return re.sub(BAD_CHARS, "-", str(text)).strip("-")[:40]


def render_one(tpl_path: Path, sample: dict, items: list, out_dir: Path) -> Path:
    """一个样品一份 Word：模板里两张表，第 1 张样品信息，第 2 张检项明细。

    邮件合并是攒成一个大文档再想辙拆，这里一个样品直接存一个文件，文件名带编号和样品名。
    """
    doc = Document(str(tpl_path))
    mapping = dict(sample)
    mapping["结论"] = make_conclusion(items)

    for para in doc.paragraphs:
        fill_paragraph(para, mapping)
    for row in doc.tables[0].rows:            # 第 1 张表是样品信息，格子里也是占位符
        for cell in row.cells:
            for para in cell.paragraphs:
                fill_paragraph(para, mapping)
    fill_detail_table(doc.tables[1], items)

    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"检验报告_{safe_name(sample[KEY_COL])}_{safe_name(sample['样品名称'])}.docx"
    doc.save(str(out))
    return out


def main() -> bool:
    samples = read_samples(SAMPLE_FILE)
    items = read_items(ITEM_FILE)
    print(f"读到 {len(samples)} 个样品、{sum(len(v) for v in items.values())} 行检项明细")

    made, item_map = {}, {}
    for sid, sample in samples.items():
        rows = sort_items(items.get(sid, []))
        item_map[sid] = rows
        made[sid] = render_one(TPL_FILE, sample, rows, OUT_DIR)
        print(f"  {made[sid].name}  （{len(rows)} 个检项，结论：{make_conclusion(rows)}）")

    # ---------- 以下为校验，正文里没有 ----------
    print("\n===== 回读校验 =====")
    zero_id = [s for s, r in item_map.items() if not r][0]                       # 0 个检项
    most_id = max(item_map, key=lambda s: len(item_map[s]))                      # 检项最多
    pre_id = [s for s in samples if s.startswith("0")][0]                        # 编号带前导零
    slash_id = [s for s in samples if "/" in samples[s]["样品名称"]][0]           # 样品名带斜杠
    bad_id = [s for s, r in item_map.items() if make_conclusion(r) == "不符合规定"][0]
    wait_id = [s for s, r in item_map.items() if r and any(not x["判定"] for x in r)][0]

    def doc_text(sid: str) -> str:
        doc = Document(str(made[sid]))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)

    det_most = Document(str(made[most_id])).tables[1]
    det_zero = Document(str(made[zero_id])).tables[1]
    det_wait = Document(str(made[wait_id])).tables[1]
    row_font = det_most.rows[1].cells[0].paragraphs[0].runs[0].font.size
    wait_result = {r.cells[0].text: r.cells[2].text for r in det_wait.rows[1:]}
    check = [
        ("每个样品出一份 Word（不是一个总文档）",
         len(made) == len(samples) and all(p.exists() for p in made.values())
         and len({p.name for p in made.values()}) == len(made),
         f"{len(made)} 份"),
        ("文件名用样品编号，前导零没丢",
         f"检验报告_{pre_id}_" in made[pre_id].name, made[pre_id].name),
        ("拆进多个 run 的 {{样品编号}} 照样换掉了",
         "{{" not in doc_text(pre_id) and samples[pre_id][KEY_COL] in doc_text(pre_id),
         f"残留占位符 {doc_text(pre_id).count('{{')} 处"),
        ("样品名里的斜杠被换掉，没建出子目录",
         "/" not in made[slash_id].name
         and "阿莫西林-克拉维酸钾片" in made[slash_id].name
         and not any(p.is_dir() for p in OUT_DIR.iterdir()),
         made[slash_id].name),
        ("明细表行数 = 检项数 + 表头",
         len(det_most.rows) == len(item_map[most_id]) + 1,
         f"{len(det_most.rows)} 行 / {len(item_map[most_id])} 个检项"),
        ("检项按药典体例排，不是拼音序",
         [r.cells[0].text for r in det_most.rows[1:4]] == ["性状", "鉴别", "溶解度"],
         [r.cells[0].text for r in det_most.rows[1:4]]),
        ("0 个检项的样品只留表头，没有空占位行",
         len(det_zero.rows) == 1 and "{{" not in doc_text(zero_id),
         f"{len(det_zero.rows)} 行"),
        ("单元格字号守住模板的 10.5pt（没被 cell.text 清掉）",
         row_font == Pt(10.5), row_font),
        ("空结果留空，没被写成 0 或 None",
         wait_result.get("鉴别") == "", repr(wait_result.get("鉴别"))),
        ("结论按结果算：有不符合写不符合规定",
         "不符合规定" in doc_text(bad_id) and make_conclusion(item_map[bad_id]) == "不符合规定",
         bad_id),
        ("结论按结果算：有没出结果的先不下结论",
         "待结果出齐后判定" in doc_text(wait_id), wait_id),
        ("全部合格才写符合规定", "符合规定" in doc_text(pre_id), pre_id),
    ]
    ok = True
    for name, passed, val in check:
        print(f"  [{'OK ' if passed else 'FAIL'}] {name}  ->  {val}")
        ok &= bool(passed)

    print("\n总判定:", "PASS" if ok else "FAIL")
    return bool(ok)


class _Tee:
    """同时写控制台和日志文件，万一控制台中文乱码还有日志可看。"""

    def __init__(self, *streams):
        self.streams = streams

    def write(self, s):
        for st in self.streams:
            try:
                st.write(s)
            except Exception:
                pass

    def flush(self):
        for st in self.streams:
            try:
                st.flush()
            except Exception:
                pass


def _note_unfinished():
    """上一条执行若被手动关窗口打断，日志里会缺完成行。开跑前补一句说明。"""
    LOG_FILE = LOG_DIR / "run_log.txt"
    if not LOG_FILE.exists():
        return
    tail = ""
    with open(LOG_FILE, encoding="utf-8", errors="replace") as f:
        for ln in f:
            if ln.strip():
                tail = ln.strip()
    if tail and not tail.startswith("执行完成"):
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write("（上一条没有留下完成时间，多半是窗口被手动关闭或中途强制中断）\n")


def _host_info() -> str:
    """主机名 + 本机 IP，用来分辨这段日志是哪台机器跑出来的。"""
    try:
        host = socket.gethostname()
    except Exception:
        host = "unknown"
    try:
        # UDP 不会真的发包，只按路由表取本机对外网卡的地址
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:
            s.connect(("8.8.8.8", 80))
            ip = s.getsockname()[0]
    except Exception:
        try:
            ip = socket.gethostbyname(host)
        except Exception:
            ip = "未知"
    return f"{host}（{ip}）"


def _run():
    """一次执行 = 追加一段：开始时间 + 过程 + 完成时间。历史日志不覆盖。"""
    LINE = "=" * 60
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    _note_unfinished()
    _log = open(LOG_DIR / "run_log.txt", "a", encoding="utf-8")   # 追加，别用 w
    _log.write(f"\n{LINE}\n")
    _log.write(f"运行主机：{_host_info()}\n")
    _log.write(f"开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}\n")
    _log.write(f"Python   ：{sys.version.split()[0]}  （{Path(sys.executable).name}）\n")
    _log.flush()

    _orig = sys.stdout
    sys.stdout = _Tee(sys.stdout, _log)
    t0, ok, tail = time.time(), False, ""
    try:
        try:
            ok = main()
        except KeyboardInterrupt:
            tail = "\n被手动中断（Ctrl+C）。"
        except Exception as e:
            import traceback
            traceback.print_exc()
            tail = f"\n跑失败了：{type(e).__name__}: {e}"
        else:
            tail = "\n跑完了。结果在 02_output/，运行记录已写入 source/run_log.txt。"
        if tail:
            print(tail)
        # 完成行必须写在暂停之前：双击后关窗口就丢不了
        print(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}"
              f"  用时 {time.time() - t0:.1f}s  判定 {'PASS' if ok else 'FAIL'}")
    finally:
        sys.stdout.flush()
        _log.close()
        sys.stdout = _orig

    _pause()          # 等回车放在日志收尾之后
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    _run()

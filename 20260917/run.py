# -*- coding: utf-8 -*-
"""引流长文《多中心检验报告批量分发：各中心命名规则不同，只能一份份改》配套脚本。

把一批出好的报告，按中心拆成能直接发出去的分发包：
    报告索引（一行 = 一份报告，带中心、受试者、访视点、日期）
  + 中心分发规则（一行 = 一个中心，命名模板、分包方式、接收人、TAT 约定）
  + 报告原件目录（实际躺在磁盘上的那些文件）
→ 每个中心的报告按本中心的命名模板改名、复制成独立目录、算出 TAT；
  编号位数不对的、原件找不到的、改名后撞名的、日期录错的，一并列出来交给人去核，
  脚本不替人发邮件，也不替人拍板。

核心代码与正文逐字一致，末尾多了自动回读校验，方便确认该标出来的都标出来了。
跨平台：Windows / macOS / Linux 都是一条命令，缺依赖会自动装。
    python run.py
    python run.py --no-pause   # 不暂停（CI 或脚本里用）

打包给别人时，把整个目录拷走即可，路径全部相对定位，无硬编码。
"""
import csv
import shutil
import socket
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

# 控制台编码随环境变，只加容错，不强行改编码。
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(errors="replace")
    except Exception:
        pass

MIRROR = "https://pypi.tuna.tsinghua.edu.cn/simple"
NO_PAUSE = "--no-pause" in sys.argv


def _pause(msg=""):
    """Windows 双击 .py 时留住窗口，避免一闪而过。"""
    if msg:
        print(msg)
    if NO_PAUSE:
        return
    try:
        if sys.platform == "win32" and sys.stdin and sys.stdin.isatty():
            input("按回车退出")
    except Exception:
        pass


def _ensure(*modules_pkgs):
    """缺哪个装哪个。先直接装，失败再退到 --user（无管理员权限的机器）。"""
    for mod, pkg in modules_pkgs:
        try:
            __import__(mod)
            continue
        except ImportError:
            pass
        print(f"缺 {pkg}，正在用清华源安装（解释器: {sys.executable}）...")
        base = [sys.executable, "-m", "pip", "install", "-i", MIRROR, pkg]
        attempts = [base]
        if sys.prefix == sys.base_prefix:      # 不在虚拟环境里，普通安装可能没权限
            attempts.append(base[:4] + ["--user"] + base[4:])
        for cmd in attempts:
            subprocess.run(cmd, check=False)
            try:
                __import__(mod)
                break
            except ImportError:
                continue
        else:
            _pause(f"自动安装失败。手动执行：\n  {sys.executable} -m pip install -i {MIRROR} {pkg}")
            sys.exit(1)


_ensure(("openpyxl", "openpyxl"))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
INDEX_FILE = HERE / "01_raw_data" / "报告索引_系统导出.csv"      # 一行一份报告
RULE_FILE = HERE / "01_raw_data" / "中心分发规则_项目组.csv"     # 一行一个中心
SRC_DIR = HERE / "01_raw_data" / "报告原件"                     # 待分发的报告原件
PACK_DIR = HERE / "02_output" / "分发包"                        # 按中心拆好的包
OUT_FILE = HERE / "02_output" / "分发清单.xlsx"                 # 分发清单 + 报告明细 + 异常
LOG_DIR = HERE / "source"                                      # 运行日志（run_log.txt 就落在 source/ 下）

INDEX_COLS = ["报告编号", "中心编号", "中心名称", "受试者编号", "访视点", "访视编号",
              "样本接收日", "报告日期", "报告文件"]
RULE_COLS = ["中心编号", "中心名称", "命名模板", "分包方式", "接收人", "接收方式",
             "TAT自然日", "编号位数"]
CENTER_COLS = ["中心编号", "中心名称", "报告份数", "待分发", "挂起", "超时", "分包方式",
               "接收人", "接收方式", "TAT约定(自然日)", "最晚报告日", "输出目录", "分发状态"]
ITEM_COLS = ["报告编号", "中心编号", "中心名称", "受试者编号", "访视点", "样本接收日",
             "报告日期", "TAT天数", "超时", "原名", "新文件名", "状态"]
ALERT_COLS = ["中心编号", "中心名称", "报告编号", "类型", "说明"]

# 命名模板里认的占位符。方案上叫「访视点」、系统里叫「访视编号」，
# 各中心的命名规矩又不一样，靠这张表把名字对上，不给每个中心写一段 if/else。
TEMPLATE_FIELDS = {
    "site": "中心编号",
    "subj": "受试者编号",
    "visit": "访视点",
    "date": "报告日期",
}

# 两种分包方式：逐份单发，或者一个受试者的几份报告放一个子目录。
PACK_ONE = "逐份"
PACK_BY_SUBJECT = "按受试者合包"

STATUS_OK = "待分发"
STATUS_HOLD = "挂起"

HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
BAD_FILL = PatternFill("solid", fgColor="FFC7CE")     # 要人拍板的（挂起、超时）

# 异常的类型词。写得具体一点，接的人一眼知道该找谁。
KIND_TPL = "命名模板缺占位符"
KIND_NO_RULE = "中心不在规则表里"
KIND_BAD_SUBJ = "受试者编号位数不对"
KIND_MISSING = "原件缺失"
KIND_EXTRA = "目录里多出来的文件"
KIND_CLASH = "改名后重名"
KIND_NO_DATE = "日期不齐"
KIND_BAD_ORDER = "日期先后反了"
KIND_LATE = "TAT 超时"


def to_int(text: str) -> int:
    """约定天数、编号位数写成空、"待定"、带单位的，一律当 0；0 表示不设这个限制。"""
    try:
        return int(str(text).strip())
    except (TypeError, ValueError):
        return 0


def to_date(text: str):
    """日期列写空、写"待出报告"的都当没有，返回 None。算 TAT 之前必须先判这个。"""
    text = (text or "").strip()
    if not text:
        return None
    try:
        return datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError:
        return None


def safe_name(text: str) -> str:
    """把文件名里的非法字符换掉。

    受试者编号带斜杠是常事（筛选失败补号 0102/2），落成文件名在 Windows 上直接建不了，
    整个中心的分发包就卡在这一份上。斜杠换短横线，其余非法字符换下划线。
    """
    out = text.strip().replace("/", "-").replace("\\", "-")
    for ch in ':*?"<>|':
        out = out.replace(ch, "_")
    return out


def alert(code: str, name: str, no: str, kind: str, note: str) -> dict:
    """一条异常记录。报告级的带上报告编号，中心级或目录级的留空。"""
    return {"中心编号": code, "中心名称": name, "报告编号": no, "类型": kind, "说明": note}


def read_rules(path: Path):
    """读中心分发规则，一行一个中心。

    命名模板必须带 {subj}：漏了的话同一个中心所有报告会拼成同一个文件名，整批互相覆盖。
    这种规则不能凑合往下跑，一发现就把这个中心标出来，它的报告整批挂起。
    表尾合计行落在中心编号为空的那一行，直接丢。
    """
    rules, alerts = {}, []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            code = (row.get("中心编号") or "").strip()
            if not code:
                continue
            rec = {c: (row.get(c) or "").strip() for c in RULE_COLS}
            rec["模板可用"] = "{subj}" in rec["命名模板"]
            if not rec["模板可用"]:
                alerts.append(alert(code, rec["中心名称"], "", KIND_TPL,
                                    f"命名模板「{rec['命名模板']}」里没有 {{subj}}，"
                                    f"整批报告会拼成同一个文件名，这个中心先整批挂起"))
            rules[code] = rec
    return rules, alerts


def read_index(path: Path):
    """读报告索引，一行一份报告。报告编号非空才算数，表尾的合计行与说明行落在这儿。"""
    out = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if not (row.get("报告编号") or "").strip():
                continue
            out.append({c: (row.get(c) or "").strip() for c in INDEX_COLS})
    return out


def scan_sources(folder: Path) -> set:
    """扫一遍报告原件目录，只取文件名。

    索引里有、目录里没有 = 这份报告根本没拿到；目录里有、索引里没有 = 上一批的残留。
    两种都要挑出来：前者发不出去，后者混进包里就是发错报告。
    """
    if not folder.exists():
        return set()
    return {p.name for p in folder.iterdir() if p.is_file()}


def tat_days(received: str, reported: str):
    """TAT = 报告日期 - 样本接收日，按自然日算。

    只数工作日会好看很多，但中心等的是自然日：周五收到的样本，周一出报告就是 3 天。
    两边有一个是空就算不出来，返回 None —— 不能当 0 天，0 天看起来像当天就出了报告。
    """
    a, b = to_date(received), to_date(reported)
    if a is None or b is None:
        return None
    return (b - a).days


def make_name(rule: dict, rec: dict) -> str:
    """按这个中心的命名模板拼文件名。

    模板写成 {site}_{subj}_{visit} 这种占位符，而不是给每个中心写一段 if/else：
    中心的命名规矩三天两头改，改模板是改一行 CSV，改代码得重新发一遍脚本。
    后缀跟着原件走，拼完再过一遍非法字符。
    """
    name = rule["命名模板"]
    for key, col in TEMPLATE_FIELDS.items():
        name = name.replace("{" + key + "}", rec.get(col, ""))
    return safe_name(name) + Path(rec["报告文件"]).suffix


def plan_center(rule: dict, reports: list, on_disk: set) -> dict:
    """一个中心：逐份改名、查重名、算 TAT。

    重名不能覆盖 —— 同一受试者同一访视点重出的报告，算出来是同一个文件名，
    覆盖等于丢一份，稽查时说不清哪一版是发出去的。发现重名，两份都挂起，交给人去核。
    """
    code, name = rule["中心编号"], rule["中心名称"]
    tpl_ok = rule["模板可用"]
    limit = to_int(rule["TAT自然日"])
    width = to_int(rule["编号位数"])
    items, alerts, seen = [], [], {}

    for rec in reports:
        item = {**rec, "TAT天数": "", "超时": "", "原名": rec["报告文件"],
                "新文件名": "", "状态": STATUS_OK}
        item["新文件名"] = make_name(rule, rec) if tpl_ok else ""

        days = tat_days(rec["样本接收日"], rec["报告日期"])
        if days is None:
            item["状态"] = STATUS_HOLD
            alerts.append(alert(code, name, rec["报告编号"], KIND_NO_DATE,
                                f"样本接收日「{rec['样本接收日'] or '空'}」、"
                                f"报告日期「{rec['报告日期'] or '空'}」，TAT 算不出来，先挂着"))
        elif days < 0:
            item["状态"] = STATUS_HOLD
            alerts.append(alert(code, name, rec["报告编号"], KIND_BAD_ORDER,
                                f"报告日期比样本接收日还早 {abs(days)} 天，日期多半录错了"))
        else:
            item["TAT天数"] = days
            if limit and days > limit:
                item["超时"] = "是"
                alerts.append(alert(code, name, rec["报告编号"], KIND_LATE,
                                    f"TAT {days} 天，超过这个中心约定的 {limit} 天"))

        if not tpl_ok:
            item["状态"] = STATUS_HOLD
        # 编号位数只校短不校长：补号本来就比常规编号长（0102/2），按"不等于"去卡会把正常的补号也拦下来。
        subj = rec["受试者编号"]
        if subj and width and len(subj) < width:
            item["状态"] = STATUS_HOLD
            alerts.append(alert(code, name, rec["报告编号"], KIND_BAD_SUBJ,
                                f"受试者编号「{subj}」只有 {len(subj)} 位，规则要求 {width} 位，"
                                f"多半是导出时前导零被吃掉，回去重导"))
        if rec["报告文件"] not in on_disk:
            item["状态"] = STATUS_HOLD
            alerts.append(alert(code, name, rec["报告编号"], KIND_MISSING,
                                f"索引里有这份报告，原件目录里没有「{rec['报告文件']}」，发不出去"))
        if tpl_ok and item["新文件名"] in seen:
            item["状态"] = STATUS_HOLD
            seen[item["新文件名"]]["状态"] = STATUS_HOLD
            alerts.append(alert(code, name, rec["报告编号"], KIND_CLASH,
                                f"和 {seen[item['新文件名']]['报告编号']} 改名后是同一个文件"
                                f"「{item['新文件名']}」，覆盖就是丢一份，两份都先挂起"))
        elif tpl_ok:
            seen[item["新文件名"]] = item
        items.append(item)

    return {"rule": rule, "items": items, "alerts": alerts}


def copy_package(plans: list, src_dir: Path, pack_dir: Path) -> int:
    """把改名后的报告复制成分发包：一个中心一个目录。

    挂起的一份都不复制 —— 包里的每一份都得是能直接发出去的。
    按受试者合包的中心再套一层受试者目录：现场是一个受试者一袋，拆包时不用在几百份里翻。
    """
    if pack_dir.exists():
        shutil.rmtree(pack_dir)                        # 重跑先清空，免得上一批的旧文件混进去
    copied = 0
    for plan in plans:
        rule = plan["rule"]
        folder = pack_dir / safe_name(f"{rule['中心编号']}_{rule['中心名称']}")
        for it in plan["items"]:
            if it["状态"] != STATUS_OK:
                continue
            target = folder / it["新文件名"]
            if rule["分包方式"] == PACK_BY_SUBJECT:
                target = folder / safe_name(it["受试者编号"]) / it["新文件名"]
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src_dir / it["报告文件"], target)
            copied += 1
    return copied


def write_workbook(plans: list, items: list, alerts: list, out_path: Path) -> Path:
    """三个页签写进一个 xlsx。

    分发清单给做分发的人，一行一个中心，照着它发邮件或传门户；
    报告明细是每一份报告改名前后的对照，名字对不上时拿它去查；
    异常与提醒给项目组 —— 要重导编号、去找原件、跟超时报告的事都在最后一页。
    """
    rows = []
    for plan in plans:
        rule, its = plan["rule"], plan["items"]
        ready = [it for it in its if it["状态"] == STATUS_OK]
        dates = [it["报告日期"] for it in its if it["报告日期"]]
        rows.append([
            rule["中心编号"], rule["中心名称"], len(its), len(ready),
            len([it for it in its if it["状态"] == STATUS_HOLD]),
            len([it for it in ready if it["超时"] == "是"]),
            rule["分包方式"], rule["接收人"], rule["接收方式"], rule["TAT自然日"],
            max(dates) if dates else "",
            f"{pack_dir_name()}/{safe_name(rule['中心编号'] + '_' + rule['中心名称'])}",
            STATUS_OK,
        ])

    wb = Workbook()
    sheets = [("分发清单", CENTER_COLS, rows), ("报告明细", ITEM_COLS, items),
              ("异常与提醒", ALERT_COLS, alerts)]
    for idx, (title, cols, data) in enumerate(sheets):
        ws = wb.active if idx == 0 else wb.create_sheet()
        ws.title = title
        ws.append(cols)
        for c in ws[1]:
            c.fill, c.font = HEAD_FILL, Font(bold=True)
        for row in data:
            if isinstance(row, dict):
                ws.append([row.get(c, "") for c in cols])
                if row.get("状态") == STATUS_HOLD or row.get("超时") == "是":
                    for c in ws[ws.max_row]:
                        c.fill = BAD_FILL
            else:
                ws.append(row)
        for i, col in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(max(len(col) + 4, 10), 30)
        ws.freeze_panes = "A2"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def pack_dir_name() -> str:
    """分发清单里写给人看的相对路径，不带本机绝对路径。"""
    return PACK_DIR.name


def main() -> bool:
    rules, alerts = read_rules(RULE_FILE)
    index = read_index(INDEX_FILE)
    on_disk = scan_sources(SRC_DIR)
    print(f"分发规则 {len(rules)} 个中心、报告索引 {len(index)} 份、原件目录 {len(on_disk)} 个文件")

    by_center, plans, items = {}, [], []
    for rec in index:
        if rec["中心编号"] not in rules:
            alerts.append(alert(rec["中心编号"], rec["中心名称"], rec["报告编号"], KIND_NO_RULE,
                                "中心分发规则表里没有这个中心，这份报告没人认领，先去补规则"))
            continue
        by_center.setdefault(rec["中心编号"], []).append(rec)

    for code in sorted(by_center):
        plan = plan_center(rules[code], by_center[code], on_disk)
        plans.append(plan)
        items += plan["items"]
        alerts += plan["alerts"]

    known = {rec["报告文件"] for rec in index}
    for name in sorted(on_disk - known):
        alerts.append(alert("", "", "", KIND_EXTRA, f"原件目录里这份文件索引里没有：{name}"))

    copied = copy_package(plans, SRC_DIR, PACK_DIR)
    write_workbook(plans, items, alerts, OUT_FILE)

    ready = [it for it in items if it["状态"] == STATUS_OK]
    hold = [it for it in items if it["状态"] == STATUS_HOLD]
    late = [it for it in items if it["超时"] == "是"]
    print(f"  待分发 {len(ready)} 份，挂起 {len(hold)} 份，其中 TAT 超时 {len(late)} 份")
    print(f"  已复制进分发包 {copied} 份；异常与提醒 {len(alerts)} 条")
    print(f"分发清单已生成：{OUT_FILE.name}")

    return verify(plans, items, alerts, copied)


def verify(plans: list, items: list, alerts: list, copied: int) -> bool:
    """回读校验。这些数字是照着造好的边界数据定的，对不上说明逻辑变了或数据被改过。"""
    from openpyxl import load_workbook
    wb = load_workbook(OUT_FILE)
    centers = list(wb["分发清单"].iter_rows(min_row=2, values_only=True))
    detail = list(wb["报告明细"].iter_rows(min_row=2, values_only=True))
    trouble = list(wb["异常与提醒"].iter_rows(min_row=2, values_only=True))
    kinds = {}
    for r in trouble:
        kinds[r[3]] = kinds.get(r[3], 0) + 1

    s01 = sorted(p.name for p in (PACK_DIR / "S01_北京中心").iterdir())
    s02 = sorted(p.name for p in (PACK_DIR / "S02_上海中心").iterdir())
    s03 = sorted(p.name for p in (PACK_DIR / "S03_广州中心").iterdir() if p.is_dir())
    held = [it for it in items if it["状态"] == STATUS_HOLD]

    checks = [
        ("分发清单 5 个中心", len(centers) == 5, f"{len(centers)} 个"),
        ("报告明细 23 份（无规则那 1 份不进明细）", len(detail) == 23, f"{len(detail)} 份"),
        ("异常与提醒 11 条", len(trouble) == 11, f"{len(trouble)} 条"),
        ("分发包共 15 个文件（挂起的 8 份没进去）", copied == 15, f"{copied} 个"),
        ("S01 只放了 4 份（重名与编号异常的都挂起）", len(s01) == 4, f"{len(s01)} 个"),
        ("S02 里斜杠编号改成了短横线",
         "S02-0204-2-V1-2026-08-14.txt" in s02,
         "、".join(s02) if len(s02) != 5 else "S02-0204-2-V1-2026-08-14.txt"),
        ("S03 按受试者分了 3 个子目录", len(s03) == 3, "、".join(s03)),
        ("挂起的 8 份都写了原因", len(held) == 8, f"{len(held)} 份"),
        ("TAT 超时 3 条", kinds.get(KIND_LATE, 0) == 3, f"{kinds.get(KIND_LATE, 0)} 条"),
        ("改名后重名 1 条（两份都挂起）", kinds.get(KIND_CLASH, 0) == 1, f"{kinds.get(KIND_CLASH, 0)} 条"),
        ("目录里多出来的文件 1 条", kinds.get(KIND_EXTRA, 0) == 1, f"{kinds.get(KIND_EXTRA, 0)} 条"),
    ]
    ok = True
    for name, passed, val in checks:
        print(f"  [{'OK ' if passed else 'FAIL'}] {name}  ->  {val}")
        ok &= bool(passed)

    print("\n总判定:", "PASS" if ok else "FAIL")
    return bool(ok)


class _Tee:
    """同时写控制台和日志文件，万一控制台中文乱码还有日志可看。"""

    def __init__(self, *streams):
        self.streams = streams

    def write(self, s):
        for st in self.streams:
            try:
                st.write(s)
            except Exception:
                pass

    def flush(self):
        for st in self.streams:
            try:
                st.flush()
            except Exception:
                pass


def _note_unfinished():
    """上一条执行若被手动关窗口打断，日志里会缺完成行。开跑前补一句说明。"""
    LOG_FILE = LOG_DIR / "run_log.txt"
    if not LOG_FILE.exists():
        return
    tail = ""
    with open(LOG_FILE, encoding="utf-8", errors="replace") as f:
        for ln in f:
            if ln.strip():
                tail = ln.strip()
    if tail and not tail.startswith("执行完成"):
        with open(LOG_FILE, "a", encoding="utf-8") as f:
            f.write("（上一条没有留下完成时间，多半是窗口被手动关闭或中途强制中断）\n")


def _host_info() -> str:
    """主机名 + 本机 IP，用来分辨这段日志是哪台机器跑出来的。"""
    try:
        host = socket.gethostname()
    except Exception:
        host = "unknown"
    try:
        # UDP 不会真的发包，只按路由表取本机对外网卡的地址
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:
            s.connect(("8.8.8.8", 80))
            ip = s.getsockname()[0]
    except Exception:
        try:
            ip = socket.gethostbyname(host)
        except Exception:
            ip = "未知"
    return f"{host}（{ip}）"


def _run():
    """一次执行 = 追加一段：开始时间 + 过程 + 完成时间。历史日志不覆盖。"""
    LINE = "=" * 60
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    _note_unfinished()
    _log = open(LOG_DIR / "run_log.txt", "a", encoding="utf-8")   # 追加，别用 w
    _log.write(f"\n{LINE}\n")
    _log.write(f"运行主机：{_host_info()}\n")
    _log.write(f"开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}\n")
    _log.write(f"Python   ：{sys.version.split()[0]}  （{Path(sys.executable).name}）\n")
    _log.flush()

    _orig = sys.stdout
    sys.stdout = _Tee(sys.stdout, _log)
    t0, ok, tail = time.time(), False, ""
    try:
        try:
            ok = main()
        except KeyboardInterrupt:
            tail = "\n被手动中断（Ctrl+C）。"
        except Exception as e:
            import traceback
            traceback.print_exc()
            tail = f"\n跑失败了：{type(e).__name__}: {e}"
        else:
            tail = "\n跑完了。结果在 02_output/，运行记录已写入 source/run_log.txt。"
        if tail:
            print(tail)
        # 完成行必须写在暂停之前：双击后关窗口就丢不了
        print(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}"
              f"  用时 {time.time() - t0:.1f}s  判定 {'PASS' if ok else 'FAIL'}")
    finally:
        sys.stdout.flush()
        _log.close()
        sys.stdout = _orig

    _pause()          # 等回车放在日志收尾之后
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    _run()

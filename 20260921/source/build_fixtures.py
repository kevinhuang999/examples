#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""重造本示例用到的模板与数据（跑坏了随时重来一遍）。

上半个文件：建「线上接收单」的空白模板（source/templates/接收单模板.xlsx）
下半个文件：造两份 CSV 原始数据（01_raw_data/ 下：纸质单人工补录 + 扫码导出明细）

用法：python source/build_fixtures.py
"""

import csv
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
TPL_DIR = ROOT / "source" / "templates"
RAW_DIR = ROOT / "01_raw_data"
TPL_FILE = TPL_DIR / "接收单模板.xlsx"
PAPER_CSV = RAW_DIR / "纸质接收单_人工补录.csv"
DETAIL_CSV = RAW_DIR / "扫码导出_样本明细.csv"


# ---------------------------------------------------------------- 上半：建模板

# 模板左边的栏目名。脚本跑的时候按这个名字反查行号，不写死坐标。
FORM_LABELS = [
    "线上单号", "运单号", "发运中心", "到达日期", "到达时间", "接收人",
    "箱数", "管数", "运输温度", "样本类型分布", "纸面对账", "核对结论", "确认状态", "生成时间",
]

FOOTER_NOTE = "本单由脚本按纸质接收单与扫码导出明细自动生成；温度仅作记录，不作合规判定。"


def build_template():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "接收单"

    thin = Side(style="thin", color="BFBFBF")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    # ★值格必须先画上样式：openpyxl 读回时"没内容也没样式"的格子根本不存在，
    #   脚本按标签定位后往空格子写，格子没有底纹也不影响；但模板肉眼要看得见在哪填。
    fill = PatternFill("solid", fgColor="F2F6FB")

    ws["A1"] = "样 本 接 收 单"
    ws["A1"].font = Font(size=16, bold=True, color="1F3864")
    ws.merge_cells("A1:B1")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws["A2"] = "Sample Receipt Form"
    ws["A2"].font = Font(size=9, color="8C8C8C")
    ws.merge_cells("A2:B2")
    ws["A2"].alignment = Alignment(horizontal="center")

    for i, label in enumerate(FORM_LABELS):
        r = 4 + i
        ws.cell(row=r, column=1, value=label).font = Font(bold=True, color="404040")
        cell = ws.cell(row=r, column=2)
        cell.fill = fill
        cell.border = box
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[r].height = 20

    note_row = 4 + len(FORM_LABELS) + 1
    ws.cell(row=note_row, column=1, value=FOOTER_NOTE).font = Font(size=8, color="8C8C8C")

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 46

    TPL_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(TPL_FILE)
    print(f"[模板] {TPL_FILE.relative_to(ROOT)}  {len(FORM_LABELS)} 个栏目")


# ---------------------------------------------------------------- 下半：造数据

# 一行 = 一批送到实验室的货，也就是手里的一张纸质接收单。
# 这些字段是从纸面上抄下来的：只有纸面上才有，扫码文件里没有。
PAPER_ROWS = [
    # 运单号,    发运中心,  到达日期,      到达时间, 纸质箱数, 纸质管数, 温上限, 温下限, 接收人
    ("0000712", "上海中心", "2026-09-14", "08:20", "2", "12", "7.8", "2.5", "王敏"),
    ("0000713", "北京中心", "2026-09-14", "09:05", "2", "8", "7.9", "2.4", "李强"),   # 纸质箱数抄多 1
    ("0000714", "广州中心", "2026-09-14", "10:30", "1", "10", "8.4", "2.6", "陈静"),   # 纸质管数抄多 1
    ("0000715", "上海中心", "2026-09-15", "08:10", "3", "60", "7.6", "2.5", "王敏"),   # 一单 60 管
    ("0000716", "北京中心", "2026-09-15", "09:40", "1", "1", "", "", "李强"),          # 1 管 + 温度未记录
    ("0000718", "广州中心", "2026-09-15", "11:15", "2", "14", "7.5", "2.3", "陈静"),   # 明细里查不到
    ("0000719", "上海中心", "", "07:50", "1", "6", "8.0", "2.7", "王敏"),              # 到达日期空
    ("0000720", "北京中心", "2026-09-16", "08:55", "2", "16", "7.7", "2.5", ""),       # 接收人空
]

# 扫码文件里实际有多少管（故意跟纸质单上的数字错开）
DETAIL_COUNT = {
    "0000712": 12,
    "0000713": 8,
    "0000714": 9,     # 纸质写 10，实际 9
    "0000715": 60,
    "0000716": 1,
    # 0000718 整单缺失
    "0000719": 6,
    "0000720": 16,
}

SAMPLE_TYPES = ["血清", "血浆", "全血", "尿液"]
DETAIL_HEADER = ["运单号", "箱号", "样本编号", "样本类型", "采集中心"]
PAPER_HEADER = ["运单号", "发运中心", "到达日期", "到达时间",
                "纸质登记箱数", "纸质登记管数", "运输温度上限", "运输温度下限", "接收人"]


def build_csv():
    RAW_DIR.mkdir(parents=True, exist_ok=True)

    with open(PAPER_CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(PAPER_HEADER)
        for row in PAPER_ROWS:
            w.writerow(row)
        # ★表尾说明行：第一列（运单号）必须留空，否则会被当成一张真单子
        w.writerow(["", "", "", "", "", "", "", "", "备注：本批为 9 月纸质接收单存量补录（共 8 张）"])

    detail_rows = []
    for waybill, count in DETAIL_COUNT.items():
        boxes = 3 if waybill == "0000715" else (2 if waybill in ("0000712", "0000720") else 1)
        for i in range(count):
            box = (i % boxes) + 1 if boxes > 1 else 1
            stype = SAMPLE_TYPES[i % len(SAMPLE_TYPES)]
            if waybill == "0000714" and i < 2:
                stype = ""            # 样本类型漏填 2 管
            detail_rows.append([
                waybill, str(box), f"{waybill}-{i + 1:03d}", stype,
                "上海" if waybill in ("0000712", "0000715", "0000719") else "北京",
            ])
    detail_rows.sort(key=lambda r: (r[0], r[2]))

    with open(DETAIL_CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(DETAIL_HEADER)
        w.writerows(detail_rows)
        # ★表尾合计行：第一列（运单号）留空，合计数字放在"样本类型"列，别去污染样本编号
        w.writerow(["", "", "", f"合计 {len(detail_rows)} 管", ""])

    print(f"[数据] {PAPER_CSV.relative_to(ROOT)}  {len(PAPER_ROWS)} 张纸质单")
    print(f"[数据] {DETAIL_CSV.relative_to(ROOT)}  {len(detail_rows)} 管明细")


if __name__ == "__main__":
    build_template()
    build_csv()
    print("模板与数据已重造完成。")

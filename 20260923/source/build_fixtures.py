#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""重造本示例的模板与数据（上半建模板，下半造 CSV）。

上半：建 `source/templates/检测委托单模板.xlsx`——单头是「标签 + 取值」两列一组，
明细区固定 8 行、**每行都画了边框**（不画边框，openpyxl 读回时这些行根本不存在）。
下半：建 `01_raw_data/委托受理登记.csv`——前台按客户委托原件登记的原始表，
故意留了日期格式不统一、客户名夹全角空格、数量写成「2 支」、缺依据标准、
明细 9 行、客户名称空着、受理日期写「待定」这些边界。

用法：python source/build_fixtures.py
"""

from csv import DictWriter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

ROOT = Path(__file__).resolve().parent.parent      # 包根 = build_fixtures.py 的上上层
TPL = ROOT / "source" / "templates" / "检测委托单模板.xlsx"
CSV = ROOT / "01_raw_data" / "委托受理登记.csv"

# 单头标签按「列 A/C/E 放标签、列 B/D/F 放取值」排；明细区是 7 列，多出一列 G 留给表头合并
HEAD_FIELDS = [("委托单号", 2), ("受理日期", 2), ("加急", 2),
               ("客户名称", 3), ("联系人", 3), ("送样方式", 3),
               ("报告交付方式", 4), ("样品数", 4), ("检测项目数", 4)]
DETAIL_HEAD = ["序号", "样品名称", "客户样品编号", "数量", "检测项目", "依据标准", "备注"]
DETAIL_ROWS = 8
WIDTHS = {"A": 18, "B": 22, "C": 14, "D": 16, "E": 14, "F": 18, "G": 16}

THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_template():
    """上半：建空白委托单模板。模板由质量部备案，脚本只能填、不能改结构。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "检测委托单"

    ws["A1"] = "检 测 委 托 单"
    ws.merge_cells("A1:G1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    for label, row in HEAD_FIELDS:
        col = 1 if label in ("委托单号", "客户名称", "报告交付方式") else \
              3 if label in ("受理日期", "联系人", "样品数") else 5
        cell = ws.cell(row=row, column=col, value=label)
        cell.font = Font(bold=True)
        ws.cell(row=row, column=col + 1).border = BOX

    ws["A5"] = "备注"
    ws["A5"].font = Font(bold=True)
    ws.merge_cells("B5:G5")
    ws["B5"].border = BOX

    head = 6
    for i, name in enumerate(DETAIL_HEAD, 1):
        cell = ws.cell(row=head, column=i, value=name)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = BOX

    # 明细区空行必须画边框：不画的话这些行在 openpyxl 眼里"不存在"，
    # 回读时 ws.max_row 只到表头那一行，靠 max_row 反推明细区的写法会一行都找不到。
    for r in range(head + 1, head + 1 + DETAIL_ROWS):
        for c in range(1, 8):
            ws.cell(row=r, column=c).border = BOX

    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width

    TPL.parent.mkdir(parents=True, exist_ok=True)
    wb.save(TPL)
    print(f"模板已生成：{TPL}")


def build_csv():
    """下半：造受理登记表——一行一个样品；同一份委托单的多行靠「受理日期 + 客户名称」归堆。"""
    head = ["受理日期", "客户名称", "客户联系人", "送样方式", "报告交付方式", "加急",
            "样品名称", "客户样品编号", "样品数量", "检测项目", "依据标准", "备注"]
    hy, xy, qh = "浙江宏远纺织有限公司", "苏州新亚电子材料有限公司", "上海启海化工有限公司"
    rows = []

    def add(day, cust, who, send, back, rush, name, scode, qty, items, std, note=""):
        rows.append([day, cust, who, send, back, rush, name, scode, qty, items, std, note])

    # ① 一张单三个样品：同一天同一客户，靠日期 + 客户名归成一张
    add("2026-09-18", hy, "李工", "自送", "电子+纸质", "", "全棉染色布", "0000713", "2",
        "甲醛含量、pH值、异味", "GB 18401-2010", "急件样品已拆封拍照")
    add("2026-09-18", hy, "李工", "自送", "电子+纸质", "", "涤纶针织布", "0000714", "1",
        "耐摩擦色牢度/耐水色牢度", "GB/T 3920-2008", "")
    add("2026-09-18", hy, "李工", "自送", "电子+纸质", "", "涤棉混纺布", "HR-2026-088", "3",
        "纤维含量", "GB/T 2910.1-2009", "客户自编编号，不是本机构格式")

    # ② 同一天同一客户，日期写成三种格式、客户名夹了全角空格——应归成**同一张**单
    add("2026-09-18", xy, "王女士", "寄样", "电子", "", "覆铜板", "XYA-001", "4",
        "剥离强度、耐浸焊性", "IPC-TM-650 2.4.9", "")
    add("2026/9/18", xy + "\u3000", "王女士", "寄样", "电子", "", "环氧树脂", "XYA-002", "2",
        "玻璃化转变温度", "IPC-TM-650 2.4.24", "后半批补寄，同一委托")
    add("2026.09.18", " " + xy + " ", "王女士", "寄样", "电子", "", "铜箔", "XYA-003", "1",
        "抗拉强度、延伸率", "GB/T 5230-2020", "")

    # ③ 加急只标在其中一行的行上——单头该写什么，要按单级规则定
    add("2026-09-18", qh, "张经理", "上门取样", "电子", "加急", "废水水样 A", "W-0918-A", "1",
        "pH、化学需氧量、氨氮", "HJ 828-2017", "加急，48 小时出报告")
    add("2026-09-18", qh, "张经理", "上门取样", "电子", "", "废水水样 B", "W-0918-B", "2",
        "pH、总磷", "HJ 828-2017", "")

    # ④ 同一客户隔天再来一次——必须是**另一张**单（日期不进分组键就会并成一张）
    add("2026-09-19", hy, "李工", "自送", "电子+纸质", "", "家纺面料", "0000721", "2",
        "甲醛含量、pH值", "GB 18401-2010", "")
    add("2026-09-19", hy, "李工", "自送", "电子+纸质", "", "装饰面料", "0000722", "1",
        "阻燃性能", "GB 8624-2012", "追加委托")

    # ⑤ 六种要挂起的情形
    add("2026-09-19", "宁波航海紧固件有限公司", "陈工", "寄样", "纸质", "", "螺栓 M12", "HB-0921", "1",
        "抗拉强度", "", "客户没写标准号，已催")
    add("2026-09-19", "台州欣达塑业有限公司", "赵工", "寄样", "电子", "", "塑料管件", "XD-0921-1", "2 支",
        "静液压强度", "GB/T 6111-2018", "客户口头说 2 支，原件上也是这么写的")
    add("2026-09-20", "杭州锦泰包装有限公司", "周工", "寄样", "电子", "", "瓦楞纸箱", "JT-0920-1", "5",
        "", "GB/T 6543-2008", "项目待客户确认")
    add("2026-09-20", "", "", "自送", "纸质", "", "土壤样品", "", "1",
        "pH、重金属", "HJ 803-2016", "前台登记时漏填客户名称")
    for i in range(1, 10):
        add("2026-09-20", "山东鲁南建材有限公司", "孙工", "自送", "电子", "", f"水泥试样 {i:02d}",
            f"LN-{i:02d}", "1", "抗压强度", "GB/T 17671-2021", "")
    add("待定", "常州华通金属制品有限公司", "吴工", "寄样", "电子", "", "铝合金型材", "HT-0923", "1",
        "化学成分", "GB/T 7999-2015", "受理日期待业务确认")

    # 表尾说明行：第一列（受理日期）留空，读的时候整行跳过。
    # 也正因为这样，说明文字千万别写到第一列——写进去就会被当成一份委托。
    rows.append(["", "说明：本表由前台按客户委托原件登记，一行一个样品；"
                 "同一份委托单的多行样品，受理日期与客户名称要填成一致。", "", "", "", "",
                 "", "", "", "", "", ""])

    CSV.parent.mkdir(parents=True, exist_ok=True)
    with open(CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = DictWriter(f, fieldnames=head)
        w.writeheader()
        for r in rows:
            w.writerow(dict(zip(head, r)))
    print(f"登记表已生成：{CSV}  数据 {len(rows) - 1} 行 + 表尾说明 1 行")


if __name__ == "__main__":
    build_template()
    build_csv()

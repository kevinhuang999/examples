#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""重造 20260927 示例的输入数据与模板。

上半：建模板（source/templates/盖章台账模板.xlsx，两个工作表）。
下半：造数据——盖章申请单.csv、用章规则表.csv、印章图/*.png、报告原件/*.pdf。

跑一次全部重建：
    python source/build_fixtures.py
正常跑 run.py 不需要跑它。
"""

import csv
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "01_raw_data"
TPL = ROOT / "source" / "templates"
STAMP_DIR = RAW / "印章图"
PDF_DIR = RAW / "报告原件"

LEDGER_HEAD = ["报告编号", "客户名称", "报告类型", "章名", "落位", "页码", "盖章时间", "操作人", "结论"]
HOLD_HEAD = ["报告编号", "客户名称", "报告类型", "挂起原因"]

# 报告原件页数。JC2026-1009 故意不建文件，用来验证「原件没到」那一支。
PAGES = {
    "JC2026-1001": 6, "JC2026-1002": 3, "0000122": 4, "JC2026-1004": 1,
    "JC2026-1005": 5, "JC2026-1006": 2, "JC2026-1007": 2, "JC2026-1008": 2,
    "JC2026-1010": 2, "JC2026-1011": 3, "JC2026-1012": 3, "JC2026-1013": 1,
}

# 申请单：申请日期,报告编号,客户名称,报告类型,审签状态,报告文件,申请人,备注
REQUESTS = [
    ("2026-09-27", "JC2026-1001", "启海环境检测", "委托检测报告", "已审签", "报告_JC2026-1001.pdf", "周慧", ""),
    ("2026-09-27", "JC2026-1002", "海门建材", "委托检测报告", "已审签", "报告_JC2026-1002.pdf", "周慧", ""),
    ("2026-09-27", "0000122", "宏远食品", "委托检测报告", "已审签", "报告_0000122.pdf", "周慧", "老系统编号，前导零照原样"),
    ("2026-09-27", "JC2026-1004", "启东土壤", "内部质控报告", "已审签", "报告_JC2026-1004.pdf", "李昶", ""),
    ("2026-09-27", "JC2026-1005", "崇明水务", "加急委托报告", "已审签", "报告_JC2026-1005.pdf", "李昶", "客户催，今天必须发"),
    ("2026-09-27", "JC2026-1006", "苏通包装", "涉外检测报告", "已审签", "报告_JC2026-1006.pdf", "周慧", ""),
    ("2026-09-27", "JC2026-1007", "江海水质", "委托检测报告", "审核中", "报告_JC2026-1007.pdf", "李昶", "技术负责人还没签"),
    ("2026-09-27", "JC2026-1008", "欣荣化工", "委托检测报告", "已作废", "报告_JC2026-1008.pdf", "周慧", "客户改了委托项目，这份作废重出"),
    ("2026-09-27", "JC2026-1009", "通州金属", "委托检测报告", "已审签", "报告_JC2026-1009.pdf", "周慧", ""),
    ("2026-09-27", "JC2026-1010", "启海环境检测", "委托检测报告", "已审签", "报告_JC2026-1010.pdf", "周慧", ""),
    ("2026-09-27", "JC2026-1011", "大丰农科", "年度评估报告", "已审签", "报告_JC2026-1011.pdf", "李昶", "这个类型今年刚开"),
    ("2026-09-27", "JC2026-1012", "金鼎检测\u3000", "委托检测报告", "已审签", "报告_JC2026-1012.pdf", "周慧", "客户名后面带了个全角空格"),
    ("2026-09-27", "JC2026-1013", "海隆电子", "委托检测报告", "已审签", "报告_JC2026-1013.pdf", "周慧", "只有一页"),
    ("2026-09-27", "JC2026-1001", "启海环境检测", "委托检测报告", "已审签", "报告_JC2026-1001.pdf", "李昶", "催得急，重复报了一次"),
    ("", "", "", "", "", "", "", "以下空白"),
]

# 用章规则表：报告类型,章名,落位,骑缝起始页
RULES = [
    ("委托检测报告", "检测专用章", "首页右下", ""),
    ("委托检测报告", "资质认定章", "末页右下", ""),
    ("委托检测报告", "骑缝章", "右侧骑缝", "2"),
    ("加急委托报告", "检测专用章", "首页右下", ""),
    ("加急委托报告", "骑缝章", "右侧骑缝", "2"),
    ("内部质控报告", "检测专用章", "首页右下", ""),
    ("涉外检测报告", "检测专用章", "首页右下", ""),
    ("涉外检测报告", "涉外认证章", "首页右下", ""),
    ("", "", "", "表格到此，下面没有"),
]

# 章图：四枚里只造三枚，「涉外认证章」故意不造 —— 申请单里那份涉外报告会挂起。
STAMP_KINDS = {
    "检测专用章": "double",
    "资质认定章": "ring",
    "骑缝章": "hatch",
}


def build_template():
    """① A4 台账模板：表头 + 边框 + 列宽，编号列一律文本格式（前导零保得住）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "盖章台账"
    thin = Side(style="thin", color="8C8C8C")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="E8EEF7")
    for i, name in enumerate(LEDGER_HEAD, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(name="微软雅黑", size=10, bold=True)
        c.fill = fill
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, w in enumerate([14, 16, 14, 12, 12, 10, 20, 10, 16], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.cell(row=2, column=1).number_format = "@"          # 报告编号 0000122 走文本
    ws.cell(row=2, column=6).number_format = "@"          # 页码可能是 «2-6» 这种区间

    hs = wb.create_sheet("挂起清单")
    for i, name in enumerate(HOLD_HEAD, 1):
        c = hs.cell(row=1, column=i, value=name)
        c.font = Font(name="微软雅黑", size=10, bold=True)
        c.fill = fill
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, w in enumerate([14, 16, 14, 44], 1):
        hs.column_dimensions[hs.cell(row=1, column=i).column_letter].width = w

    TPL.mkdir(parents=True, exist_ok=True)
    wb.save(TPL / "盖章台账模板.xlsx")
    print(f"模板：{TPL / '盖章台账模板.xlsx'}")


def draw_star(d, cx, cy, r, fill):
    pts = []
    for i in range(10):
        ang = math.pi / 2 + i * math.pi / 5
        rr = r if i % 2 == 0 else r * 0.42
        pts.append((cx + rr * math.cos(ang), cy - rr * math.sin(ang)))
    d.polygon(pts, fill=fill)


def build_stamp(path, kind):
    """画一枚纯图形的红章（透明底，不写一个字——写字就要挑字体，跨平台必漏）。"""
    S = 600
    img = Image.new("RGBA", (S, S), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    red = (198, 32, 32, 235)
    if kind == "double":
        d.ellipse([16, 16, S - 16, S - 16], outline=red, width=16)
        d.ellipse([62, 62, S - 62, S - 62], outline=red, width=6)
        for i in range(28):                                  # 一圈小方块，模拟印章外圈的字
            ang = i * 2 * math.pi / 28
            cx, cy = 300 + 248 * math.cos(ang), 300 + 248 * math.sin(ang)
            d.rectangle([cx - 9, cy - 9, cx + 9, cy + 9], fill=red)
        draw_star(d, 300, 300, 118, red)
    elif kind == "ring":
        d.ellipse([20, 20, S - 20, S - 20], outline=red, width=14)
        for i in range(3):
            r = 190 - i * 52
            d.ellipse([300 - r, 300 - r, 300 + r, 300 + r], outline=red, width=5)
    else:                                                    # hatch：骑缝章，斜纹
        d.ellipse([20, 20, S - 20, S - 20], outline=red, width=14)
        d.ellipse([64, 64, S - 64, S - 64], outline=red, width=4)
        for i in range(-7, 8):
            x = 300 + i * 34
            d.line([x - 150, 480, x + 150, 120], fill=red, width=7)
    img.save(path)
    print(f"章图：{path.name}")


def build_pdf(path, no, pages):
    """造一份待盖章报告。中文字体走 reportlab 自带的 CID 字体，不依赖系统字体文件。"""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except Exception:
        pass
    W, H = A4
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle(f"检测报告 {no}")
    c.setAuthor("某某检测技术有限公司")                    # 不留本机用户名
    c.setCreator("build_fixtures.py")
    for p in range(1, pages + 1):
        c.setFont("STSong-Light", 9)
        c.drawString(56, H - 46, f"报告编号：{no}　　第 {p} 页 / 共 {pages} 页")
        c.setStrokeColorRGB(0.6, 0.6, 0.6)
        c.line(56, H - 54, W - 56, H - 54)
        if p == 1:
            c.setFont("STSong-Light", 22)
            c.drawCentredString(W / 2, H - 140, "检 测 报 告")
            c.setFont("STSong-Light", 11)
            body = [
                f"报告编号：{no}",
                f"受检单位：见委托单",
                "检测类别：委托检测",
                "签发日期：2026-09-26",
                "",
                "本报告仅对本次委托送检样品负责。未经本机构书面同意，不得部分复制本报告。",
                "报告正文共 %d 页，附件另附。" % pages,
            ]
            y = H - 190
            for line in body:
                c.drawString(90, y, line)
                y -= 26
        else:
            c.setFont("STSong-Light", 11)
            y = H - 110
            for i in range(1, 13):
                c.drawString(70, y, f"{p}-{i}　检测项目数据记录（仪器导出后逐项核对，数据以原始记录为准）。")
                y -= 34
        c.setFont("STSong-Light", 8)
        c.drawString(56, 40, "某某检测技术有限公司　检测专用　计量认证合格单位")
        c.showPage()
    c.save()
    print(f"报告：{path.name}（{pages} 页）")


def main():
    RAW.mkdir(parents=True, exist_ok=True)
    STAMP_DIR.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)

    build_template()

    with open(RAW / "盖章申请单.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["申请日期", "报告编号", "客户名称", "报告类型", "审签状态", "报告文件", "申请人", "备注"])
        w.writerows(REQUESTS)

    with open(RAW / "用章规则表.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["报告类型", "章名", "落位", "骑缝起始页"])
        w.writerows(RULES)

    for name, kind in STAMP_KINDS.items():
        build_stamp(STAMP_DIR / f"{name}.png", kind)

    for no, pages in PAGES.items():
        build_pdf(PDF_DIR / f"报告_{no}.pdf", no, pages)

    print("\n造数据完成。")


if __name__ == "__main__":
    main()

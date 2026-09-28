#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""第三方检测 · 开票与关账：把当月已出报告、已开发票摆到一起，算出这一批能开哪几张。

一句话：开票的动作不难，难的是**谁可以开、按什么口径开、是不是已经开过、合同额度还剩多少**——
到了关账那天，「报告已经出了但票还没开」的那笔金额对不上，才是财务被追着问的那件事。

用法：
    python run.py              # 跑完停住等回车（Windows 双击也行）
    python run.py --no-pause   # 跑完直接退出
"""

import csv
import re
import socket
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

# ----------------------------------------------------------------- ① 配置区（跑之前主要改这里）

ROOT = Path(__file__).resolve().parent
RAW_DIR = ROOT / "01_raw_data"
OUT_DIR = ROOT / "02_output"
TPL_DIR = ROOT / "source" / "templates"
LOG_FILE = ROOT / "source" / "run_log.txt"

REQ_CSV = RAW_DIR / "开票申请单.csv"          # 客服/业务交上来的开票申请
REP_CSV = RAW_DIR / "报告完成台账.csv"        # 已出报告（出了报告才有开票依据）
INV_CSV = RAW_DIR / "已开发票台账.csv"        # 财务已经开出去的票（含作废）
RULE_CSV = RAW_DIR / "客户开票规则表.csv"     # 客户 → 开票方式 / 税率 / 结算周期 / 是否必须 PO

TPL_XLSX = TPL_DIR / "开票对账台账模板.xlsx"
LEDGER_XLSX = OUT_DIR / "开票对账台账.xlsx"
IMPORT_CSV = OUT_DIR / "开票数据_导入用.csv"  # 给开票软件导入用

# 判空列：每张表都要指定「哪一列一定有值」，表尾那行说明靠它跳过。
# 换一张表先问一句"靠哪一列判空"——判空列挑错，整张表会被静默读空。
REQ_HEAD = "申请日期"
REP_HEAD = "委托单号"
INV_HEAD = "发票号码"
RULE_HEAD = "客户名称"

# 货物名称固定写成一项：检测服务费的税收分类编码各家有自己的口径，交给开票软件去带。
GOODS_NAME = "检测服务费"

BLANK_RE = re.compile(r"[\s\u3000]+")   # 全角空格 strip() 不管它，要显式替换


# ----------------------------------------------------------------- 装依赖 / 日志


def ensure_deps():
    """缺 openpyxl 就自己装（走清华源）。"""
    need = []
    for mod, pkg in (("openpyxl", "openpyxl"),):
        try:
            __import__(mod)
        except ImportError:
            need.append(pkg)
    if not need:
        return
    pkgs = ["-i", "https://pypi.tuna.tsinghua.edu.cn/simple", *need]
    for extra in ([], ["--user"]):
        r = subprocess.run([sys.executable, "-m", "pip", "install", *pkgs, *extra],
                           capture_output=True, text=True)
        if r.returncode == 0:
            return
    print("依赖安装失败，请手动执行：\n"
          "pip install openpyxl -i https://pypi.tuna.tsinghua.edu.cn/simple")
    sys.exit(1)


def host_line():
    """运行主机：主机名 + 本机对外网卡地址，用来分辨这段日志是哪台机器跑出来的。"""
    try:
        name = socket.gethostname()
    except Exception:
        name = "未知"
    ip = "未知"
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
    except Exception:
        try:
            ip = socket.gethostbyname(name)
        except Exception:
            ip = "未知"
    return f"运行主机：{name}（{ip}）"


def log(text):
    with open(LOG_FILE, "a", encoding="utf-8") as f:
        f.write(text + "\n")


# ----------------------------------------------------------------- ② 读数据


def norm(value):
    """比较用的键：去空白（含全角空格）。显示照原样，只压比较用的那一份。"""
    return BLANK_RE.sub("", str(value if value is not None else ""))


def to_money(text):
    """把表里的金额文本转成两位小数。带千分位逗号和 ￥ 的照样读——系统导出来就长这样。"""
    s = str(text).strip().replace(",", "").replace("￥", "").replace("¥", "")
    if not s:
        return 0.0
    return round(float(s), 2)


def tax_rate(text):
    """税率写成「6%」也能读——导出表里常带百分号，直接 float 会炸。"""
    return round(float(str(text).strip().replace("%", "")) / 100, 4)


def read_csv_rows(path, head_col):
    """读一张 CSV 成 dict 列表，全部按文本取。

    **一律不转类型**——委托单号 WT2026-0000715、发票号 00012301 一转 int 前导零就没了，
    跟开票软件里的单号对不上。表尾的说明行靠 `head_col` 那一列为空跳过，所以说明文字别写进那一列。
    """
    rows = []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        for raw in csv.DictReader(f):
            if not raw or not norm(raw.get(head_col, "")):
                continue
            rows.append({k: ("" if v is None else str(v)) for k, v in raw.items()})
    return rows


def read_requests():
    """开票申请单：行序就是客服交单的顺序，出来的清单跟着这个顺序走。"""
    return read_csv_rows(REQ_CSV, REQ_HEAD)


def read_reports():
    """报告完成台账：委托单号 → 这一单的报告信息。**只有出过报告的单子才在这一张表里**。"""
    reports = {}
    for row in read_csv_rows(REP_CSV, REP_HEAD):
        reports[norm(row["委托单号"])] = row
    return reports


def read_invoices():
    """已开发票台账：全部历史票据，作废的也在里面（作废票不占额度，但要看得见）。"""
    return read_csv_rows(INV_CSV, INV_HEAD)


def read_rules():
    """客户开票规则表：客户 → 开票方式 / 税率 / 结算周期 / 是否必须 PO。

    客户名当键时压成"去全角空格"的比较键，显示名另外存一份干净的——
    同一家客户在申请单里多打一个空格，本来会被当成两家。
    """
    rules = {}
    for row in read_csv_rows(RULE_CSV, RULE_HEAD):
        rules[norm(row["客户名称"])] = {
            "客户名称": BLANK_RE.sub("", row["客户名称"].strip()),
            "开票方式": row["开票方式"].strip(),
            "税率": row["税率"].strip(),
            "结算周期": row["结算周期"].strip(),
            "是否必须PO": row["是否必须PO"].strip(),
        }
    return rules


def contract_used(invoices, reports):
    """各合同已经占掉的额度（不含税）：只算状态正常的票，作废票不占。"""
    used = {}
    for inv in invoices:
        if inv["状态"].strip() != "正常":
            continue
        rep = reports.get(norm(inv["关联委托单"]))
        if rep:
            key = norm(rep["合同号"])
            used[key] = used.get(key, 0.0) + to_money(rep["不含税金额"])
    return used


# ----------------------------------------------------------------- ③ 判够不够格


def plan_one(req, reports, rules, opened, used_net, row_no):
    """逐单判定能不能开票，返回 (可开票记录 或 None, 卡在哪一步, 说明)。

    判定顺序 = 业务顺序，不许调：先看客户规则，再看资料全不全，再看报告出了没，
    再看有没有 PO，然后才是金额、重票、合同额度。顺序反了理由就会张冠李戴——
    拿着"额度不够"去找客服，客服补完额度发现这单报告根本没出。
    """
    if not norm(req["委托单号"]):
        return None, "委托单号", "委托单号空着，和开票软件里的单子对不上"

    cust = norm(req["客户名称"])
    if cust not in rules:
        return None, "客户规则", "这家客户没进开票规则表，先定开票方式与税率再开"
    rule = rules[cust]
    no = norm(req["委托单号"])

    if not req["发票抬头"].strip() or not req["纳税人识别号"].strip():
        return None, "开票信息", "发票抬头或纳税人识别号空着，开票软件这一步就过不去"

    rep = reports.get(no)
    if rep is None:                                  # 出了报告才有开票依据
        return None, "报告未出", "报告完成台账里没有这一单，报告没出不能先开票"

    if rule["是否必须PO"] == "是" and not req["PO号"].strip():
        return None, "缺PO", f"这家客户必须凭 PO 开票，申请单上 PO 号空着"

    tax = tax_rate(rule["税率"])
    net = to_money(rep["不含税金额"])
    due = round(net * (1 + tax), 2)                  # 申请金额应当等于 不含税 ×（1+税率）
    want = to_money(req["申请开票金额"])
    if abs(want - due) > 0.005:
        return None, "金额不符", f"申请金额 {want:,.2f}，台账算出来是 {due:,.2f}，差 {abs(want - due):,.2f}"

    if no in opened:                                 # 作废票不算开过，重开要走到这里
        return None, "重复开票", f"这一单已经开过票（{opened[no]}），再开一次就是重票"

    key = norm(rep["合同号"])
    cap = to_money(rep["合同额"])
    used = used_net.get(key, 0.0)
    if used + net > cap + 0.005:
        return None, "超合同额度", (f"合同 {key} 额度 {cap:,.2f}（不含税），已占 {used:,.2f}，"
                                    f"这一单 {net:,.2f} 会超")

    return {
        "委托单号": req["委托单号"].strip(),
        "客户名称": rule["客户名称"],
        "合同号": rep["合同号"].strip(),
        "PO号": req["PO号"].strip(),
        "不含税金额": net,
        "税率": rule["税率"].strip(),
        "税额": round(net * tax, 2),
        "价税合计": due,
        "发票抬头": req["发票抬头"].strip(),
        "纳税人识别号": req["纳税人识别号"].strip(),
        "开票方式": rule["开票方式"].strip(),
        "结算周期": rule["结算周期"].strip(),
    }, "", ""


# ----------------------------------------------------------------- ④ 输出


def build_recon(reports, invoices):
    """关账对账：按客户把「已出报告的不含税金额」和「已开正常票的不含税金额」摆一起。

    差额那一列就是关账那天被问得最多的数——报告出了多少、票开了多少、还差多少没开。
    作废票不计入已开，否则差额会被算小，看着像开完了。
    """
    opened = {}
    for inv in invoices:
        if inv["状态"].strip() != "正常":
            continue
        rep = reports.get(norm(inv["关联委托单"]))
        if rep:
            opened[norm(inv["关联委托单"])] = to_money(rep["不含税金额"])

    names, total, invo, cnt = {}, {}, {}, {}
    for rep in reports.values():
        key = norm(rep["客户名称"])
        names[key] = BLANK_RE.sub("", rep["客户名称"].strip())
        net = to_money(rep["不含税金额"])
        total[key] = total.get(key, 0.0) + net
        if norm(rep["委托单号"]) not in opened:
            cnt[key] = cnt.get(key, 0) + 1
    for no_ in opened:
        rep = reports[no_]
        key = norm(rep["客户名称"])
        invo[key] = invo.get(key, 0.0) + opened[no_]

    rows = []
    for key in sorted(names, key=lambda k: names[k]):
        rows.append([names[key], round(total[key], 2), round(invo.get(key, 0.0), 2),
                     round(total[key] - invo.get(key, 0.0), 2), cnt.get(key, 0)])
    return rows


def write_ledger(items, holds, recon):
    """把三张表写进台账：可开票 / 挂起 / 关账对账。"""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = load_workbook(TPL_XLSX)
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws = wb["可开票清单"]
    for item in items:
        ws.append([item["委托单号"], item["客户名称"], item["合同号"], item["PO号"],
                   item["不含税金额"], item["税率"], item["税额"], item["价税合计"],
                   item["发票抬头"], item["纳税人识别号"], item["开票方式"], item["结算周期"]])
        for c in ws[ws.max_row]:
            c.font = Font(name="微软雅黑", size=10)
            c.border = border
            if c.column == 1:
                c.number_format = "@"          # 委托单号带前导零，当文本存
                c.alignment = Alignment(horizontal="left")
            elif c.column in (5, 7, 8):
                c.number_format = "#,##0.00"

    hs = wb["挂起清单"]
    for hold in holds:
        hs.append(hold)
        for c in hs[hs.max_row]:
            c.font = Font(name="微软雅黑", size=10)
            c.border = border
            if c.column == 1:
                c.number_format = "@"

    rc = wb["关账对账"]
    for row in recon:
        rc.append(row)
        for c in rc[rc.max_row]:
            c.font = Font(name="微软雅黑", size=10)
            c.border = border
            if c.column in (2, 3, 4):
                c.number_format = "#,##0.00"

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(LEDGER_XLSX)


def write_import_csv(items):
    """给开票软件的导入文件：一行一张票，抬头、税号、金额、备注都齐了。"""
    head = ["发票抬头", "纳税人识别号", "货物或应税劳务名称", "金额", "税率", "税额", "价税合计", "备注"]
    lines = [",".join(head)]
    for item in items:
        note = f"合同号 {item['合同号']}；委托单 {item['委托单号']}"
        if item["PO号"]:
            note += f"；PO {item['PO号']}"
        cells = [item["发票抬头"], item["纳税人识别号"], GOODS_NAME, f"{item['不含税金额']:.2f}",
                 item["税率"], f"{item['税额']:.2f}", f"{item['价税合计']:.2f}", note]
        lines.append(",".join('"%s"' % str(x).replace('"', '""') for x in cells))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    IMPORT_CSV.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8-sig")


# ----------------------------------------------------------------- 回读校验


def verify(items, holds, recon):
    from openpyxl import load_workbook

    checks = []
    def ok(name, cond):
        checks.append((name, bool(cond)))

    wb = load_workbook(LEDGER_XLSX, data_only=True)
    ws = wb["可开票清单"]
    read = [dict(zip([c.value for c in ws[1]], [c.value for c in row]))
            for row in ws.iter_rows(min_row=2) if row[0].value]
    hs = wb["挂起清单"]
    hold_read = [(str(r[0].value), str(r[2].value), str(r[3].value))
                 for r in hs.iter_rows(min_row=2) if r[0].value]
    rc = wb["关账对账"]
    recon_read = [dict(zip([c.value for c in rc[1]], [c.value for c in row]))
                  for row in rc.iter_rows(min_row=2) if row[0].value]

    def hold_of(no):
        """这一单卡在哪一步（取不到就返回空串，让断言直接失败）。"""
        for n, step, _ in hold_read:
            if n == no:
                return step
        return ""

    def recon_of(cust):
        for row in recon_read:
            if row["客户名称"] == cust:
                return row
        return {}

    # 数量
    ok("可开票清单 7 张", len(read) == 7)
    ok("挂起清单 9 行", len(hold_read) == 9)
    ok("申请单读进来 16 行（表尾说明行没被当数据）", len(items) + len(holds) == 16)
    ok("关账对账 6 家客户", len(recon_read) == 6)

    # 判定顺序：资料先于额度、报告先于 PO
    ok("0727 卡在开票信息（不是超额度）——资料检查在合同额度之前", hold_of("WT2026-0000727") == "开票信息")
    ok("0712 卡在报告未出", hold_of("WT2026-0000712") == "报告未出")
    ok("0731 卡在缺PO", hold_of("WT2026-0000731") == "缺PO")
    ok("0730 卡在开票信息（税号空）", hold_of("WT2026-0000730") == "开票信息")
    ok("0728 卡在客户规则（新客户没进规则表）", hold_of("WT2026-0000728") == "客户规则")
    ok("0716 卡在重复开票", hold_of("WT2026-0000716") == "重复开票")
    ok("0725 卡在超合同额度", hold_of("WT2026-0000725") == "超合同额度")
    ok("0726 卡在超合同额度（C1 累计 55000+12000 > 60000）", hold_of("WT2026-0000726") == "超合同额度")
    ok("0729 卡在金额不符，理由里带差额 230.00",
       hold_of("WT2026-0000729") == "金额不符" and any("230.00" in d for n, _, d in hold_read if n == "WT2026-0000729"))
    ok("超额度理由里同时写了额度和已占金额",
       any(n == "WT2026-0000726" and "60,000.00" in d and "55,000.00" in d for n, _, d in hold_read))

    # 作废票不挡重开
    ok("0721（原票已作废）进了可开票清单", any(x["委托单号"] == "WT2026-0000721" for x in read))
    ok("已作废的那张票没被算成开过（理由里不出现它）",
       not any(n == "WT2026-0000721" for n, _, _ in hold_read))

    # 金额与税率
    r15 = next(x for x in read if x["委托单号"] == "WT2026-0000715")
    ok("0715 税额 900.00、价税合计 15,900.00（6%）",
       abs(r15["税额"] - 900.0) < 1e-6 and abs(r15["价税合计"] - 15900.0) < 1e-6)
    r33 = next(x for x in read if x["委托单号"] == "WT2026-0000733")
    ok("0733 按 13% 算：税额 390.00、价税合计 3,390.00",
       abs(r33["税额"] - 390.0) < 1e-6 and abs(r33["价税合计"] - 3390.0) < 1e-6)
    r23 = next(x for x in read if x["委托单号"] == "WT2026-0000723")
    ok("0723 价税合计 4,028.00（千分位逗号与空 PO 不影响）",
       abs(r23["价税合计"] - 4028.0) < 1e-6 and not str(r23["PO号"] or "").strip())

    # 客户名归一化：规则表里 C2 带全角空格，申请单里不带，仍然匹配上
    ok("带全角空格的客户名归一化后照样匹配到规则表（0717 跨月汇总客户）",
       any(x["委托单号"] == "WT2026-0000717" and x["开票方式"] == "按月汇总" for x in read))
    ok("台账里客户名是干净写法，不带全角空格",
       all("\u3000" not in str(x["客户名称"]) for x in read))

    # 合同额度按不含税累计
    ok("C1 前面 4 单合计 55,000.00 不含税，第 5 单 12,000 才判超（额度 60,000）",
       recon_of("上海鼎衡环境检测有限公司")["已出报告不含税"] == 67000.0)

    # 关账对账
    ok("C1 差额 47,000.00（报告 67,000 − 已开 20,000）",
       recon_of("上海鼎衡环境检测有限公司")["差额"] == 47000.0)
    ok("C4 已开票按 0 算（那张票作废了），差额 7,000.00",
       recon_of("深圳博测新材科技有限公司")["已开发票不含税"] == 0.0
       and recon_of("深圳博测新材科技有限公司")["差额"] == 7000.0)
    ok("C3 已出报告 33,500.00（含那单被额度挡住的 10,000）",
       recon_of("浙江嘉澳新材料有限公司")["已出报告不含税"] == 33500.0)
    ok("C3 差额 18,000.00（超额度那单也算没开）",
       recon_of("浙江嘉澳新材料有限公司")["差额"] == 18000.0)
    ok("未开票委托单数合计 = 14 单",
       sum(r["未开票委托单数"] for r in recon_read) == 14)
    ok("全部客户差额合计 175,300.00",
       abs(sum(r["差额"] for r in recon_read) - 175300.0) < 1e-6)

    # 清单顺序跟申请单走，不是按单号字符串排
    order = [x["委托单号"] for x in read]
    ok("清单按申请单行序排，不是按委托单号字符串排", order != sorted(order))
    ok("第一张是申请单第一行 WT2026-0000715", order[0] == "WT2026-0000715")
    ok("委托单号前导零没丢（台账里还是 WT2026-0000715）", order[0] == "WT2026-0000715")

    # 导入文件
    import_rows = [r for r in IMPORT_CSV.read_text(encoding="utf-8-sig").splitlines() if r.strip()]
    ok("开票导入文件 = 表头 + 7 行", len(import_rows) == 8)
    ok("导入文件里带合同号与委托单号备注",
       any("合同号 HT2026-A01" in r and "WT2026-0000715" in r for r in import_rows))
    total = sum(x["价税合计"] for x in read)
    ok("待开票价税合计 96,738.00", abs(total - 96738.0) < 1e-6)

    print()
    for name, good in checks:
        print(f"  {'OK  ' if good else 'FAIL'} {name}")
    passed = sum(1 for _, g in checks if g)
    print(f"\n回读校验：{passed}/{len(checks)} 项通过")
    return passed == len(checks), len(checks), passed


# ----------------------------------------------------------------- 主线


def main():
    ensure_deps()
    t0 = time.time()
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    log(host_line())
    log(f"开始执行：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    log(f"Python：{sys.version.split()[0]}（{sys.executable}）")

    requests = read_requests()
    reports = read_reports()
    invoices = read_invoices()
    rules = read_rules()

    # 合同已占额度先按历史正常票算一遍，跑的过程中把这一批批准的累加进去。
    used_net = contract_used(invoices, reports)
    opened = {}
    for inv in invoices:
        if inv["状态"].strip() == "正常":
            opened[norm(inv["关联委托单"])] = inv["发票号码"].strip()

    items, holds = [], []
    for row_no, req in enumerate(requests, 1):
        item, step, reason = plan_one(req, reports, rules, opened, used_net, row_no)
        if item is None:
            holds.append([req["委托单号"].strip(), req["客户名称"].strip(), step, reason])
            continue
        items.append(item)
        used_net[norm(item["合同号"])] = used_net.get(norm(item["合同号"]), 0.0) + item["不含税金额"]
        log(f"可开票：{item['委托单号']}  {item['客户名称']}  价税合计 {item['价税合计']:,.2f}")

    recon = build_recon(reports, invoices)
    write_ledger(items, holds, recon)
    write_import_csv(items)
    ok, total_chk, passed = verify(items, holds, recon)

    total_amount = sum(x["价税合计"] for x in items)
    gap = sum(r[3] for r in recon)
    used_s = time.time() - t0
    print(f"\n申请单 {len(requests)} 行 → 可开票 {len(items)} 张 / 价税合计 {total_amount:,.2f} 元 "
          f"→ 挂起 {len(holds)} 行")
    print(f"关账对账：{len(recon)} 家客户，报告已出未开票合计 {gap:,.2f} 元（不含税）")
    print(f"结果目录：{OUT_DIR}")
    log(f"申请单 {len(requests)} 行 → 可开票 {len(items)} 张 / 价税合计 {total_amount:,.2f} 元 → 挂起 {len(holds)} 行")
    log(f"关账对账：{len(recon)} 家客户，报告已出未开票合计 {gap:,.2f} 元")
    log(f"执行完成：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  用时 {used_s:.1f}s  "
        f"判定 {'PASS' if ok else 'FAIL'}（{passed}/{total_chk}）")
    print(f"判定：{'PASS' if ok else 'FAIL'}（{passed}/{total_chk} 项）")

    if "--no-pause" not in sys.argv:
        input("\n按回车退出…")


if __name__ == "__main__":
    main()

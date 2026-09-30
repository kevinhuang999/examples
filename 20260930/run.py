# -*- coding: utf-8 -*-
"""第三方检测「检测周期进度跟踪台账」自动生成。

读 01_raw_data/ 里的四张表（检测委托台账 / 检测项目周期表 / 环节节点记录 / 工作日历），
按工作日口径从收样日倒推出每张委托单的承诺交期，还原当前卡在哪个环节，
再标出还剩几个工作日、是否超期，最后出一份进度台账和一份催办清单。
跑完自动回读校验，打印 PASS/FAIL。

用法：python run.py      （Windows 双击 run.py 也行；加 --no-pause 跳过结尾回车）
"""

import csv
import socket
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path

BASE = Path(__file__).resolve().parent
RAW_DIR = BASE / "01_raw_data"
OUT_DIR = BASE / "02_output"
SRC_DIR = BASE / "source"
TPL_FILE = SRC_DIR / "templates" / "检测进度台账模板.xlsx"
LOG_FILE = SRC_DIR / "run_log.txt"

ORDER_FILE = RAW_DIR / "检测委托台账.csv"
STANDARD_FILE = RAW_DIR / "检测项目周期表.csv"
STEP_FILE = RAW_DIR / "环节节点记录.csv"
CALENDAR_FILE = RAW_DIR / "工作日历.csv"
LEDGER_FILE = OUT_DIR / "检测进度台账.xlsx"
REMIND_FILE = OUT_DIR / "催办清单.csv"

CUTOFF = date(2026, 9, 30)      # 数据截止日：进度与剩余工作日都算到这一天

# 第三方检测的真实流转环节，顺序即业务顺序，不要按时间戳排序
STEPS = ["收样登记", "样品前处理", "上机检测", "数据审核", "报告编制", "报告审核", "报告发出"]
STEP_OWNER = {
    "样品前处理": "前处理组",
    "上机检测": "检测组",
    "数据审核": "数据审核（技术负责人）",
    "报告编制": "报告编制岗",
    "报告审核": "授权签字人",
    "报告发出": "客服",
}

STATUS_HANG = "挂起"
STATUS_DONE = "已出报告"
STATUS_OVER = "超期"
STATUS_SOON = "临近超期"
STATUS_OK = "正常"

# 台账列名（表头按栏目名反查，写死列号必漂）
COL_NO = "委托单号"
COL_CLIENT = "客户名称"
COL_SAMPLE = "样品编号"
COL_ITEM = "检测项目"
COL_SUB = "是否分包"
COL_RECV = "收样日期"


def ensure_openpyxl():
    """缺依赖自动装（走清华源），装不上再退到 --user。"""
    try:
        import openpyxl  # noqa: F401
        return
    except ImportError:
        print("缺少 openpyxl，正在用清华源安装 …")
    mirror = ["-i", "https://pypi.tuna.tsinghua.edu.cn/simple"]
    for extra in ([], ["--user"]):
        args = [sys.executable, "-m", "pip", "install", "openpyxl"] + extra + mirror
        if subprocess.call(args) == 0:
            return
    raise SystemExit("openpyxl 装不上，请手动执行：pip install openpyxl")


def machine_tag():
    """运行主机标识：主机名 + 本机对外网卡 IP，用来分辨是哪台机器跑的。"""
    host = socket.gethostname()
    ip = "未知"
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sock.connect(("8.8.8.8", 80))
        ip = sock.getsockname()[0]
        sock.close()
    except OSError:
        try:
            ip = socket.gethostbyname(host)
        except OSError:
            ip = "未知"
    return f"{host}（{ip}）"


def log_open():
    """开一段追加日志：先记运行主机，再记开始时间与解释器。"""
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    fh = open(LOG_FILE, "a", encoding="utf-8")
    fh.write(f"\n运行主机：{machine_tag()}\n")
    fh.write(f"开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}  "
             f"Python {sys.version.split()[0]}（{Path(sys.executable).name}）\n")
    return fh


def norm(text):
    """比较用的归一化：去掉所有空白（含全角空格），统一成字符串。"""
    return "".join(str(text).split())


def parse_date(text):
    """认日期；认不出返回 None，绝不拿当天日期凑。"""
    text = str(text).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_dt(text):
    """认环节完成时间；只给到日也认，认不出返回 None。"""
    text = str(text).strip()
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%Y/%m/%d %H:%M"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def read_table(path, key_col):
    """读 CSV。key_col 是「靠哪一列判空」——每张表传它自己的键列，
    写死某一列会把没有该列的表整张读空，而且一行错都不报。"""
    if not path.exists():
        raise SystemExit(f"找不到数据文件：{path.name}，先跑 python source/build_fixtures.py")
    rows = []
    with open(path, encoding="utf-8-sig", newline="") as f:
        for raw in csv.DictReader(f):
            key = (raw.get(key_col) or "").strip()
            if not key:
                continue          # 表尾合计行 / 说明行：键列为空就跳过
            rows.append({k: (v or "").strip() for k, v in raw.items()})
    if not rows:
        raise SystemExit(f"按「{key_col}」判空后一行都没读到：{path.name} 的列名可能对不上")
    return rows


def load_workdays():
    """读工作日历 → 工作日集合。周六不一定不上班（调休）、工作日不一定不是假日。"""
    days = set()
    for row in read_table(CALENDAR_FILE, "日期"):
        if row.get("是否工作日") == "是":
            days.add(row["日期"])
    return days


def add_workdays(start, count, workdays):
    """从 start 的次日往后数 count 个工作日，得到承诺交期。"""
    from datetime import timedelta
    cur, left = start, count
    while left > 0:
        cur += timedelta(days=1)
        if cur.isoformat() in workdays:
            left -= 1
    return cur


def count_workdays(after, upto, workdays):
    """数 (after, upto] 区间里有几个工作日 = 这张单已经耗掉几天。"""
    from datetime import timedelta
    n, cur = 0, after
    while cur < upto:
        cur += timedelta(days=1)
        if cur.isoformat() in workdays:
            n += 1
    return n


def load_standards():
    """检测项目 → 周期。分包件走「分包周期」那一列，没有就退回本所周期。"""
    std = {}
    for row in read_table(STANDARD_FILE, "检测项目"):
        std[norm(row["检测项目"])] = {
            "cycle": int(row["检测周期(工作日)"]),
            "sub": int(row["分包周期(工作日)"]) if row.get("分包周期(工作日)") else None,
        }
    return std


def load_orders():
    """读委托台账，按委托单归堆：一张单一行，含它的样品与检测项目。"""
    orders = {}
    for row in read_table(ORDER_FILE, COL_NO):
        no = row[COL_NO]
        o = orders.setdefault(no, {
            "客户": norm(row[COL_CLIENT]),
            "样品": [], "项目": [], "分包": row[COL_SUB] == "是",
            "收样日": parse_date(row[COL_RECV]),
        })
        if norm(row[COL_CLIENT]) != o["客户"]:
            o["客户"] = norm(row[COL_CLIENT])       # 同一单客户名不一致，按最后一条显示
        if row[COL_SAMPLE] not in o["样品"]:
            o["样品"].append(row[COL_SAMPLE])
        if row[COL_ITEM] not in o["项目"]:
            o["项目"].append(row[COL_ITEM])
    return orders


def load_steps():
    """读环节节点记录，按委托单归堆；认不出的时间保留为 None，后面单独挂起。"""
    steps = {}
    for row in read_table(STEP_FILE, COL_NO):
        steps.setdefault(row[COL_NO], []).append({
            "环节": norm(row["环节"]),
            "时间": parse_dt(row["完成时间"]),
        })
    return steps


def latest_steps(items):
    """还原这单的环节进度。返回 (已完成环节列表, 卡住原因)。
    只按环节顺序还原，不按时间排序——时间转抄错了，排序会把倒挂洗成正常。"""
    seen = {}
    for it in items:
        if it["环节"] not in STEPS:
            return [], f"环节「{it['环节']}」不在流程里"
        if it["环节"] in seen:
            return [], f"环节「{it['环节']}」有两条记录，无法判断哪条为准"
        if it["时间"] is None:
            return [], f"环节「{it['环节']}」的完成时间认不出"
        seen[it["环节"]] = it["时间"]
    ordered = [s for s in STEPS if s in seen]
    for prev, cur in zip(ordered, ordered[1:]):
        if seen[cur] < seen[prev]:
            return [], f"环节时间倒挂：{cur} 早于 {prev}"
    return ordered, None


def due_of(order, standards, workdays):
    """从收样日按工作日倒推销出承诺交期。返回 (周期, 交期, 卡住原因)。"""
    if order["收样日"] is None:
        return None, None, "缺收样日期"
    cycles = []
    for item in order["项目"]:
        std = standards.get(norm(item))
        if std is None:
            return None, None, f"检测项目「{item}」未登记检测周期"
        cycles.append(std["sub"] if (order["分包"] and std["sub"]) else std["cycle"])
    cycle = max(cycles)      # 一张单多项检测：承诺周期取最长的那项
    return cycle, add_workdays(order["收样日"], cycle, workdays), None


def build_states(orders, standards, steps, workdays):
    """逐单判定：先查能不能算出交期，再看环节，最后才排超期与否。
    顺序就是业务顺序——反了的话，挂起理由会指向不对的补料动作。"""
    states = []
    for no, order in orders.items():
        rec = {
            "委托单号": no, "客户名称": order["客户"], "样品编号": "、".join(order["样品"]),
            "检测项目": "、".join(order["项目"]), "项目数": len(order["项目"]),
            "收样日期": order["收样日"], "是否分包": order["分包"],
        }
        cycle, due, why = due_of(order, standards, workdays)
        done, step_why = latest_steps(steps.get(no, []))
        if why is None and not steps.get(no):
            step_why = "一条环节记录都没有"
        if why is None and step_why is not None:
            why, cycle, due = step_why, None, None
        if why is not None:
            rec.update({"承诺周期": "", "承诺交期": "", "当前环节": "", "已完成环节": "",
                        "已耗工作日": "", "剩余工作日": "", "进度": "",
                        "状态": STATUS_HANG, "下一环节负责岗位": "", "备注": why})
            states.append(rec)
            continue

        used = count_workdays(order["收样日"], CUTOFF, workdays)
        current = STEPS[len(done)] if len(done) < len(STEPS) else ""
        rec.update({
            "承诺周期": cycle, "承诺交期": due, "已完成环节": len(done),
            "当前环节": current or "（全部完成）", "已耗工作日": used,
            "进度": f"{round(len(done) / len(STEPS) * 100)}%",
            "下一环节负责岗位": STEP_OWNER.get(current, ""),
        })
        if "报告发出" in done:
            rec.update({"剩余工作日": "", "状态": STATUS_DONE, "备注": "报告已发出"})
        else:
            left = cycle - used
            rec["剩余工作日"] = left
            if left < 0:
                rec.update({"状态": STATUS_OVER, "备注": f"已超出承诺交期 {-left} 个工作日"})
            elif left <= 1:
                rec.update({"状态": STATUS_SOON, "备注": "就要到承诺交期了"})
            else:
                rec.update({"状态": STATUS_OK, "备注": ""})
        states.append(rec)
    states.sort(key=lambda r: r["委托单号"])
    return states


def head_col(ws, name):
    """按栏目名反查列号（模板里调过一次顺序也不怕）。"""
    for cell in ws[2]:
        if norm(cell.value) == norm(name):
            return cell.column
    raise SystemExit(f"模板里找不到栏目「{name}」")


def write_ledger(states, out_file):
    """把结果写回模板：明细行不够就补行，状态列刷底色。"""
    from openpyxl import load_workbook
    from openpyxl.styles import Border, PatternFill, Side

    wb = load_workbook(TPL_FILE)
    ws = wb["检测进度台账"]
    cols = ["委托单号", "客户名称", "样品编号", "检测项目", "项目数", "收样日期",
            "承诺周期(工作日)", "承诺交期", "当前环节", "已完成环节", "已耗工作日",
            "剩余工作日", "进度", "状态", "下一环节负责岗位", "备注"]
    fill_of = {STATUS_OVER: "FFC7CE", STATUS_SOON: "FFEB9C", STATUS_OK: "C6EFCE",
               STATUS_DONE: "D9D9D9", STATUS_HANG: "F2F2F2"}
    thin = Side(style="thin", color="9E9E9E")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i, rec in enumerate(states):
        r = 3 + i
        for name in cols:
            key = name.replace("(工作日)", "")
            cell = ws.cell(row=r, column=head_col(ws, name))
            cell.border = border
            if name == "收样日期":
                cell.value = rec["收样日期"]
                cell.number_format = "yyyy-mm-dd"
            elif name == "承诺交期":
                cell.value = rec["承诺交期"]
                cell.number_format = "yyyy-mm-dd"
            elif name == "进度":
                cell.value = rec["进度"]
            elif name == "状态":
                cell.value = rec["状态"]
                cell.fill = PatternFill("solid", fgColor=fill_of.get(rec["状态"], "FFFFFF"))
            else:
                cell.value = rec.get(key if key in rec else name, "")
    out_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_file)
    return cols


def export_remind(states, out_file):
    """催办清单：还没出报告、且超期或就要超期的单，按剩余工作日升序。"""
    todos = [r for r in states if r["状态"] in (STATUS_OVER, STATUS_SOON)]
    todos.sort(key=lambda r: r["剩余工作日"])
    header = ["委托单号", "客户名称", "承诺交期", "当前环节", "下一环节负责岗位",
              "剩余工作日", "状态", "备注"]
    for r in todos:
        r["_行"] = [r["委托单号"], r["客户名称"], r["承诺交期"], r["当前环节"],
                    r["下一环节负责岗位"], r["剩余工作日"], r["状态"], r["备注"]]
    out_file.parent.mkdir(parents=True, exist_ok=True)
    with open(out_file, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in todos:
            w.writerow(r["_行"])
    return todos


def verify(states, todos, ledger_file, workdays):
    """回读生成的台账逐项断言，打印 OK 与总判定。"""
    from openpyxl import load_workbook
    wb = load_workbook(ledger_file)
    ws = wb["检测进度台账"]
    col = {norm(c.value): c.column for c in ws[2] if c.value}
    read = []
    for r in range(3, ws.max_row + 1):
        no = ws.cell(row=r, column=col[norm("委托单号")]).value
        if not no:
            continue
        read.append({k: ws.cell(row=r, column=v).value for k, v in col.items()})

    by_no = {r["委托单号"]: r for r in read}
    checks = []

    def ok(name, cond, extra=""):
        checks.append((name, bool(cond), extra))

    hang = [r for r in states if r["状态"] == STATUS_HANG]
    ok("回读行数 = 委托单数", len(read) == len(states) == 11, f"{len(read)} 行 / {len(states)} 单")
    ok("挂起 4 单", len(hang) == 4, "、".join(r["委托单号"] for r in hang))
    reasons = {r["备注"] for r in hang}
    ok("挂起理由四类齐全", len(reasons) == 4, " / ".join(sorted(reasons)))
    ok("已出报告 1 单", sum(1 for r in states if r["状态"] == STATUS_DONE) == 1)
    ok("正常 3 单", sum(1 for r in states if r["状态"] == STATUS_OK) == 3)
    ok("催办清单 3 条", len(todos) == 3, "、".join(r["委托单号"] for r in todos))
    ok("催办按剩余升序（超期的排最前）", todos and todos[0]["委托单号"] == "WT2026-0000123")

    r121 = by_no["WT2026-0000121"]
    ok("一单多项目周期取最长（5 而非 3）", r121["承诺周期(工作日)"] == 5, str(r121["承诺周期(工作日)"]))
    ok("交期倒排按调休后的工作日历", str(r121["承诺交期"])[:10] == "2026-09-30", str(r121["承诺交期"])[:10])
    ok("剩余 0 归「临近超期」而不是正常", r121["剩余工作日"] == 0 and r121["状态"] == STATUS_SOON)
    ok("当前环节还原成「报告编制」", r121["当前环节"] == "报告编制", str(r121["当前环节"]))
    ok("进度 4/7 ≈ 57%", r121["进度"] == "57%", str(r121["进度"]))

    r123 = by_no["WT2026-0000123"]
    ok("超期 1 个工作日 = 剩余 -1", r123["剩余工作日"] == -1 and r123["状态"] == STATUS_OVER)
    r122 = by_no["WT2026-0000122"]
    ok("跨国庆假期倒排（交期 10-08）", str(r122["承诺交期"])[:10] == "2026-10-08", str(r122["承诺交期"])[:10])
    r125 = by_no["WT2026-0000125"]
    ok("分包件走分包周期 7", r125["承诺周期(工作日)"] == 7, str(r125["承诺周期(工作日)"]))
    ok("分包件交期按分包周期倒排（10-09）", str(r125["承诺交期"])[:10] == "2026-10-09", str(r125["承诺交期"])[:10])
    r124 = by_no["WT2026-0000124"]
    ok("客户名全角空格已归一", r124["客户名称"] == "华泰环境", repr(str(r124["客户名称"])))
    r128 = by_no["WT2026-0000128"]
    ok("已出报告的单剩余列留空", r128["状态"] == STATUS_DONE and r128["剩余工作日"] in (None, ""))
    ok("09-26 调休算工作日", "2026-09-26" in workdays)
    ok("10-01 假期不算工作日", "2026-10-01" not in workdays)
    ok("承诺交期是日期格式", str(r121["承诺交期"])[:4] == "2026")

    passed = sum(1 for _, c, _ in checks if c)
    print("\n---- 回读校验 ----")
    for name, cond, extra in checks:
        print(f"[{'OK ' if cond else 'BAD'}] {name}" + (f"  （{extra}）" if extra else ""))
    print(f"\n共 {len(checks)} 项，通过 {passed} 项")
    return passed == len(checks)


def main():
    ensure_openpyxl()
    fh = log_open()
    started = datetime.now()
    passed = False
    try:
        workdays = load_workdays()
        standards = load_standards()
        orders = load_orders()
        steps = load_steps()
        states = build_states(orders, standards, steps, workdays)
        write_ledger(states, LEDGER_FILE)
        todos = export_remind(states, REMIND_FILE)
        passed = verify(states, todos, LEDGER_FILE, workdays)
        print(f"\n台账：{LEDGER_FILE.relative_to(BASE)}（{len(states)} 张委托单）")
        print(f"催办：{REMIND_FILE.relative_to(BASE)}（{len(todos)} 条）")
    except SystemExit as exc:
        print(f"\n{exc}")
    finally:
        used = (datetime.now() - started).total_seconds()
        verdict = "PASS" if passed else "FAIL"
        print(f"\n执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}  用时 {used:.1f}s  判定 {verdict}")
        fh.write(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}  用时 {used:.1f}s  判定 {verdict}\n")
        fh.close()
    if "--no-pause" not in sys.argv:
        input("\n按回车键退出 …")


if __name__ == "__main__":
    main()

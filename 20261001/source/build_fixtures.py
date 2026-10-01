# -*- coding: utf-8 -*-
"""重造本示例的原始数据与模板。

上半个文件建模板：source/templates/检测结果报表模板.xlsx
（抬头三格 + 表头行 + 4 行预置明细，带边框与公式 + 汇总行 + 说明行）。
下半个文件造数据：01_raw_data/仪器导出结果.csv（仪器软件导出的结果明细）
与 01_raw_data/受检单位清单.csv。

跑法：python source/build_fixtures.py
"""

import csv
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent      # 包根：source/ 的上一层
RAW_DIR = BASE / "01_raw_data"
SRC_DIR = BASE / "source"
TPL_DIR = SRC_DIR / "templates"
TPL_FILE = TPL_DIR / "检测结果报表模板.xlsx"
SHEET_NAME = "检测结果报表"

HEADERS = ["序号", "样品编号", "检测项目", "称样量(g)", "定容体积(mL)", "稀释倍数",
           "仪器读数(mg/L)", "含量(mg/kg)", "限值(mg/kg)", "判定"]
PRESET_ROWS = 4        # 模板预置明细行数：数据够了照用，少了拷行，多了清空
HEAD_ROW = 6

RESULT_HEAD = ["样品编号", "受检单位", "检测项目", "称样量(g)", "定容体积(mL)",
               "稀释倍数", "仪器读数(mg/L)", "限值(mg/kg)"]

# 仪器软件导出的结果明细。故意造的边界：
#   ① 样品编号带前导零（QC-0001234）② 同一家单位一处写成全角空格版（上海丰味　食品有限公司）
#   ③ 仪器读数为空 / 称样量为空 / 稀释倍数为空 / 读数写成「<0.01」
#   ④ 含量正好等于限值（0.5000 g 称样、读数 5.000 → 含量 0.500）
#   ⑤ 表尾一条说明行，第一列留空（判空列就是第一列）
RESULT_ROWS = [
    ("QC-0001234", "上海丰味\u3000食品有限公司", "铅", "0.5120", "50", "1", "0.048", "0.5"),
    ("QC-0001234", "上海丰味食品有限公司", "镉", "0.5120", "50", "1", "0.012", "0.1"),
    ("QC-0001238", "上海丰味食品有限公司", "铅", "0.5000", "50", "1", "", "0.5"),
    ("QC-0001235", "上海丰味食品有限公司", "铅", "0.5098", "50", "2", "6.10", "0.5"),
    ("QC-0001236", "上海丰味食品有限公司", "铅", "0.5010", "100", "1", "0.210", "0.5"),
    ("QC-0001239", "上海丰味食品有限公司", "铅", "", "50", "1", "0.150", "0.5"),
    ("QC-0001237", "上海丰味食品有限公司", "铅", "0.5150", "50", "1", "0.330", "0.5"),
    ("QC-0001243", "上海丰味食品有限公司", "铅", "0.5000", "50", "1", "5.000", "0.5"),
    ("QC-0001241", "苏州清源饮料有限公司", "铅", "0.5080", "50", "1", "0.088", "0.5"),
    ("QC-0001240", "苏州清源饮料有限公司", "铅", "0.5050", "50", "1", "<0.01", "0.5"),
    ("QC-0001241", "苏州清源饮料有限公司", "镉", "0.5080", "50", "1", "0.150", "0.1"),
    ("QC-0001244", "上海丰味食品有限公司", "铅", "0.5080", "50", "1", "0.102", "0.5"),
    ("QC-0001242", "上海丰味食品有限公司", "铅", "0.5100", "50", "", "0.400", "0.5"),
]
RESULT_TAIL = ["", "本表由仪器软件导出，未检出结果以检出限形式给出，未做检出限换算。", "", "", "", "", "", ""]

UNIT_HEAD = ["受检单位", "受检单位简称", "本批报告编号"]
UNIT_ROWS = [
    ("上海丰味食品有限公司", "丰味", "RPT-2026-0912"),
    ("苏州清源饮料有限公司", "清源", "RPT-2026-0913"),
]


def build_template():
    """建模板：预置 4 行明细，公式列写成相对引用形式，插行/拷行时行号才跟得上。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    TPL_DIR.mkdir(parents=True, exist_ok=True)
    thin = Side(style="thin", color="808080")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    ws["A1"] = "检测结果报表"
    ws["A1"].font = Font(size=14, bold=True)
    for r, label in ((2, "报告编号"), (3, "受检单位"), (4, "检测日期")):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2).border = box

    for i, name in enumerate(HEADERS, 1):
        c = ws.cell(row=HEAD_ROW, column=i, value=name)
        c.font = Font(bold=True)
        c.fill = head_fill
        c.border = box
        c.alignment = Alignment(horizontal="center")

    first = HEAD_ROW + 1
    for r in range(first, first + PRESET_ROWS):
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=r, column=col).border = box
        # 公式列只写公式、不写值：含量按「读数 × 定容 × 稀释 ÷（称样量 × 1000）」算
        ws.cell(row=r, column=8, value=f"=ROUND(G{r}*E{r}*F{r}/(D{r}*1000),3)").number_format = "0.000"
        ws.cell(row=r, column=10, value=f'=IF(H{r}<=I{r},"合格","不合格")')
        ws.cell(row=r, column=4).number_format = "0.0000"
        ws.cell(row=r, column=7).number_format = "0.000"
        ws.cell(row=r, column=9).number_format = "0.000"

    sum_row = first + PRESET_ROWS
    ws.cell(row=sum_row, column=2, value="合格率").font = Font(bold=True)
    rate = ws.cell(row=sum_row, column=10, value=(
        f'=IF(COUNTIF(J{first}:J{sum_row - 1},"合格")+COUNTIF(J{first}:J{sum_row - 1},"不合格")=0,"",'
        f'ROUND(COUNTIF(J{first}:J{sum_row - 1},"合格")/'
        f'(COUNTIF(J{first}:J{sum_row - 1},"合格")+COUNTIF(J{first}:J{sum_row - 1},"不合格")),4))'))
    rate.number_format = "0.00%"
    ws.cell(row=sum_row + 1, column=2, value="说明：含量＝仪器读数 × 定容体积 × 稀释倍数 ÷（称样量 × 1000）；判定按「含量 ≤ 限值」为合格。")

    for i, w in enumerate([6, 16, 12, 12, 14, 10, 16, 14, 14, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(TPL_FILE)
    print(f"模板已生成：{TPL_FILE.name}（预置明细 {PRESET_ROWS} 行，汇总行在第 {sum_row} 行）")


def write_csv(path, head, rows, tail=None):
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(head)
        w.writerows(rows)
        if tail:
            w.writerow(tail)
    print(f"数据已生成：{path.name}（{len(rows)} 行数据）")


def main():
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    build_template()
    write_csv(RAW_DIR / "仪器导出结果.csv", RESULT_HEAD, RESULT_ROWS, RESULT_TAIL)
    write_csv(RAW_DIR / "受检单位清单.csv", UNIT_HEAD, UNIT_ROWS)
    return 0


if __name__ == "__main__":
    sys.exit(main())

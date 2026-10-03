#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""色谱峰表批量汇总：把仪器导出的峰表文件解析成一张 Excel。

产物（02_output/色谱峰汇总.xlsx）三页：
  峰明细   —— 一个峰一行（含不参与归一的溶剂峰，留痕）
  样品汇总 —— 一次进样一行，给目标峰的面积归一含量与判定
  挂起清单 —— 解析不了、算不出来的文件，一个一行，写清原因

用法：
    python run.py            跑完暂停，双击也能跑
    python run.py --no-pause 不暂停
首次运行会自动用清华源补 openpyxl。
"""
import csv
import re
import shutil
import socket
import subprocess
import sys
import time
from pathlib import Path

# ===== ① 路径与字段配置区：换仪器、换判据，只改这一段 =====
BASE = Path(__file__).resolve().parent
RAW = BASE / "01_raw_data"
OUT = BASE / "02_output"
RULES = BASE / "source" / "templates" / "峰归属规则.csv"
LOG = BASE / "source" / "run_log.txt"
OUT_NAME = "色谱峰汇总.xlsx"

# 仪器软件导出文件的编码候选：中文 Windows 上常吐 GBK/ANSI
ENCODINGS = ("utf-8-sig", "gbk", "utf-8")

# 峰表列名候选：各家仪器叫法不一，一律按列名反查，不写死列序
COLS = {
    "seq": ("序号", "编号", "no.", "#"),
    "name": ("峰名", "名称", "化合物", "组分"),
    "rt": ("保留时间", "时间", "rt"),
    "area": ("峰面积", "面积", "area"),
}

# 峰面积写成这些的，代表没测到 / 低于检出限
AREA_EMPTY = ("", "-", "—", "n.d.", "nd", "n.a.", "na")
UNDER_NOTE = "低于线性下限，按半值估"

NORM_RANGE = (98.0, 102.0)      # 目标峰面积归一含量的接受范围（示例口径）
PARALLEL_LIMIT = 2.0            # 同一样品两次进样，归一含量的相对偏差上限（%）

DETAIL_SHEET, SUMMARY_SHEET, HANG_SHEET = "峰明细", "样品汇总", "挂起清单"


class HangError(Exception):
    """这条文件处理不了：不进明细、不进汇总，只进挂起清单。"""

    def __init__(self, msg, no=""):
        super().__init__(msg)
        self.no = no


# ===== ② 读数据：一份文件 = 一次进样，先切成头字段与峰行 =====
def read_lines(path):
    """按候选编码依次试读。中文峰名用 utf-8 读会直接抛异常，不是乱码。"""
    for enc in ENCODINGS:
        try:
            return path.read_text(encoding=enc).splitlines(), enc
        except UnicodeDecodeError:
            continue
    raise HangError(f"编码认不出来（试过 {'/'.join(ENCODINGS)}）")


def split_file(lines):
    """把一份仪器导出文件切成三段：头部字段、峰表表头、峰行。

    头部字段＝冒号行；峰表起点＝第一列写着「序号」的那一行。
    「总计」行的第一列不是数字，正好当峰表终点——别用「行数减一」去猜。
    """
    fields, head, peaks, started = {}, None, [], False
    for raw in lines:
        line = raw.rstrip()
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        cells = [c.strip() for c in re.split(r"\t|\s{2,}", line)]
        if not started:
            if cells and cells[0] in COLS["seq"]:
                head, started = cells, True
            elif ":" in line:
                key, val = line.split(":", 1)
                fields[key.strip().lstrip("\ufeff")] = val.strip()
            continue
        if cells and cells[0].isdigit():
            peaks.append(cells)
        else:
            break
    if head is None:
        raise HangError("找不到峰表表头（第一列应为「序号」）")
    return fields, head, peaks


def col_of(head, key):
    """按列名反查列号：列叫「RT」还是「保留时间/min」都能认。"""
    for i, cell in enumerate(head):
        plain = re.sub(r"[\s/（）()]", "", cell).lower()
        for alias in COLS[key]:
            if alias in plain:
                return i
    raise HangError(f"峰表里找不到「{COLS[key][0]}」列")


def load_rules(path):
    """峰归属规则表：峰名关键词 → 是否参与归一 / 是否目标峰。先匹配到的先用。"""
    with path.open(encoding="utf-8-sig", newline="") as fh:
        rows = [r for r in csv.DictReader(fh) if (r.get("峰名关键词") or "").strip()]
    return [{"key": r["峰名关键词"].strip(),
             "norm": r["参与归一"].strip() == "是",
             "target": r["是否目标峰"].strip() == "是"} for r in rows]


def peak_rule(name, rules):
    for r in rules:
        if r["key"] in name:
            return r
    return {"key": "", "norm": True, "target": False}   # 规则表没覆盖的，默认参与归一


def parse_area(text):
    """峰面积取值，返回 (数值, 备注)。

    低于线性下限写成「<1000」的，按 1000 的一半记并标出来——当成 0 会把
    归一含量的分母悄悄改小，跑出来不报错、结果全偏。
    写成 n.d./空 的，说明这个峰本身没测出来，整条挂起，不许当 0。
    """
    s = (text or "").strip()
    if s.lower() in AREA_EMPTY:
        raise HangError(f"峰面积是「{s or '空'}」，没法参与计算")
    if s[:1] in ("<", "＜"):
        try:
            return float(s[1:].strip()) / 2, UNDER_NOTE
        except ValueError:
            raise HangError(f"峰面积「{s}」解析不了")
    try:
        return float(s), ""
    except ValueError:
        raise HangError(f"峰面积「{s}」解析不了")


def parse_one(path, rules):
    """一份文件＝一次进样：读样品信息、剔溶剂峰、按峰面积算归一含量。"""
    lines, enc = read_lines(path)
    fields, head, rows = split_file(lines)
    no = fields.get("样品编号", "").strip()
    sample = fields.get("样品名称", "").strip()
    seq = fields.get("进样次数", "").strip()
    if not no:
        raise HangError("文件头没有「样品编号」")
    if not seq.isdigit():
        raise HangError(f"「进样次数」不是数字：{seq or '空'}", no)

    idx = {k: col_of(head, k) for k in ("name", "rt", "area")}
    peaks = []
    for r in rows:
        name = r[idx["name"]].strip()
        rule = peak_rule(name, rules)
        try:
            area, note = parse_area(r[idx["area"]])
        except HangError as e:
            raise HangError(f"「{name}」{e}", no)
        try:
            rt = float(r[idx["rt"]].strip().strip("<＞>"))
        except ValueError:
            rt = None
        peaks.append({"峰名": name, "rt": rt, "area": area,
                      "norm": rule["norm"], "target": rule["target"], "备注": note})

    used = [p for p in peaks if p["norm"]]
    if not used:
        raise HangError("没有可参与归一的峰（溶剂峰/系统峰按规则表剔除了）", no)
    total = sum(p["area"] for p in used)
    if total <= 0:
        raise HangError("可参与归一的峰面积合计为 0，算不出归一含量", no)
    for p in peaks:
        p["pct"] = round(p["area"] / total * 100, 2) if p["norm"] else None

    target = next((p for p in used if p["target"]), None)
    if target is None:
        target = max(used, key=lambda p: p["area"])
        extra = "目标峰按最大峰推定（规则表里没匹配到目标峰名）"
        target = {**target, "备注": "；".join(x for x in (target["备注"], extra) if x)}

    return {"文件": path.name, "编码": enc, "样品编号": no, "样品名称": sample,
            "进样次数": int(seq), "峰": peaks, "目标峰": target, "有效峰数": len(used)}


def collect(rules):
    """一趟扫完 01_raw_data：跑得动的进 injections，跑不动的进 hangs。"""
    injections, hangs = [], []
    for path in sorted(RAW.glob("*.txt")):
        try:
            injections.append(parse_one(path, rules))
        except HangError as e:
            hangs.append({"文件": path.name, "样品编号": e.no, "原因": str(e)})
    return injections, hangs


# ===== ③ 填数：判定、平行样偏差 =====
def judge(pct):
    if pct is None:
        return "—"
    lo, hi = NORM_RANGE
    if pct < lo:
        return "低于下限"
    if pct > hi:
        return "高于上限"
    return "合格"


def build_summaries(injections):
    out = []
    for inj in injections:
        t = inj["目标峰"]
        out.append({"文件": inj["文件"], "样品编号": inj["样品编号"],
                    "样品名称": inj["样品名称"], "进样次数": inj["进样次数"],
                    "有效峰数": inj["有效峰数"], "目标峰名": t["峰名"],
                    "目标峰保留时间": t["rt"], "目标峰面积": t["area"],
                    "主峰归一": t["pct"], "含量判定": judge(t["pct"]),
                    "平行样偏差": None, "备注": t["备注"]})
    return out


def mark_parallel(summaries):
    """同一编号的多次进样：不该并成一条（并了面积就翻倍），只算偏差、超限提示复测。"""
    groups = {}
    for s in summaries:
        groups.setdefault(s["样品编号"], []).append(s)
    for group in groups.values():
        if len(group) < 2:
            continue
        vals = [g["主峰归一"] for g in group]
        mean = sum(vals) / len(vals)
        rpd = round(abs(max(vals) - min(vals)) / mean * 100, 2) if mean else None
        for g in group:
            g["平行样偏差"] = rpd
            if rpd is not None and rpd > PARALLEL_LIMIT:
                g["备注"] = "；".join(x for x in
                                     (g["备注"], f"两次进样相差 {rpd}%，建议复测") if x)


# ===== ④ 输出：三页 Excel =====
def write_workbook(injections, summaries, hangs, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="B0B0B0")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_font = Font(bold=True, size=10)
    head_fill = PatternFill("solid", fgColor="DCE6F1")
    body_font = Font(size=10)

    def write_sheet(ws, headers, rows, widths, text_cols=()):
        ws.append(headers)
        for i, name in enumerate(headers, 1):
            cell = ws.cell(row=1, column=i)
            cell.font = head_font
            cell.fill = head_fill
            cell.border = box
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
        for r in rows:
            ws.append(["—" if v is None else v for v in r])
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
            for cell in row:
                cell.border = box
                cell.font = body_font
                if headers[cell.column - 1] in text_cols:
                    cell.number_format = "@"
                    cell.alignment = Alignment(horizontal="left")

    wb = Workbook()

    ws1 = wb.active
    ws1.title = DETAIL_SHEET
    detail = []
    for inj in injections:
        for p in inj["峰"]:
            note = "；".join(x for x in (p["备注"],
                                        "" if p["norm"] else "按规则表剔除，不参与归一") if x)
            detail.append([inj["文件"], inj["样品编号"], inj["样品名称"], inj["进样次数"],
                           p["峰名"], p["rt"], p["area"], p["pct"],
                           "是" if p["norm"] else "否", note])
    write_sheet(ws1, ["文件名", "样品编号", "样品名称", "进样次数", "峰名",
                      "保留时间(min)", "峰面积", "面积归一(%)", "参与归一", "备注"],
                detail, [24, 12, 20, 9, 14, 13, 12, 13, 10, 30],
                text_cols=("样品编号",))

    ws2 = wb.create_sheet(SUMMARY_SHEET)
    write_sheet(ws2, ["样品编号", "样品名称", "进样次数", "有效峰数", "目标峰名",
                      "目标峰保留时间(min)", "目标峰面积", "主峰归一(%)",
                      "含量判定", "平行样偏差(%)", "备注"],
                [[s["样品编号"], s["样品名称"], s["进样次数"], s["有效峰数"],
                  s["目标峰名"], s["目标峰保留时间"], s["目标峰面积"], s["主峰归一"],
                  s["含量判定"], s["平行样偏差"], s["备注"]] for s in summaries],
                [12, 20, 9, 9, 16, 18, 13, 13, 11, 15, 34],
                text_cols=("样品编号",))

    ws3 = wb.create_sheet(HANG_SHEET)
    write_sheet(ws3, ["文件名", "样品编号", "原因"],
                [[h["文件"], h["样品编号"], h["原因"]] for h in hangs],
                [24, 12, 56], text_cols=("样品编号",))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return ws1.max_row - 1, ws2.max_row - 1, ws3.max_row - 1


# ===== 跑完回读一遍，对不上就 FAIL =====
def verify(path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    checks = []

    def load(name):
        rows = list(wb[name].iter_rows(values_only=True))
        head = [re.sub(r"（[^）]*）|\([^)]*\)", "", str(c)) if c is not None else ""
                for c in rows[0]]
        return [dict(zip(head, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]

    detail, summary, hang = load(DETAIL_SHEET), load(SUMMARY_SHEET), load(HANG_SHEET)

    def check(label, cond):
        checks.append((label, bool(cond)))

    def pct(no, seq=1):
        row = next(s for s in summary if s["样品编号"] == no and s["进样次数"] == seq)
        return row["主峰归一"]

    check(f"明细 {len(detail)} 行（应为 12）", len(detail) == 12)
    check(f"汇总 {len(summary)} 行（应为 4）", len(summary) == 4)
    check(f"挂起 {len(hang)} 行（应为 2）", len(hang) == 2)
    check("挂起清单：空白对照＝没有可参与归一的峰",
          any(h["样品编号"] == "0000184" and "没有可参与归一的峰" in h["原因"] for h in hang))
    check("挂起清单：回收率样品＝n.d. 不当作 0",
          any(h["样品编号"] == "0000185" and "n.d." in h["原因"] for h in hang))
    check("0000182 主峰归一＝98.53", pct("0000182") == 98.53)
    check("0000183 主峰归一＝99.94", pct("0000183") == 99.94)
    check("0000186 两次进样分别 98.53 / 94.83",
          pct("0000186", 1) == 98.53 and pct("0000186", 2) == 94.83)
    check("0000186 平行样偏差＝3.83 且提示复测",
          all(s["平行样偏差"] == 3.83 and "建议复测" in (s["备注"] or "")
              for s in summary if s["样品编号"] == "0000186"))
    check("0000186 第二次进样判定＝低于下限",
          next(s for s in summary if s["样品编号"] == "0000186" and s["进样次数"] == 2)
          ["含量判定"] == "低于下限")
    check("0000182 样品编号前导零没丢（读回仍是字符串）",
          any(s["样品编号"] == "0000182" for s in summary))
    check("GBK 文件的中文样品名读对了",
          any(s["样品名称"] == "复方磺胺甲噁唑片" for s in summary))
    check("0000183 目标峰按最大峰推定",
          next(s for s in summary if s["样品编号"] == "0000183")["目标峰名"] == "磺胺甲噁唑"
          and "推定" in (next(s for s in summary if s["样品编号"] == "0000183")["备注"] or ""))
    check("0000182 目标峰保留时间＝3.482",
          next(s for s in summary if s["样品编号"] == "0000182")["目标峰保留时间"] == 3.482)
    check("溶剂峰行：参与归一＝否、面积归一留空",
          all(d["参与归一"] == "否" and d["面积归一"] in (None, "—")
              for d in detail if d["峰名"] == "溶剂峰"))
    check("溶剂峰不是没有，是留在明细里（共 4 条）",
          sum(1 for d in detail if d["峰名"] == "溶剂峰") == 4)
    check("甲氧苄啶按半值估＝500 且备注写明",
          any(d["峰名"] == "甲氧苄啶" and d["峰面积"] == 500 and "半值" in (d["备注"] or "")
              for d in detail))
    check("「总计」行没被当成峰",
          all(d["峰名"] not in ("总计",) for d in detail))
    check("挂起的两条没混进明细（0000184/0000185 不出现）",
          all(d["样品编号"] not in ("0000184", "0000185") for d in detail))
    check("0000182 参与归一的峰面积归一合计＝100.00",
          round(sum(d["面积归一"] for d in detail
                    if d["样品编号"] == "0000182" and d["参与归一"] == "是"), 2) == 100.00)
    check("0000183 参与归一的峰面积归一合计＝100.00",
          round(sum(d["面积归一"] for d in detail
                    if d["样品编号"] == "0000183" and d["参与归一"] == "是"), 2) == 100.00)
    check("0000182 含量判定＝合格",
          next(s for s in summary if s["样品编号"] == "0000182")["含量判定"] == "合格")
    check("单次进样的样品，平行样偏差留空",
          next(s for s in summary
               if s["样品编号"] == "0000182")["平行样偏差"] in (None, "—"))

    failed = [c for c in checks if not c[1]]
    for label, ok in checks:
        print(f"  {'OK  ' if ok else 'FAIL'} {label}")
    print(f"回读校验：{len(checks) - len(failed)}/{len(checks)} 项通过")
    return not failed


# ===== 工程外壳：日志、依赖、暂停 =====
def host_info():
    name = socket.gethostname()
    ip = "未知"
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
    except OSError:
        try:
            ip = socket.gethostbyname(name)
        except OSError:
            pass
    return f"{name}（{ip}）"


def ensure_deps():
    try:
        import openpyxl  # noqa: F401
        return
    except ImportError:
        pass
    for extra in ([], ["--user"]):
        cmd = [sys.executable, "-m", "pip", "install", "-q", "openpyxl",
               "-i", "https://pypi.tuna.tsinghua.edu.cn/simple"] + extra
        if subprocess.call(cmd) == 0:
            return
    raise SystemExit("openpyxl 装不上，请手动执行：pip install openpyxl")


def log(text):
    LOG.parent.mkdir(parents=True, exist_ok=True)
    with LOG.open("a", encoding="utf-8") as fh:
        fh.write(text + "\n")


def main():
    ensure_deps()
    t0 = time.time()
    log(f"运行主机：{host_info()}")
    log(f"开始执行：{time.strftime('%Y-%m-%d %H:%M:%S')}  "
        f"Python {sys.version.split()[0]}（{Path(sys.executable).name}）")

    if OUT.exists():
        shutil.rmtree(OUT)          # 重跑先清空，免得上一版的文件被当成本次结果
    OUT.mkdir(parents=True)

    rules = load_rules(RULES)
    injections, hangs = collect(rules)
    summaries = build_summaries(injections)
    mark_parallel(summaries)
    n_detail, n_sum, n_hang = write_workbook(injections, summaries, hangs, OUT / OUT_NAME)

    print(f"峰表文件 {len(injections) + len(hangs)} 份："
          f"进明细 {len(injections)} 次进样 / 挂起 {len(hangs)} 份")
    print(f"写盘：{OUT / OUT_NAME}（峰明细 {n_detail} 行、样品汇总 {n_sum} 行、"
          f"挂起清单 {n_hang} 行）")

    ok = verify(OUT / OUT_NAME)
    cost = time.time() - t0
    log(f"执行完成：{time.strftime('%Y-%m-%d %H:%M:%S')}  用时 {cost:.1f}s  "
        f"判定 {'PASS' if ok else 'FAIL'}")
    print(f"判定：{'PASS' if ok else 'FAIL'}  用时 {cost:.1f}s")

    if "--no-pause" not in sys.argv:
        input("按回车键关闭窗口……")


if __name__ == "__main__":
    main()

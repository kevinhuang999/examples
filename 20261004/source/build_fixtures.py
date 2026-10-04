# -*- coding: utf-8 -*-
"""重造本示例用到的模板与数据。

跑法：python source/build_fixtures.py
产物：
  01_raw_data/报表口径.xlsx              业务侧维护的口径表（中心字典 / 项目口径 / 报表定义）
  01_raw_data/lims_export_结果明细.csv   系统导出的结果明细（本示例的输入数据）
  source/templates/月度质量报表模板.xlsx 报表模板（表头 + 12 行明细区 + 合计行）

上半部分建两张表（口径表、报表模板），下半部分造结果明细 CSV。
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent      # 包根 = 20261004/
RAW = ROOT / "01_raw_data"
TPL = ROOT / "source" / "templates"

BOLD = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center")
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
THIN = Side(style="thin", color="9BAFC4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

REPORT_COLS = ["检测项目", "单位", "检测项数", "合格", "不合格", "未判定", "合格率(%)", "结果均值"]
DETAIL_COLS = ["样本编号", "中心代码", "检测项目", "结果", "单位", "判定", "状态", "检测日期"]

# ---------------------------------------------------------------- 口径表数据
# 中心字典：代码 -> 名称。注意 03 那行名称尾巴上有个全角空格，用来测归一化。
CENTER_DICT = [
    ["01", "上海中心"],
    ["02", "北京中心"],
    ["03", "广州中心\u3000"],
]

# 项目口径：来源单位 -> 目标单位 的换算系数，以及报表要求的小数位。
ITEM_SPEC = [
    # 检测项目, 来源单位, 目标单位, 换算系数, 小数位
    ["pH", "", "", 1.0, 2],
    ["残留溶剂", "mg/L", "mg/dL", 0.1, 2],
    ["含量", "g/L", "g/L", 1.0, 2],
    ["有关物质", "%", "%", 1.0, 3],
    ["重金属", "mg/kg", "mg/kg", 1.0, 2],
    ["微生物限度", "CFU/g", "CFU/g", 1.0, 1],
    ["干燥失重", "%", "%", 1.0, 2],
    ["炽灼残渣", "%", "%", 1.0, 3],
]

REPORT_DEF = [
    ["月度质量报表", "月度质量报表_2026-09.xlsx", "01|02|03"],
]

# ---------------------------------------------------------------- 明细数据
# 样本编号, 中心代码, 检测项目, 结果, 单位, 判定, 状态, 检测日期
DETAIL = [
    # —— 上海中心：8 个项目，编号是纯数字带前导零 ——
    ["0000713", "01", "pH", "6.82", "", "合格", "已出报告", "2026-09-02"],
    ["0000714", "01", "pH", "6.34", "", "合格", "已出报告", "2026-09-02"],
    ["0000021", "01", "残留溶剂", "12.5", "mg/L", "合格", "已出报告", "2026-09-03"],
    ["0000022", "01", "残留溶剂", "<0.01", "mg/L", "合格", "已出报告", "2026-09-03"],
    ["0000030", "01", "含量", "98.6", "g/L", "合格", "已出报告", "2026-09-04"],
    ["0000031", "01", "含量", "95.2", "g/L", "不合格", "已出报告", "2026-09-04"],
    ["0000040", "01", "有关物质", "0.42", "%", "合格", "已出报告", "2026-09-05"],
    ["0000041", "01", "有关物质", "1.86", "%", "不合格", "已出报告", "2026-09-05"],
    ["0000050", "01", "重金属", "未检出", "mg/kg", "合格", "已出报告", "2026-09-06"],
    ["0000051", "01", "重金属", "2.4", "mg/kg", "合格", "已出报告", "2026-09-06"],
    ["0000060", "01", "微生物限度", "120", "CFU/g", "合格", "已出报告", "2026-09-07"],
    ["0000070", "01", "干燥失重", "0.31", "%", "合格", "已出报告", "2026-09-07"],
    ["0000080", "01", "炽灼残渣", "0.08", "%", "—", "已出报告", "2026-09-08"],
    ["0000081", "01", "炽灼残渣", "", "", "", "已作废", "2026-09-08"],
    ["0000090", "01", "pH", "6.50", "", "", "复检中", "2026-09-09"],
    ["0000091", "01", "残留溶剂", "11.0", "mg/L", "合格", "", "2026-09-09"],
    # —— 北京中心：只有 2 个项目，明细区大半是空的 ——
    ["BJ-0001", "02", "pH", "7.05", "", "合格", "已出报告", "2026-09-03"],
    ["BJ-0002", "02", "含量", "99.1", "g/L", "合格", "已出报告", "2026-09-04"],
    ["BJ-0003", "02", "pH", "6.97", "", "合格", "已出报告", "2026-09-05"],
    ["BJ-0004", "02", "含量", "101.5", "g/L", "不合格", "已出报告", "2026-09-06"],
    # —— 广州中心：单量最多；有一条把单位写进了结果列 ——
    ["GZ-1001", "03", "pH", "6.55", "", "合格", "已出报告", "2026-09-02"],
    ["GZ-1002", "03", "pH", "6.48", "", "合格", "已出报告", "2026-09-03"],
    ["GZ-1003", "03", "pH", "6.72", "", "合格", "已出报告", "2026-09-04"],
    ["GZ-1010", "03", "含量", "98.6 g/L", "", "合格", "已出报告", "2026-09-05"],
    ["GZ-1011", "03", "含量", "97.4", "g/L", "合格", "已出报告", "2026-09-05"],
    ["GZ-1012", "03", "含量", "94.8", "g/L", "不合格", "已出报告", "2026-09-06"],
    ["GZ-1013", "03", "含量", "99.2", "g/L", "合格", "已出报告", "2026-09-06"],
    ["GZ-1020", "03", "重金属", "2.1", "mg/kg", "合格", "已出报告", "2026-09-07"],
    ["GZ-1021", "03", "重金属", "未检出", "mg/kg", "合格", "已出报告", "2026-09-07"],
    ["GZ-1022", "03", "重金属", "<0.5", "mg/kg", "合格", "已出报告", "2026-09-08"],
    ["GZ-1030", "03", "干燥失重", "0.25", "%", "合格", "已出报告", "2026-09-08"],
    ["GZ-1031", "03", "干燥失重", "0.33", "%", "合格", "已出报告", "2026-09-09"],
    ["GZ-1040", "03", "炽灼残渣", "0.06", "%", "—", "已出报告", "2026-09-09"],
    ["GZ-1041", "03", "炽灼残渣", "0.10", "%", "合格", "已出报告", "2026-09-09"],
    ["GZ-1050", "03", "pH", "6.60", "", "合格", "已作废", "2026-09-09"],
]


def write_sheet(ws, head, rows, widths, note):
    """写一张小表：表头 + 数据行 + 表尾说明行（说明行的判空列留空）。"""
    ws.append(head)
    for c in range(1, len(head) + 1):
        ws.cell(1, c).font = BOLD
        ws.cell(1, c).fill = HEAD_FILL
        ws.cell(1, c).border = BOX
    for r in rows:
        ws.append(r)
    for r in range(2, ws.max_row + 1):
        for c in range(1, len(head) + 1):
            ws.cell(r, c).border = BOX
    ws.append([""] + [note])              # 判空列留空才会被当成表尾说明跳过
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w


def build_spec():
    """建口径表：中心字典 / 项目口径 / 报表定义。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "中心字典"
    write_sheet(ws, ["中心代码", "中心名称"], CENTER_DICT, [10, 16],
                "说明：中心名称按系统里的叫法维护，报表按这里的名字建工作表")

    ws = wb.create_sheet("项目口径")
    write_sheet(ws, ["检测项目", "来源单位", "目标单位", "换算系数", "小数位"],
                ITEM_SPEC, [14, 10, 10, 10, 8],
                "说明：换算系数 = 目标单位 ÷ 来源单位；小数位是报表里的显示位数")

    ws = wb.create_sheet("报表定义")
    write_sheet(ws, ["报表名", "输出文件名", "中心列表"], REPORT_DEF, [16, 28, 14],
                "说明：中心列表用竖线分隔，几个中心就出几个工作表")

    out = RAW / "报表口径.xlsx"
    wb.save(out)
    wb.close()
    print(f"已生成口径表：{out.name}（3 个工作表）")


def build_template():
    """建报表模板：标题行 + 中心标签行 + 表头 + 12 行明细区 + 合计行。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "报表模板"

    ws["A1"] = "月度检测汇总报表"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")
    ws["A1"].alignment = CENTER

    ws["A3"] = "中心："
    ws["A3"].font = BOLD
    ws["A4"] = "口径：只统计状态为「已出报告」的记录；合格率 = 合格 ÷（合格 + 不合格）"
    ws["A4"].font = Font(size=9, color="808080")

    head_row = 6
    for j, col in enumerate(REPORT_COLS, 1):
        cell = ws.cell(head_row, j, col)
        cell.font = BOLD
        cell.fill = HEAD_FILL
        cell.border = BOX
        cell.alignment = CENTER

    detail_rows = 12
    for r in range(head_row + 1, head_row + 1 + detail_rows):
        for j in range(1, len(REPORT_COLS) + 1):
            ws.cell(r, j).border = BOX          # 明细区先画好边框，脚本按行填
    total_row = head_row + 1 + detail_rows
    ws.cell(total_row, 1, "合计").font = BOLD

    for col, w in zip("ABCDEFGH", [14, 10, 10, 8, 8, 8, 11, 11]):
        ws.column_dimensions[col].width = w

    out = TPL / "月度质量报表模板.xlsx"
    wb.save(out)
    wb.close()
    print(f"已生成报表模板：{out.name}（表头第 {head_row} 行，明细区 {detail_rows} 行，合计第 {total_row} 行）")


def build_detail_csv():
    """造结果明细 CSV。表尾加一行说明，用来测「判空列留空就跳过」。"""
    out = RAW / "lims_export_结果明细.csv"
    with out.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(DETAIL_COLS)
        w.writerows(DETAIL)
        w.writerow([])
        w.writerow(["", "", "说明：本文件由系统导出，判定栏「—」表示未判定", "", "", "", "", ""])
    print(f"已生成结果明细：{out.name}（数据 {len(DETAIL)} 行 + 1 行表尾说明）")


if __name__ == "__main__":
    RAW.mkdir(parents=True, exist_ok=True)
    TPL.mkdir(parents=True, exist_ok=True)
    build_spec()
    build_template()
    build_detail_csv()
    print("造数完成：模板与口径表在上半段，明细 CSV 在下半段。")

#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把一个文件夹里各台仪器导出的结果文件批量解析成一张统一的汇总表。

用法：
    python run.py              # 跑完暂停，方便双击看结果
    python run.py --no-pause   # 不暂停

依赖：openpyxl，缺了会自动走清华源安装。

输入：01_raw_data/ 下各仪器导出的 .txt / .csv / .xlsx（编码、分隔符、表头位置都不一样）
输出：02_output/仪器数据汇总_20261005.xlsx（汇总 + 按仪器统计 + 按项目统计 + 挂起清单）
"""

import csv
import re
import socket
import subprocess
import sys
import time
from pathlib import Path

BASE = Path(__file__).resolve().parent
RAW_DIR = BASE / "01_raw_data"
OUT_DIR = BASE / "02_output"
LOG_FILE = BASE / "source" / "run_log.txt"
OUT_FILE = OUT_DIR / "仪器数据汇总_20261005.xlsx"

# —— 汇总表的列：不管源文件长什么样，落到这张表里就这几列 ——
COLUMNS = ["仪器", "源文件", "样品编号", "样品名称", "批号",
           "检验项目", "结果原值", "数值", "单位", "形态", "分析日期"]
TEXT_COLS = {"源文件", "样品编号", "批号"}
DATE_COL = "分析日期"

# 结果列的三种形态：纯数值、只给边界的 <0.01、文字结论（未检出 / n.d.）
RE_LIMIT = re.compile(r"^([<>]=?)\s*([-+]?\d+(?:\.\d+)?)$")
RE_NUM_UNIT = re.compile(r"^([-+]?\d+(?:\.\d+)?)\s+([A-Za-z%][A-Za-z0-9/%.]*)$")

LAB_COLS = ["样品编号", "样品名称", "批号", "检验项目", "结果", "单位", "分析日期"]


def make_spec(instrument, suffix, encoding, delimiter, marks, mapping=None, key="样品编号"):
    """一种仪器导出文件的「样子」：靠后缀 + 表头特征行认，不靠文件名里的仪器名。"""
    if mapping is None:
        mapping = {name: name for name in LAB_COLS}
    return {"仪器": instrument, "后缀": suffix, "编码": encoding, "分隔符": delimiter,
            "表头特征": marks, "列映射": mapping, "判空列": key}


FORMATS = [
    make_spec("HPLC 工作站", ".txt", "gbk", "\t", ["样品编号", "检验项目", "结果"]),
    make_spec("GC 工作站", ".csv", "utf-8-sig", ",", ["样品编号", "检验项目", "结果"]),
    make_spec("紫外分光光度计", ".xlsx", None, None, ["Sample", "Item", "Result"],
              {"Sample": "样品编号", "Item": "检验项目", "Result": "结果",
               "Unit": "单位", "Date": "分析日期"},
              key="Sample"),
]


def norm(text):
    """去掉所有空白（含全角空格），比较列名 / 样品编号 / 批号之前一律先过它。"""
    return "".join(str(text if text is not None else "").split())


def host_tag():
    """本段日志跑在哪台机器上：主机名（本机对外网卡 IP）。"""
    name = "未知"
    try:
        name = socket.gethostname()
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        try:
            sock.connect(("8.8.8.8", 80))
            return f"{name}（{sock.getsockname()[0]}）"
        finally:
            sock.close()
    except OSError:
        try:
            return f"{name}（{socket.gethostbyname(name)}）"
        except OSError:
            return f"{name}（未知）"


def log_line(text):
    with LOG_FILE.open("a", encoding="utf-8") as fh:
        fh.write(text + "\n")


def ensure_deps():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("缺少 openpyxl，改用清华源安装……")
        subprocess.check_call([sys.executable, "-m", "pip", "install",
                               "-i", "https://pypi.tuna.tsinghua.edu.cn/simple", "openpyxl"])


def read_text_any(path):
    """仪器工作站导出的 txt 常是 GBK，直接按 utf-8 读会抛 UnicodeDecodeError（不是乱码）。"""
    for enc in ("utf-8-sig", "gbk", "utf-8"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"{path.name}：三种编码都读不出来，先确认文件没有损坏")


def find_header(rows, marks):
    """找表头行：同一行里出现全部特征列名才算。

    不按「第几行」找——各仪器前面压的仪器信息行数不一样。
    """
    for i, row in enumerate(rows):
        cells = [norm(c) for c in row]
        if all(any(mark == cell for cell in cells) for mark in marks):
            return i
    return None


def read_one(path, spec):
    """按格式表读一个文件，返回 (记录列表, 跳过原因)。"""
    if spec["后缀"] == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        rows = [[c.value for c in row] for row in wb.active.iter_rows()]
        wb.close()
    else:
        rows = list(csv.reader(read_text_any(path).splitlines(), delimiter=spec["分隔符"]))

    head = find_header(rows, spec["表头特征"])
    if head is None:
        return [], "表头特征行找不到，不是这种仪器的导出文件"

    titles = [norm(c) for c in rows[head]]
    out = []
    for row in rows[head + 1:]:
        cells = {titles[j]: (row[j] if j < len(row) else None) for j in range(len(titles))}
        if norm(cells.get(norm(spec["判空列"]))) == "":
            continue                      # 判空列一空，就是表尾的合计行 / 说明行
        rec = {"仪器": spec["仪器"], "源文件": path.name}
        for src, dst in spec["列映射"].items():
            rec[dst] = cells.get(norm(src))
        out.append(rec)

    if not out:
        return [], "表头下面没有数据行"
    return out, None


def parse_result(raw, unit_cell):
    """结果原文 → {结果原值, 数值, 形态, 单位}。

    仪器导出的结果列有三种写法：98.6、<0.01、「未检出」。
    第三种不能当 0 算，第二种也不能当 0.01 算，都要分开对待。
    """
    text = "" if raw is None else str(raw).strip()
    unit = "" if unit_cell is None else str(unit_cell).strip()
    if text == "":
        return {"结果原值": "", "数值": None, "形态": "空", "单位": unit}

    hit = RE_LIMIT.match(text)
    if hit:
        value = float(hit.group(2)) / 2 if hit.group(1) == "<" else float(hit.group(2))
        return {"结果原值": text, "数值": round(value, 6), "形态": "限度", "单位": unit}

    hit = RE_NUM_UNIT.match(text)
    if hit:
        return {"结果原值": text, "数值": float(hit.group(1)), "形态": "数值",
                "单位": unit or hit.group(2)}

    try:
        return {"结果原值": text, "数值": float(text), "形态": "数值", "单位": unit}
    except ValueError:
        return {"结果原值": text, "数值": None, "形态": "文字", "单位": unit}


def collect(raw_dir):
    """一个文件夹里的导出文件逐个读进来，去重后归一成一张宽表。"""
    records, pendings, seen, conflicted = [], [], {}, set()
    for path in sorted(raw_dir.iterdir()):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        spec = next((c for c in FORMATS if path.suffix.lower() == c["后缀"]), None)
        if spec is None:
            pendings.append({"源文件": path.name, "样品编号": "", "检验项目": "",
                             "原因": "后缀不认识，只认 .txt/.csv/.xlsx"})
            continue

        rows, why = read_one(path, spec)
        if why:
            pendings.append({"源文件": path.name, "样品编号": "", "检验项目": "", "原因": why})
            continue

        for rec in rows:
            parsed = parse_result(rec.pop("结果", None), rec.pop("单位", None))
            rec.update(parsed)
            if parsed["形态"] == "空":
                pendings.append({**rec, "原因": "结果列为空，凑不出这条记录"})
                continue

            key = (rec["仪器"], norm(rec["样品编号"]), norm(rec["检验项目"]))
            if key in conflicted:
                continue
            if key in seen:
                first = seen[key]
                if norm(first["结果原值"]) == norm(rec["结果原值"]):
                    pendings.append({**rec, "原因": "重复导出，内容一致，只留先读到的那份"})
                else:
                    records.remove(first)
                    seen.pop(key)
                    conflicted.add(key)
                    both = "同一样品同一项目结果不一致，两条都不汇总"
                    pendings.append({**first, "原因": both})
                    pendings.append({**rec, "原因": both})
                continue

            seen[key] = rec
            records.append(rec)
    return records, pendings


def write_sheet(ws, head, rows):
    """一张表 = 表头 + 数据；按列名决定写文本格式还是日期格式。"""
    from openpyxl.styles import Border, Font, PatternFill, Side

    thin = Side(style="thin", color="9BAFC4")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, title in enumerate(head, 1):
        cell = ws.cell(1, j, title)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
        cell.border = box

    for i, row in enumerate(rows, start=2):
        for j, title in enumerate(head, 1):
            cell = ws.cell(i, j, row.get(title))
            cell.border = box
            if title in TEXT_COLS:
                cell.number_format = "@"          # 批号 00903、样品编号 0000713 都带前导零
            elif title == DATE_COL:
                cell.number_format = "yyyy-mm-dd"  # 不然 Excel 显示成 46265

    for j, title in enumerate(head, 1):
        ws.column_dimensions[ws.cell(1, j).column_letter].width = max(10, len(title) * 2 + 4)


def build_output(records, pendings):
    """三个部分：汇总宽表、两张统计表、挂起清单。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    write_sheet(ws, COLUMNS, records)

    by_inst = {}
    for rec in records:
        stat = by_inst.setdefault(rec["仪器"], [0, 0, 0])
        stat[0] += 1
        stat[1 if rec["形态"] in ("数值", "限度") else 2] += 1
    inst_rows = [{"仪器": name, "条数": stat[0], "可算数值": stat[1], "文字结论": stat[2]}
                 for name, stat in sorted(by_inst.items(), key=lambda kv: -kv[1][0])]
    write_sheet(wb.create_sheet("按仪器统计"), ["仪器", "条数", "可算数值", "文字结论"], inst_rows)

    by_item = {}
    for rec in records:
        by_item.setdefault((rec["检验项目"], rec["单位"] or "—"), []).append(rec["数值"])
    item_rows = []
    for (name, unit), values in sorted(by_item.items(), key=lambda kv: (-len(kv[1]), kv[0])):
        nums = [v for v in values if v is not None]
        item_rows.append({"检验项目": name, "单位": unit, "条数": len(values),
                          "数值均值": round(sum(nums) / len(nums), 2) if nums else "—"})
    write_sheet(wb.create_sheet("按项目统计"), ["检验项目", "单位", "条数", "数值均值"], item_rows)

    write_sheet(wb.create_sheet("挂起清单"), ["源文件", "样品编号", "检验项目", "原因"], pendings)

    wb.save(OUT_FILE)
    wb.close()
    return inst_rows, item_rows


def verify():
    """生成后回读断言：不读回来一遍，不敢说解析对了。"""
    from openpyxl import load_workbook

    stats = {"ok": 0, "fail": 0}

    def check(label, got, want):
        if got == want:
            stats["ok"] += 1
            print(f"  OK   {label} = {got!r}")
        else:
            stats["fail"] += 1
            print(f"  FAIL {label} = {got!r}，期望 {want!r}")

    wb = load_workbook(OUT_FILE)
    check("工作表清单", wb.sheetnames, ["汇总", "按仪器统计", "按项目统计", "挂起清单"])

    ws = wb["汇总"]
    head = {norm(ws.cell(1, j).value): j for j in range(1, ws.max_column + 1)}
    rows = {}
    for r in range(2, ws.max_row + 1):
        key = (ws.cell(r, head["样品编号"]).value, norm(ws.cell(r, head["检验项目"]).value))
        rows[key] = {name: ws.cell(r, j).value for name, j in head.items()}

    check("汇总条数", len(rows), 9)
    check("HPLC 含量 数值", rows[("0000713", "含量")]["数值"], 99.6)
    check("只给边界的两列",
          [rows[("0000714", "有关物质")]["结果原值"], rows[("0000714", "有关物质")]["数值"]],
          ["<0.01", 0.005])
    check("只给边界 形态", rows[("0000714", "有关物质")]["形态"], "限度")
    check("未检出不算 0",
          [rows[("00903", "甲醇")]["数值"], rows[("00903", "甲醇")]["形态"]], [None, "文字"])
    check("样品编号前导零没丢", rows[("00903", "甲醇")]["样品编号"], "00903")
    check("批号按文本格式写", ws.cell(2, head["批号"]).number_format, "@")
    check("单位从结果里拆出来", rows[("UV26-001", "含量")]["单位"], "mg/L")
    check("单位列有值时优先", rows[("UV26-001", "吸收度")]["单位"], "-")
    check("日期按日期格式写", ws.cell(2, head["分析日期"]).number_format, "yyyy-mm-dd")
    check("日期读回来还是那一天",
          rows[("UV26-001", "吸收度")]["分析日期"].strftime("%Y-%m-%d"), "2026-10-05")

    wsi = wb["按仪器统计"]
    inst = {wsi.cell(r, 1).value: [wsi.cell(r, c).value for c in (2, 3, 4)]
            for r in range(2, wsi.max_row + 1)}
    check("HPLC 计数", inst["HPLC 工作站"], [4, 4, 0])
    check("GC 计数", inst["GC 工作站"], [3, 2, 1])
    check("紫外 计数", inst["紫外分光光度计"], [2, 2, 0])

    wsm = wb["按项目统计"]
    item = {(wsm.cell(r, 1).value, wsm.cell(r, 2).value):
            [wsm.cell(r, 3).value, wsm.cell(r, 4).value] for r in range(2, wsm.max_row + 1)}
    check("同名项目按单位分开", sorted(item),
          [("含量", "%"), ("含量", "mg/L"), ("吸收度", "-"),
           ("有关物质", "%"), ("残留溶剂", "%"), ("甲醇", "%")])
    check("含量 % 均值", item[("含量", "%")], [2, 99.25])
    check("含量 mg/L 不跟 % 混算", item[("含量", "mg/L")], [1, 12.3])
    check("有关物质 均值只吃数值", item[("有关物质", "%")], [2, round((0.12 + 0.005) / 2, 2)])
    check("残留溶剂 均值", item[("残留溶剂", "%")], [2, 0.17])
    check("甲醇 均值留空", item[("甲醇", "%")], [1, "—"])

    wsp = wb["挂起清单"]
    pend = [(wsp.cell(r, 1).value, wsp.cell(r, 4).value) for r in range(2, wsp.max_row + 1)]
    check("挂起条数", len(pend), 6)
    check("认不出的文件写了原因",
          any(name == "天平导出_原始记录_20261005.txt" and "找不到" in str(why)
              for name, why in pend), True)
    check("只有表头的文件挂起",
          any(name == "酶标仪_结果_20261005.csv" and why == "表头下面没有数据行"
              for name, why in pend), True)
    check("结果为空挂起", sum(1 for n, w in pend if w == "结果列为空，凑不出这条记录"), 1)
    check("冲突两条都挂起", sum(1 for n, w in pend if w and "不一致" in w), 2)
    check("重复导出只留一份", sum(1 for n, w in pend if w and "重复导出" in w), 1)

    wb.close()
    print(f"  回读断言：{stats['ok']} 项 OK，{stats['fail']} 项 FAIL")
    return stats


def main():
    started = time.time()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
    for old in OUT_DIR.iterdir():
        if old.is_file():
            old.unlink()                       # 重跑先清，免得上一版多出来的文件被当这一版

    ensure_deps()
    log_line(f"运行主机：{host_tag()}")
    log_line(f"开始执行：{time.strftime('%Y-%m-%d %H:%M:%S')}  "
             f"Python {sys.version.split()[0]}（{Path(sys.executable).name}）")

    records, pendings = collect(RAW_DIR)
    log_line(f"读入：{len(records)} 条记录 / {len(pendings)} 条挂起")
    print(f"解析出 {len(records)} 条记录，挂起 {len(pendings)} 条")

    inst_rows, item_rows = build_output(records, pendings)
    print(f"生成 {OUT_FILE.name}：汇总 {len(records)} 条 / 仪器 {len(inst_rows)} 台 / "
          f"项目 {len(item_rows)} 组 / 挂起 {len(pendings)} 条")

    print("生成后回读校验：")
    stats = verify()
    passed = stats["fail"] == 0
    verdict = "PASS" if passed else "FAIL"
    span = time.time() - started
    log_line(f"执行完成：{time.strftime('%Y-%m-%d %H:%M:%S')}  用时 {span:.1f}s  判定 {verdict}")
    print(f"判定：{verdict}（用时 {span:.1f}s）")

    if "--no-pause" not in sys.argv and sys.stdin and sys.stdin.isatty():
        input("按回车结束……")
    return 0 if passed else 1


if __name__ == "__main__":
    sys.exit(main())

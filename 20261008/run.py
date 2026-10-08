# -*- coding: utf-8 -*-
"""ELN 导出数据整合成一张 Excel 汇总表 —— 可运行示例

场景：ELN（电子实验记录本）按批次分次导出实验记录，各次导出用的模板版本不同，
字段名、列顺序、编码都不一样。本示例按一张外置的字段映射表，把它们的共同字段
归并成一张汇总表；认不出的文件与缺主键的行挂起，列级映射全部留痕。

用法：
    python run.py              # 跑完暂停（双击也能用）
    python run.py --no-pause   # 不暂停（脚本 / CI 里跑）
"""
import argparse
import csv
import io
import re
import socket
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

BASE = Path(__file__).resolve().parent
RAW_DIR = BASE / "01_raw_data"
OUT_DIR = BASE / "02_output"
SRC_DIR = BASE / "source"
TPL_DIR = SRC_DIR / "templates"
MAP_FILE = TPL_DIR / "字段映射表.csv"
LOG_FILE = SRC_DIR / "run_log.txt"

DATA_SUFFIX = {".csv", ".xlsx", ".xls"}
ROW_KEY = "实验编号"          # 主键：缺了没法定这一行是谁
TAIL_COL = "批号"             # 判空列：这一列为空即数据区结束（表尾说明行）
REQUIRED = ["实验编号", "批号", "原料"]
NUM_COLS = ["投料量_g", "温度_C", "时间_h", "收率_pct"]
OUT_NAME = "ELN实验数据汇总表.xlsx"


def ensure_deps():
    """缺依赖自动装（清华源）；已装则直接返回。"""
    need = []
    for mod, pkg in (("pandas", "pandas"), ("openpyxl", "openpyxl")):
        try:
            __import__(mod)
        except ImportError:
            need.append(pkg)
    if not need:
        return
    print("缺少依赖：", ", ".join(need), "，正在从清华源安装……")
    cmd = [sys.executable, "-m", "pip", "install", "-i",
           "https://pypi.tuna.tsinghua.edu.cn/simple", *need]
    if subprocess.call(cmd) != 0:
        subprocess.call(cmd + ["--user"])


def norm(name):
    """列名归一：去空白、全角括号转半角、英文转小写。

    ELN 不同版本对同一字段的写法会飘（`投料量(g)` / `投料量（g）`），
    所以匹配前先归一，而不是要求导出的列名一尘不染。
    """
    s = str(name) if name is not None else ""
    s = s.replace("（", "(").replace("）", ")")
    s = re.sub(r"[\s\u00a0\u3000]+", "", s)
    return s.lower()


def load_mapping(path):
    """读字段映射表，返回 (标准字段顺序, 归一别名 → 标准字段)。

    映射表是业务资产：写在表格里，不写死在代码里——ELN 改了模板名改表不改代码，
    哪一版加的别名表里也看得见。
    """
    with path.open(encoding="utf-8-sig", newline="") as f:
        rows = [r for r in csv.reader(f) if any(str(c).strip() for c in r)]
    order, alias_map = [], {}
    for r in rows[1:]:
        std = r[0].strip()
        if not std:
            continue
        order.append(std)
        for a in str(r[1]).split("|"):
            a = norm(a)
            if not a:
                continue
            if a in alias_map and alias_map[a] != std:
                raise ValueError(f"别名 {a!r} 同时指向 {alias_map[a]!r} 与 {std!r}，映射表有冲突")
            alias_map[a] = std
    return order, alias_map


def read_text(path):
    """按 utf-8-sig → gbk → utf-8 依次试；GBK 文件用 utf-8 读会直接抛异常，不是读出乱码。"""
    last = None
    for enc in ("utf-8-sig", "gbk", "utf-8"):
        try:
            return path.read_text(encoding=enc), enc
        except UnicodeDecodeError as e:
            last = e
    raise last


def read_matrix(path):
    """把一个导出文件读成二维表（不做任何解释），并返回实际用的编码。"""
    if path.suffix.lower() == ".csv":
        text, enc = read_text(path)
        return list(csv.reader(io.StringIO(text))), enc
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows, "xlsx"


def find_header(rows, alias_map):
    """找表头行：第一个「命中已知别名 >= 2」的行。

    不能写死 header=0——ELN 导出常在表头上压 2~3 行导出信息；
    也不能靠行号猜，模板改一版行号就全错。
    """
    for i, r in enumerate(rows):
        hits = sum(1 for c in r if norm(c) in alias_map)
        if hits >= 2:
            return i
    return -1


def map_columns(rows, hdr, alias_map):
    """把原始表头按别名映射成标准字段，返回 (标准字段 → 列下标, 留痕, 丢弃列)。"""
    headers = [str(c) if c is not None else "" for c in rows[hdr]]
    std_cols, trace, dropped = {}, [], []
    for i, h in enumerate(headers):
        if not norm(h):
            continue
        std = alias_map.get(norm(h))
        if std is None:
            dropped.append(h)
            trace.append((h, "", "丢弃", "映射表里没有这个名字"))
        elif std in std_cols:
            dropped.append(h)
            trace.append((h, std, "丢弃", "同名字段已在本表出现"))
        else:
            std_cols[std] = i
            trace.append((h, std, "映射", ""))
    return std_cols, trace, dropped


def cell(cells, i):
    return cells[i] if i < len(cells) else ""


def normalize_value(v):
    """单元格 → 去首尾空白的字符串；空值一律返回 ''。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return re.sub(r"^[\s\u00a0\u3000]+|[\s\u00a0\u3000]+$", "", str(v))


def merge_rows(rows, hdr, std_cols, fname):
    """切表尾、抽数据行、缺主键的行挂起。

    判定顺序＝业务顺序：先按判空列切掉表尾说明行，再判主键——
    顺序反了，表尾说明行会被当成"缺主键"混进挂起清单。
    """
    data, suspended = [], []
    for k, r in enumerate(rows[hdr + 1:]):
        cells = [normalize_value(c) for c in r]
        if not any(cells):
            continue
        if cell(cells, std_cols[TAIL_COL]) == "":
            break                      # 表尾说明行：数据区到此为止
        row = {std: cell(cells, i) for std, i in std_cols.items()}
        row["来源文件"] = fname
        if row[ROW_KEY] == "":
            suspended.append((fname, f"第 {hdr + k + 2} 行", f"缺主键 {ROW_KEY}"))
            continue
        data.append(row)
    return data, suspended


def convert_numbers(data, trace):
    """数值列能转就转成 float（Excel 里要能排序、算均值），转不了原样保留。

    只做机械判断，不做业务判断：不换算单位、不填补缺失、不把「未测」当 0。
    """
    for row in data:
        for col in NUM_COLS:
            v = row.get(col, "")
            if v == "":
                continue
            try:
                row[col] = float(v)
            except ValueError:
                trace.append((row["来源文件"], col, col, "原样保留", f"「{v}」不是数值"))
    return data


def collect(order, alias_map):
    """遍历 01_raw_data，逐个文件读、映射、归并。"""
    data, trace, suspended, files = [], [], [], []
    for path in sorted(RAW_DIR.iterdir()):
        if path.suffix.lower() not in DATA_SUFFIX:
            continue                   # 说明文件之类不参与
        files.append(path.name)
        rows, enc = read_matrix(path)
        hdr = find_header(rows, alias_map)
        if hdr < 0:
            suspended.append((path.name, "整个文件", "找不到表头（命中已知字段名的行不足 2 个）"))
            continue
        std_cols, tr, dropped = map_columns(rows, hdr, alias_map)
        missing = [c for c in REQUIRED if c not in std_cols]
        if missing:
            suspended.append((path.name, "整个文件", "缺必需列：" + "、".join(missing)))
            continue
        for h, std, act, note in tr:
            trace.append((path.name, h, std, act, note))
        part, susp = merge_rows(rows, hdr, std_cols, path.name)
        convert_numbers(part, trace)
        data.extend(part)
        suspended.extend(susp)
    return data, trace, suspended, files


def write_outputs(order, data, trace, suspended):
    """汇总表 + 字段映射留痕 + 挂起清单，三份都落在 02_output/。"""
    import pandas as pd
    from openpyxl import load_workbook
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    cols = order + ["来源文件"]
    df = pd.DataFrame([{c: r.get(c, "") for c in cols} for r in data], columns=cols)
    out = OUT_DIR / OUT_NAME
    df.to_excel(out, index=False, sheet_name="汇总")
    wb = load_workbook(out)
    ws = wb["汇总"]
    for j, name in enumerate(cols, 1):      # 文本列设成文本格式，前导零才不会被吃掉
        if name in NUM_COLS:
            continue
        for i in range(2, ws.max_row + 1):
            ws.cell(row=i, column=j).number_format = "@"
    wb.save(out)
    pd.DataFrame(trace, columns=["来源文件", "原列名", "标准字段", "动作", "说明"]).to_csv(
        OUT_DIR / "字段映射留痕.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(suspended, columns=["来源文件", "位置", "原因"]).to_csv(
        OUT_DIR / "挂起清单.csv", index=False, encoding="utf-8-sig")
    return out


def verify(order, data, trace, suspended):
    """生成后回读：逐项断言，打印 OK / FAIL，返回总判定。"""
    import pandas as pd
    from openpyxl import load_workbook
    out = OUT_DIR / OUT_NAME
    df = pd.read_excel(out, dtype=str) if out.exists() else pd.DataFrame()
    checks = [("汇总表已生成", out.exists() and out.stat().st_size > 0)]
    if not df.empty:
        checks.append(("汇总行数 = 8", len(df) == 8))
        checks.append(("表头 = 映射表顺序 + 来源文件", list(df.columns) == order + ["来源文件"]))
        checks.append(("前导零保留（0098 仍是文本）", "0098" in set(df["实验编号"].astype(str))))
        counts = df["来源文件"].value_counts().to_dict()
        checks.append(("各来源文件行数 2/2/3/1", counts == {
            "ELN_BR-2026-001.csv": 2, "ELN_BR-2026-002.csv": 2,
            "ELN_BR-2026-003.xlsx": 3, "ELN_BR-2026-004.csv": 1}))
        checks.append(("「未测」原样保留", "未测" in set(df["收率_pct"].astype(str))))
        r = df[df["实验编号"].astype(str) == "0104"]
        checks.append(("空投料量原样留空", len(r) == 1 and str(r.iloc[0]["投料量_g"]) in ("", "nan")))
        checks.append(("表尾说明行未进汇总",
                       not df["原料"].astype(str).str.contains("导出").any()))
    drops = [t for t in trace if t[3] == "丢弃"]
    checks.append(("丢弃列留痕 1 条（备注）", len(drops) == 1 and drops[0][1] == "备注"))
    checks.append(("挂起 = 1 文件级 + 1 行级", len([s for s in suspended if s[1] == "整个文件"]) == 1
                   and len([s for s in suspended if s[1] != "整个文件"]) == 1))
    checks.append(("留痕 34 条", len(trace) == 34))
    if out.exists():
        ws = load_workbook(out)["汇总"]
        checks.append(("实验编号列已设文本格式", ws.cell(row=2, column=1).number_format == "@"))
    ok = True
    for name, passed in checks:
        print(("  OK   " if passed else "  FAIL ") + name)
        ok = ok and bool(passed)
    return ok


def host_info():
    """主机名 + 本机对外 IP，用来分辨这段日志是哪台机器跑的。"""
    name = socket.gethostname()
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
    except OSError:
        try:
            ip = socket.gethostbyname(name)
        except OSError:
            ip = "未知"
    return name, ip


def append_log(text):
    """日志追加不覆盖；每段开头记运行主机，方便分辨机器。"""
    with LOG_FILE.open("a", encoding="utf-8") as f:
        f.write(text)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pause", action="store_true", help="跑完不暂停")
    args = ap.parse_args()

    ensure_deps()
    if not MAP_FILE.exists():
        print("找不到字段映射表：", MAP_FILE)
        print("先跑 source/build_fixtures.py 重造模板与数据。")
        return 1
    SRC_DIR.mkdir(parents=True, exist_ok=True)

    t0 = time.time()
    host, ip = host_info()
    append_log(f"\n运行主机：{host}（{ip}）\n开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}\n"
               f"Python {sys.version.split()[0]}（{Path(sys.executable).name}）\n")

    print("=" * 62)
    print("ELN 导出数据 → 一张 Excel 汇总表")
    print("=" * 62)
    order, alias_map = load_mapping(MAP_FILE)
    print(f"字段映射表：{len(order)} 个标准字段，{len(alias_map)} 个别名")

    data, trace, suspended, files = collect(order, alias_map)
    print(f"扫描文件 {len(files)} 个 → 整合 {len(data)} 行 / 留痕 {len(trace)} 条 / 挂起 {len(suspended)} 条")

    out = write_outputs(order, data, trace, suspended)
    print(f"已生成：{out.name}")
    ok = verify(order, data, trace, suspended)

    used = time.time() - t0
    verdict = "PASS" if ok else "FAIL"
    print(f"\n总判定：{verdict}（用时 {used:.1f}s）")
    append_log(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}  用时 {used:.1f}s  判定 {verdict}\n")
    if not args.no_pause:
        input("\n按回车键退出……")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())

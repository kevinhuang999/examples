#!/usr/bin/env python3
"""药企 QC 月报自动生成：检项清单 + 检验结果明细 → QC 月报一份。

跑法：
    python run.py                # Windows 上双击 run.py 也行
    python run.py --no-pause     # 跑完不暂停（脚本调用时用）

输入：source/templates/检项清单.csv（品种 → 检项规范，受控清单的简化版）
      01_raw_data/检验结果.csv（仪器与检验员录入的结果明细导出）
输出：02_output/QC月报.xlsx（批次汇总 / 不合格明细 / 复检留痕 / 品种合格率 /
      检项完成情况 / 待核清单 六个页签）
"""
import csv
import socket
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font
except ImportError:                     # 缺依赖自动装，走国内源
    subprocess.check_call([sys.executable, "-m", "pip", "install",
                           "-i", "https://pypi.tuna.tsinghua.edu.cn/simple", "openpyxl"])
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font

# ---------- 顶部的路径与字段常量区 ----------
BASE = Path(__file__).resolve().parent
RAW_DIR = BASE / "01_raw_data"
OUT_DIR = BASE / "02_output"
SRC_DIR = BASE / "source"
LOG_FILE = SRC_DIR / "run_log.txt"

SPEC_FILE = SRC_DIR / "templates" / "检项清单.csv"
RESULT_FILE = RAW_DIR / "检验结果.csv"
REPORT_NAME = "QC月报.xlsx"
MONTH = "2026-07"                       # 报表月份：只统计这个月的检验记录

# 列名一律按表头取，不按第几列——清单加一列，按序号读就整体错位
SP_PRODUCT, SP_ITEM, SP_KIND = "品种", "检项", "判定类型"
SP_LOW, SP_HIGH, SP_TEXT = "下限", "上限", "合格短语"
RS_NO, RS_LOT, RS_PRODUCT = "检验单号", "批号", "品种"
RS_ITEM, RS_VALUE = "检项", "结果"
RS_UNIT, RS_DAY, RS_RETEST = "单位", "检验日期", "复检"

NUM, TEXT = "数值", "文字"
OK, NG, PEND, OTHER = "合格", "不合格", "待检", "待核"
EMPTY = ""
RETEST_YES = "是"


def to_num(v):
    """限度和数值结果：空、空格、「—」都不是 0，是没填。区分不开，判定全错。"""
    s = str(v or "").strip().replace("—", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def read_rows(path):
    """读一张 CSV。用 utf-8-sig——导出文件的 BOM 不剥掉，第一个列名会带上隐藏字符。"""
    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load_specs():
    """检项清单 → {(品种, 检项): 判定规则}，另带全部检项名集合。

    规则必须按「品种 + 检项」两级取：同一个检项名在不同品种上的限度不同
    （阿莫西林含量 90~110、布洛芬 93~107），按检项名一层取会把两家的限度
    搅在一起，判定结果看着有依据、其实全错。检项名集合用来把「清单外检项」
    和「该品种缺规范」分成两类待核，别混成一堆。
    """
    specs, names = {}, set()
    for r in read_rows(SPEC_FILE):
        product = str(r.get(SP_PRODUCT) or "").strip()
        item = str(r.get(SP_ITEM) or "").strip()
        names.add(item)
        specs[(product, item)] = {
            "kind": str(r.get(SP_KIND) or "").strip(),
            "low": to_num(r.get(SP_LOW)),
            "high": to_num(r.get(SP_HIGH)),
            "text": str(r.get(SP_TEXT) or "").strip(),
        }
    return specs, names


def load_results():
    """结果明细 → 本月的行 + 不在本月的行。

    批号、检验单号全程按字符串走：`007`、`0000121` 这类带前导零的编号一转
    int 就变成 7 和 121，跟清单和台账里别的写法对不上，回读也还原不回来。
    日期只取到年月，年月不等于 MONTH 的行单独放——直接过滤掉，月底对不上数
    就再也查不出是哪几行没进来。
    """
    rows, other_month = [], []
    for r in read_rows(RESULT_FILE):
        day = str(r.get(RS_DAY) or "").strip()[:7]
        row = {
            "no": str(r.get(RS_NO) or "").strip(),
            "lot": str(r.get(RS_LOT) or "").strip(),
            "product": str(r.get(RS_PRODUCT) or "").strip(),
            "item": str(r.get(RS_ITEM) or "").strip(),
            "value": str(r.get(RS_VALUE) or "").strip(),
            "unit": str(r.get(RS_UNIT) or "").strip(),
            "day": day,
            "retest": str(r.get(RS_RETEST) or "").strip() == RETEST_YES,
        }
        (rows if day == MONTH else other_month).append(row)
    return rows, other_month


def spec_note(spec):
    """把限度写成一行字，报表「判定依据」列直接能看。"""
    if spec["low"] is None and spec["high"] is not None:
        return f"≤ {spec['high']}"
    if spec["low"] is not None and spec["high"] is None:
        return f"≥ {spec['low']}"
    return f"{spec['low']} ~ {spec['high']}"


def judge_one(row, spec):
    """单条结果判定 → (判定, 判定依据)。空结果算待检，不算不合格。

    空值当不合格是最常见的一种错：那管还没检，报表先把批次打成不合格，
    再让人回来挨个解释。所以空值单独一档，既不进分母也不进不合格明细。
    数值型走闭区间——结果正好等于上限或下限，判合格。
    """
    if not row["value"]:
        return PEND, "结果为空，还没检"
    if spec is None:
        return OTHER, "清单里没有这个品种的这条检项"
    if spec["kind"] == TEXT:
        want = spec["text"]
        if not want:
            return OTHER, "清单里没写合格短语"
        return (OK if row["value"] == want else NG), f"判定短语「{want}」"
    num = to_num(row["value"])
    if num is None:
        return OTHER, f"数值型结果「{row['value']}」不是数字"
    if spec["low"] is None and spec["high"] is None:
        return OTHER, "清单里下限上限都是空的"
    if spec["low"] is not None and num < spec["low"]:
        return NG, f"低于下限 {spec['low']}"
    if spec["high"] is not None and num > spec["high"]:
        return NG, f"高于上限 {spec['high']}"
    return OK, spec_note(spec)


def last_one(rows):
    """同一条检项的多条结果：按检验日期取最后一次。

    初检和复检不在同一天时靠日期分；同一批号同一条检项在同一天写了两行，
    说明是补录，表里靠后写的那行才是最终值，所以按日期 + 表内顺序一起排。
    """
    return max(enumerate(rows), key=lambda t: (t[1]["day"], t[0]))[1]


def by_lot(rows):
    """按批号归堆：{ (批号, 品种): { 检项: [结果行, ...] } }，保持表内顺序。

    归堆要在判定之前做完。逐行判完再合并，复检那条就会变成一个独立结果，
    批次结论被初检那行盖住——判定取哪一次是本文第一条口径，不能靠运气。
    """
    lots = {}
    for r in rows:
        lots.setdefault((r["lot"], r["product"]), {}).setdefault(r["item"], []).append(r)
    return lots


def batch_rows(lots, specs):
    """按批号出一行：检项数、合格 / 不合格 / 待检 / 待核、批次结论。

    结论的优先顺序是 不合格 > 待检 > 合格：有不合格项就是不合格；没有不合格
    但还有没检完的项，就不能写「合格」——那等于替检验员签字；待核项（清单外、
    缺规范）不进结论，它属于清单要修的问题，不是这一批的问题。
    """
    out = []
    for (lot, product), items in lots.items():
        counts = {OK: 0, NG: 0, PEND: 0, OTHER: 0}
        for item, rows in items.items():
            judge, _ = judge_one(last_one(rows), specs.get((product, item)))
            counts[judge] += 1
        verdict = NG if counts[NG] else (PEND if counts[PEND] else OK)
        out.append([lot, product, len(items), counts[OK], counts[NG],
                    counts[PEND], counts[OTHER], verdict])
    out.sort(key=lambda r: (r[1], r[0]))
    return out


def pass_rate_rows(lots, specs):
    """按品种算合格率：分母只算「有规范、已出结果」的项。

    分母里混进没检的项和清单外的项，合格率就成了一个谁也解释不清的数：
    少检一项分母变小看着变差、多填一条清单外项分母又变大看着变好。
    分子分母都是判定过的项，这个数才敢往质量会上拿。
    """
    stat = {}
    for (lot, product), items in lots.items():
        s = stat.setdefault(product, {OK: 0, NG: 0})
        for item, rows in items.items():
            judge, _ = judge_one(last_one(rows), specs.get((product, item)))
            if judge in (OK, NG):
                s[judge] += 1
    out = []
    for product, s in stat.items():
        n = s[OK] + s[NG]
        out.append([product, n, s[OK], s[NG], round(s[OK] / n, 3) if n else None])
    out.sort(key=lambda r: r[0])
    return out


def defect_rows(lots, specs):
    """不合格明细：最终判定为不合格的项，逐条列出结果与判定依据。

    明细必须带「依据」一列。只给一个「不合格」，质量会上第一个问题就是
    「超的哪个限度、超了多少」，报表里没有就得回去翻原始数据。
    """
    out = []
    for (lot, product), items in lots.items():
        for item, rows in items.items():
            last = last_one(rows)
            judge, why = judge_one(last, specs.get((product, item)))
            if judge == NG:
                out.append([lot, product, item, last["value"], last["unit"],
                            last["day"], why])
    out.sort(key=lambda r: (r[1], r[0], r[2]))
    return out


def retest_rows(lots, specs):
    """复检留痕：同一条检项写了两行以上的，把初检和最终值并排列出。

    最终判定以最后一次为准，但初检那次必须留痕——质量上要能回答
    「为什么会有第二次检验」。静默丢掉，事后谁也说不清。
    """
    out = []
    for (lot, product), items in lots.items():
        for item, rows in items.items():
            if len(rows) < 2:
                continue
            ordered = sorted(rows, key=lambda r: (r["day"], rows.index(r)))
            first, last = ordered[0], last_one(rows)
            judge, why = judge_one(last, specs.get((product, item)))
            out.append([lot, product, item, f"{first['value']}（{first['day']}）",
                        f"{last['value']}（{last['day']}）", judge, why])
    out.sort(key=lambda r: (r[1], r[0], r[2]))
    return out


def completion_rows(lots, specs):
    """按品种 × 检项算完成情况：这个品种本月几批，其中几批出了结果。

    完成情况要按「品种 × 检项」看，不能只报一个总完成率：总完成率 90% 看着
    挺好，底下可能是某一项全批没检——那才是要追的。
    """
    lots_of = {}
    for lot, product in lots:
        lots_of.setdefault(product, set()).add(lot)
    out = []
    for product, lot_set in lots_of.items():
        for item in sorted(i for p, i in specs if p == product):
            done = sum(1 for lot in lot_set
                       if lots.get((lot, product), {}).get(item)
                       and last_one(lots[(lot, product)][item])["value"])
            out.append([product, item, len(lot_set), done, round(done / len(lot_set), 3)])
    out.sort(key=lambda r: (r[0], r[1]))
    return out


def todo_rows(lots, specs, known, other_month):
    """该有人核一遍的，全汇成一张清单：待检、清单外检项、该品种缺规范、不在本月。

    这几类都不能静默丢掉。待检要有人催、清单外检项要有人决定加不加进清单、
    缺规范要找质量部补，不在本月的行更要写清楚——不然月底对不上数，
    谁都说不清是漏统计还是真没做。
    """
    out = []
    for (lot, product), items in lots.items():
        for item, rows in items.items():
            judge, why = judge_one(last_one(rows), specs.get((product, item)))
            if judge == PEND:
                out.append(["待检", lot, f"{product} 的 {item} 还没出结果", item])
            elif judge == OTHER:
                kind = "清单外检项" if item not in known else "该品种缺规范"
                out.append([kind, lot, f"{product} 的 {item}：{why}", item])
    miss = {}
    for r in other_month:
        miss[(r["lot"], r["product"])] = miss.get((r["lot"], r["product"]), 0) + 1
    for (lot, product), n in sorted(miss.items()):
        out.append(["不在本月", lot,
                    f"{product} 有 {n} 行检验日期不在 {MONTH}，未计入本月报表", EMPTY])
    return out


def _sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([EMPTY if v is None else v for v in r])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"


def write_report(batch, defect, retest, rate, done_stat, todo):
    """一份月报六个页签：先看汇总，出问题再往下翻到明细和清单。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "批次汇总"
    _sheet(ws, ["批号", "品种", "检项数", "合格", "不合格", "待检", "待核", "结论"],
           batch, [12, 18, 8, 7, 9, 7, 7, 9])
    _sheet(wb.create_sheet("不合格明细"),
           ["批号", "品种", "检项", "结果", "单位", "检验日期", "判定依据"],
           defect, [12, 18, 12, 10, 7, 12, 18])
    _sheet(wb.create_sheet("复检留痕"),
           ["批号", "品种", "检项", "初检", "最终值", "最终判定", "判定依据"],
           retest, [12, 18, 12, 20, 20, 10, 18])
    _sheet(wb.create_sheet("品种合格率"),
           ["品种", "判定项数", "合格项", "不合格项", "合格率"], rate, [18, 10, 9, 10, 9])
    _sheet(wb.create_sheet("检项完成情况"),
           ["品种", "检项", "本批次数", "已出结果", "完成率"], done_stat, [18, 12, 10, 10, 9])
    _sheet(wb.create_sheet("待核清单"), ["类型", "批号", "说明", "关联检项"],
           todo, [12, 12, 52, 12])
    wb.save(OUT_DIR / REPORT_NAME)


def verify(batch, defect, retest, rate, done_stat, todo):
    """生成之后重新打开成品回读——内存里对不算对。"""
    oks = []
    wb = load_workbook(OUT_DIR / REPORT_NAME)
    oks.append(("六个页签都在", wb.sheetnames == ["批次汇总", "不合格明细", "复检留痕",
                                                 "品种合格率", "检项完成情况", "待核清单"]))
    ws = wb["批次汇总"]
    rows = {str(ws.cell(row=i + 2, column=1).value):
            [ws.cell(row=i + 2, column=c).value for c in range(1, 9)]
            for i in range(ws.max_row - 1)}
    oks.append(("本月批次 6 个，上月最后一天的批号不进汇总",
                len(rows) == 6 and "AML260630" not in rows))
    oks.append(("前导零批号回读仍是文本「007」", rows["007"][0] == "007"))
    oks.append(("复检按最后一次判定：AML260702 含量测定最终合格、不合格项 0",
                rows["AML260702"][3] == 3 and rows["AML260702"][4] == 0
                and rows["AML260702"][7] == "合格"))
    oks.append(("等于下限也判合格（MET260703 含量 95.0、007 释放度 70.0）",
                rows["MET260703"][4] == 1 and rows["007"][3] == 1))
    oks.append(("空结果算待检不算不合格（MET260703 溶出度、2607011 释放度）",
                rows["MET260703"][5] == 1 and rows["2607011"][5] == 1))
    oks.append(("清单外检项只进待核、不改批次结论（MET260704 检项数 5、待核 1、结论合格）",
                rows["MET260704"][2] == 5 and rows["MET260704"][6] == 1
                and rows["MET260704"][7] == "合格"))

    ws2 = wb["不合格明细"]
    oks.append(("不合格明细 2 条，且不含复检后合格的 AML260702 含量测定",
                ws2.max_row - 1 == len(defect) == 2
                and all(str(ws2.cell(row=i + 2, column=1).value) != "AML260702"
                        for i in range(ws2.max_row - 1))))
    ws3 = wb["复检留痕"]
    oks.append(("复检留痕 1 条，初检值 88.4 在里面",
                ws3.max_row - 1 == len(retest) == 1
                and "88.4" in str(ws3.cell(row=2, column=4).value)))
    ws4 = wb["品种合格率"]
    rates = {str(ws4.cell(row=i + 2, column=1).value): ws4.cell(row=i + 2, column=5).value
             for i in range(ws4.max_row - 1)}
    oks.append(("合格率分母不含待检与待核项（1.0 / 0.857 / 0.667）",
                abs(rates["阿莫西林胶囊"] - 1.0) < 1e-9
                and abs(rates["盐酸二甲双胍片"] - round(6 / 7, 3)) < 1e-9
                and abs(rates["布洛芬缓释胶囊"] - round(2 / 3, 3)) < 1e-9))
    ws5 = wb["检项完成情况"]
    oks.append(("检项完成情况 9 行（品种 × 检项）", ws5.max_row - 1 == len(done_stat) == 9))
    ws6 = wb["待核清单"]
    kinds = {}
    for i in range(ws6.max_row - 1):
        k = str(ws6.cell(row=i + 2, column=1).value)
        kinds[k] = kinds.get(k, 0) + 1
    oks.append(("待核清单 5 条（待检 2 / 清单外 1 / 缺规范 1 / 不在本月 1）",
                ws6.max_row - 1 == len(todo) == 5
                and kinds == {"待检": 2, "清单外检项": 1, "该品种缺规范": 1, "不在本月": 1}))
    return oks


def host_ip():
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))                    # 不真发包，只为拿到本机对外网卡的地址
        return s.getsockname()[0]
    except OSError:
        try:
            return socket.gethostbyname(socket.gethostname())
        except OSError:
            return "未知"
    finally:
        s.close()


def log(line):
    with open(LOG_FILE, "a", encoding="utf-8") as f:   # 追加不覆盖，每次跑留一段
        f.write(line + "\n")


def main():
    t0 = time.time()
    log(f"运行主机：{socket.gethostname()}（{host_ip()}）")
    log(f"开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}  Python {sys.version.split()[0]} / "
        f"{Path(sys.executable).name}")
    OUT_DIR.mkdir(exist_ok=True)

    specs, known = load_specs()
    results, other_month = load_results()
    lots = by_lot(results)
    batch = batch_rows(lots, specs)
    defect = defect_rows(lots, specs)
    retest = retest_rows(lots, specs)
    rate = pass_rate_rows(lots, specs)
    done_stat = completion_rows(lots, specs)
    todo = todo_rows(lots, specs, known, other_month)
    write_report(batch, defect, retest, rate, done_stat, todo)

    oks = verify(batch, defect, retest, rate, done_stat, todo)
    for name, ok in oks:
        print(("OK   " if ok else "FAIL ") + name)
    all_ok = all(ok for _, ok in oks)
    print(f"\n本月批次 {len(batch)} 个 / 判定项 {sum(r[1] for r in rate)} 条 / "
          f"不合格 {len(defect)} 条 / 复检留痕 {len(retest)} 条 / 待核 {len(todo)} 条")
    print("总判定：" + ("PASS" if all_ok else "FAIL"))
    log(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}  用时 {time.time() - t0:.1f}s  "
        f"判定 {'PASS' if all_ok else 'FAIL'}")

    if "--no-pause" not in sys.argv and sys.stdin and sys.stdin.isatty():
        input("按回车键结束...")


if __name__ == "__main__":
    main()

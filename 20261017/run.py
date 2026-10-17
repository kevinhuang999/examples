"""稳定性考察数据自动汇总：仪器导出的结果长表 → 汇总长表 / 宽表 / 异常清单。

用法：python run.py            # 跑完停在窗口里
      python run.py --no-pause # 跑完直接退出（脚本调用）
"""

import csv
import socket
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
RAW_DIR = BASE_DIR / "01_raw_data"
OUT_DIR = BASE_DIR / "02_output"
LOG_PATH = BASE_DIR / "source" / "run_log.txt"

PLAN_FILE = RAW_DIR / "稳定性考察方案.xlsx"
SPEC_FILE = RAW_DIR / "质量标准.xlsx"
RESULT_FILE = RAW_DIR / "稳定性考察结果_仪器导出.csv"

OUT_LONG = OUT_DIR / "稳定性考察汇总_长表.xlsx"
OUT_WIDE = OUT_DIR / "稳定性考察汇总_宽表.xlsx"
OUT_ALERT = OUT_DIR / "异常清单.xlsx"

# 仪器导出的列名，按名字取、不按列序取（列序会随仪器软件版本变）
COL_BATCH = "批号"
COL_COND = "考察条件"
COL_ITEM = "检项"
COL_TP = "计划时间点"
COL_DATE = "取样日期"
COL_VALUE = "结果"
COL_UNIT = "单位"
COL_N = "检测次数"

# 判定口径，以后要改就改这三行
WINDOW_DAYS = 15                                  # 取样窗口：计划月龄 ±15 天
REBOUND_LIMIT = 0.5                               # 含量比上一时间点回升多少算可疑（百分点）
TREND_LIMIT = {"含量": -2.0, "有关物质总杂": 0.3}  # 相对 0 月的变化阈值
DAYS_PER_MONTH = 30.4375                          # 平均月长，只用来算月龄，不做合规依据


def ensure_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return
    except ImportError:
        print("缺少 openpyxl，正在用清华源安装……")
        subprocess.check_call([
            sys.executable, "-m", "pip", "install",
            "-i", "https://pypi.tuna.tsinghua.edu.cn/simple", "openpyxl",
        ])


ensure_openpyxl()
import openpyxl  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402


def host_label():
    """日志里记下是哪台机器跑的：主机名 + 本机对外网卡地址。"""
    name = socket.gethostname()
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sock.connect(("8.8.8.8", 80))
        ip = sock.getsockname()[0]
        sock.close()
    except OSError:
        try:
            ip = socket.gethostbyname(name)
        except OSError:
            ip = "未知"
    return f"{name}（{ip}）"


def write_log(text):
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(LOG_PATH, "a", encoding="utf-8") as f:
        f.write(text + "\n")


def to_number(text):
    """空值、斜杠、'ND' 一律当没结果，别让 float('') 把整批跑崩。"""
    s = ("" if text is None else str(text)).strip().replace(",", "")
    if s in ("", "/", "-", "ND", "未检出"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_timepoints(text):
    """方案表里的计划时间点是一串文本 '0,1,2,3,6'，拆成整数列表。"""
    return [int(part) for part in str(text).replace("，", ",").split(",") if part.strip()]


def read_plan(path):
    """方案表：一行一条考察线（批号 + 条件），带起始日和计划时间点。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    plan = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        start = row[2]
        if isinstance(start, datetime):
            start = start.date()
        else:
            start = datetime.strptime(str(start)[:10], "%Y-%m-%d").date()
        plan.append({
            "批号": str(row[0]).strip(),
            "考察条件": str(row[1]).strip(),
            "起始日": start,
            "时间点": parse_timepoints(row[3]),
        })
    return plan


def read_spec(path):
    """质量标准：上下限留空表示那一侧不限。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    spec = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        spec[str(row[0]).strip()] = {
            "单位": "" if row[1] is None else str(row[1]).strip(),
            "下限": to_number(row[2]),
            "上限": to_number(row[3]),
        }
    return spec


def read_results(path):
    """结果长表按列名取，返回 dict 列表——后面全按字段名访问，不认列序。"""
    with open(path, encoding="utf-8-sig", newline="") as f:
        return [{k: (v or "").strip() for k, v in raw.items()} for raw in csv.DictReader(f)]


def pick_latest(rows):
    """同一批号 + 条件 + 检项 + 时间点有复测时，只留检测次数最大的那条。"""
    picked, dup = {}, 0
    for row in rows:
        key = (row[COL_BATCH], row[COL_COND], row[COL_ITEM], row[COL_TP])
        seq = int(row[COL_N] or 1)
        if key not in picked:
            picked[key] = (seq, row)
            continue
        dup += 1
        if seq >= picked[key][0]:
            picked[key] = (seq, row)
    return [item[1] for item in picked.values()], dup


def parse_value(text):
    """'<0.05' 这类低于定量限的结果取下限当数值，并留一个标记。"""
    s = (text or "").strip()
    if s.startswith("<"):
        return to_number(s[1:]), "低于定量限"
    return to_number(s), ""


def judge(value, spec_row):
    """拿质量标准上下限判合格；任何一侧留空就不管那一侧。"""
    if value is None:
        return "待确认"
    low, high = spec_row["下限"], spec_row["上限"]
    if low is not None and value < low:
        return "不合格"
    if high is not None and value > high:
        return "不合格"
    return "合格"


def normalize(rows, plan, spec):
    """把仪器导出的原始行摊成一条条判定记录：算月龄、查窗口、判合格。"""
    plan_map = {(p["批号"], p["考察条件"]): p for p in plan}
    long_rows = []
    for row in rows:
        batch, cond = row[COL_BATCH], row[COL_COND]
        item, tp_text = row[COL_ITEM], row[COL_TP]
        month = int(tp_text.replace("月", "") or 0)
        value, flag = parse_value(row[COL_VALUE])
        plan_row = plan_map.get((batch, cond))
        if plan_row is None:
            judge_text = "待确认"
            marks = ["方案外"]
            actual = ""
        else:
            days = (datetime.strptime(row[COL_DATE][:10], "%Y-%m-%d").date() - plan_row["起始日"]).days
            actual = round(days / DAYS_PER_MONTH, 1)
            marks = [flag] if flag else []
            if abs(days - month * DAYS_PER_MONTH) > WINDOW_DAYS:
                marks.append("超窗")
            judge_text = judge(value, spec.get(item, {"下限": None, "上限": None}))
            if judge_text == "不合格":
                marks.append("超限")
        if not marks:
            marks.append("复测取末次" if int(row[COL_N] or 1) > 1 else "")
        long_rows.append({
            "批号": batch, "考察条件": cond, "检项": item,
            "计划时间点": tp_text, "计划月龄": month, "实际月龄": actual,
            "取样日期": row[COL_DATE], "结果原文": row[COL_VALUE],
            "结果数值": value, "单位": row[COL_UNIT] or spec.get(item, {}).get("单位", ""),
            "判定": judge_text, "标记": "、".join(m for m in marks if m),
        })
    return sorted(long_rows, key=lambda r: (r["批号"], r["考察条件"], r["检项"], r["计划月龄"]))


def trend_check(long_rows):
    """按考察线看趋势：相对 0 月变化超阈、含量比上一时间点回升。"""
    alerts = []
    groups = {}
    for row in long_rows:
        if row["判定"] == "待确认":
            continue
        groups.setdefault((row["批号"], row["考察条件"]), []).append(row)
    for (batch, cond), items in groups.items():
        for item, limit in TREND_LIMIT.items():
            series = sorted(
                [r for r in items if r["检项"] == item and r["结果数值"] is not None],
                key=lambda r: r["计划月龄"],
            )
            if len(series) < 2:
                continue
            base = series[0]["结果数值"]
            for row in series[1:]:
                change = round(row["结果数值"] - base, 3)
                hit = change <= limit if limit < 0 else change >= limit
                if hit:
                    alerts.append({
                        "类型": "变化超阈", "批号": batch, "考察条件": cond,
                        "检项": item, "计划时间点": row["计划时间点"],
                        "说明": f"相对 0 月变化 {change:+.2f}，阈值 {limit:+.1f}",
                    })
            if item != "含量":
                continue
            for prev, cur in zip(series, series[1:]):
                step = round(cur["结果数值"] - prev["结果数值"], 3)
                if step > REBOUND_LIMIT:
                    alerts.append({
                        "类型": "回跳可疑", "批号": batch, "考察条件": cond,
                        "检项": item, "计划时间点": cur["计划时间点"],
                        "说明": f"较上一时间点回升 {step:+.1f} 个百分点，先查取样与复测记录",
                    })
    return alerts


def build_alerts(long_rows, trend_alerts):
    """异常清单 = 缺结果 / 方案外 / 超窗 / 不合格 + 趋势类。"""
    alerts = []
    for row in long_rows:
        if row["结果数值"] is None:
            alerts.append({
                "类型": "缺结果", "批号": row["批号"], "考察条件": row["考察条件"],
                "检项": row["检项"], "计划时间点": row["计划时间点"],
                "说明": "结果为空，等仪器出数或确认是否漏做",
            })
        elif "方案外" in row["标记"]:
            alerts.append({
                "类型": "方案外", "批号": row["批号"], "考察条件": row["考察条件"],
                "检项": row["检项"], "计划时间点": row["计划时间点"],
                "说明": "方案表里查不到这条考察线，先确认批号与条件写法",
            })
        if "超窗" in row["标记"]:
            alerts.append({
                "类型": "超窗", "批号": row["批号"], "考察条件": row["考察条件"],
                "检项": row["检项"], "计划时间点": row["计划时间点"],
                "说明": f"实际月龄 {row['实际月龄']}，超出计划点 ±{WINDOW_DAYS} 天",
            })
        if row["判定"] == "不合格":
            alerts.append({
                "类型": "不合格", "批号": row["批号"], "考察条件": row["考察条件"],
                "检项": row["检项"], "计划时间点": row["计划时间点"],
                "说明": f"结果 {row['结果原文']} 超出质量标准",
            })
    return alerts + trend_alerts


def pivot_wide(long_rows, plan, spec):
    """宽表：一行一条考察线，一列一个时间点，一个检项一张表。"""
    months = sorted({tp for p in plan for tp in p["时间点"]})
    index = {(r["批号"], r["考察条件"], r["检项"], r["计划月龄"]): r["结果数值"] for r in long_rows}
    tables = {}
    for item in spec:
        table = []
        for p in plan:
            line = {"批号": p["批号"], "考察条件": p["考察条件"]}
            for month in months:
                line[f"{month}月"] = index.get((p["批号"], p["考察条件"], item, month))
            table.append(line)
        tables[item] = (months, table)
    return tables


HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
FAIL_FONT = Font(color="C00000", bold=True)
PENDING_FONT = Font(color="BF8F00")


def write_sheet(ws, header, rows, widths):
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(k) for k in header])
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=ws.max_row, column=len(header)).coordinate}"


def write_outputs(long_rows, alerts, wide_tables):
    """三张表一次落盘：汇总长表、宽表、异常清单。"""
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    long_header = ["批号", "考察条件", "检项", "计划时间点", "计划月龄", "实际月龄",
                   "取样日期", "结果原文", "结果数值", "单位", "判定", "标记"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "汇总"
    write_sheet(ws, long_header, long_rows, [8, 20, 14, 12, 10, 10, 12, 10, 10, 8, 10, 16])
    for idx, row in enumerate(long_rows, start=2):
        if row["判定"] == "不合格":
            for cell in ws[idx]:
                cell.font = FAIL_FONT
        elif row["判定"] == "待确认":
            for cell in ws[idx]:
                cell.font = PENDING_FONT
    wb.save(OUT_LONG)

    wb_wide = openpyxl.Workbook()
    wb_wide.remove(wb_wide.active)
    for item, (months, table) in wide_tables.items():
        header = ["批号", "考察条件"] + [f"{m}月" for m in months]
        ws_wide = wb_wide.create_sheet(item[:28])
        write_sheet(ws_wide, header, table, [8, 20] + [10] * len(months))
        for row in ws_wide.iter_rows(min_row=2):
            for cell in row[2:]:
                cell.number_format = "0.00"
    wb_wide.save(OUT_WIDE)

    alert_header = ["类型", "批号", "考察条件", "检项", "计划时间点", "说明"]
    wb_alert = openpyxl.Workbook()
    ws_alert = wb_alert.active
    ws_alert.title = "异常清单"
    write_sheet(ws_alert, alert_header, alerts, [12, 8, 20, 14, 12, 60])
    wb_alert.save(OUT_ALERT)


def verify(long_rows, alerts, dup):
    """生成完回读三张表，逐项断言——不靠肉眼。"""
    checks = []

    def check(name, ok):
        checks.append((name, bool(ok)))

    wb = openpyxl.load_workbook(OUT_LONG, data_only=True)
    ws = wb["汇总"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]

    def count(keyword):
        return sum(1 for r in rows if keyword in str(r["标记"]))

    check("汇总长表 27 行（28 条原始结果合并 1 组复测）", len(rows) == 27)
    check("复测合并 1 组，且没丢行", dup == 1 and len(rows) + dup == 28)
    check("缺结果 1 条", sum(1 for r in rows if r["结果数值"] is None) == 1)
    check("方案外 1 条", count("方案外") == 1)
    check("超窗 1 条", count("超窗") == 1)
    check("不合格 1 条", sum(1 for r in rows if r["判定"] == "不合格") == 1)
    check("低于定量限 1 条", count("低于定量限") == 1)
    check("待确认 2 条（缺结果 + 方案外）", sum(1 for r in rows if r["判定"] == "待确认") == 2)

    target = [r for r in rows if r["批号"] == "0202" and r["考察条件"].startswith("加速")
              and r["检项"] == "含量" and r["计划月龄"] == 6][0]
    check("复测取末次值 100.1 且判合格", target["结果数值"] == 100.1 and target["判定"] == "合格")

    kinds = [a["类型"] for a in alerts]
    check("异常清单 8 条", len(alerts) == 8)
    check("回跳可疑 1 条", kinds.count("回跳可疑") == 1)
    check("变化超阈 3 条", kinds.count("变化超阈") == 3)

    wb_wide = openpyxl.load_workbook(OUT_WIDE, data_only=True)
    ws_wide = wb_wide["含量"]
    wide_header = [c.value for c in ws_wide[1]]
    wide_rows = [dict(zip(wide_header, [c.value for c in r])) for r in ws_wide.iter_rows(min_row=2)]
    hit = [r for r in wide_rows if r["批号"] == "0202" and r["考察条件"].startswith("加速")][0]
    check("宽表含量 8 列（批号 + 条件 + 6 个时间点）", len(wide_header) == 8)
    check("宽表 0202 加速 6 月取到复测值", hit["6月"] == 100.1)

    for name, ok in checks:
        print(("OK   " if ok else "FAIL ") + name)
    return all(ok for _, ok in checks)


def main():
    started = time.time()
    write_log(f"运行主机：{host_label()}")
    write_log(f"开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}  "
              f"Python {sys.version.split()[0]}（{Path(sys.executable).name}）")

    plan = read_plan(PLAN_FILE)
    spec = read_spec(SPEC_FILE)
    raw = read_results(RESULT_FILE)
    print(f"读到：考察线 {len(plan)} 条、质量标准 {len(spec)} 项、原始结果 {len(raw)} 行")

    rows, dup = pick_latest(raw)
    long_rows = normalize(rows, plan, spec)
    alerts = build_alerts(long_rows, trend_check(long_rows))
    write_outputs(long_rows, alerts, pivot_wide(long_rows, plan, spec))

    print(f"合并复测 {dup} 组，汇总 {len(long_rows)} 行；异常 {len(alerts)} 条：")
    for item in ("缺结果", "方案外", "超窗", "不合格", "变化超阈", "回跳可疑"):
        n = sum(1 for a in alerts if a["类型"] == item)
        print(f"  {item}：{n}")
    print(f"输出目录：{OUT_DIR}")

    passed = verify(long_rows, alerts, dup)
    cost = time.time() - started
    print(f"判定：{'PASS' if passed else 'FAIL'}（用时 {cost:.1f}s）")
    write_log(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}  用时 {cost:.1f}s  "
              f"判定 {'PASS' if passed else 'FAIL'}")

    if "--no-pause" not in sys.argv:
        input("按回车结束……")


if __name__ == "__main__":
    main()

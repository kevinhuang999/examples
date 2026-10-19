# -*- coding: utf-8 -*-
"""SOP 受控文档批量渲染：一份模板 + 一张台账 -> 一批带受控页眉的 docx。

跑法：python run.py            （Windows 也可以直接双击）
      python run.py --no-pause
"""

import re
import socket
import subprocess
import sys
from datetime import datetime
from pathlib import Path

BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / "source" / "templates" / "SOP模板.docx"
LEDGER = BASE / "01_raw_data" / "SOP台账.xlsx"
OUT_DIR = BASE / "02_output"
LOG = BASE / "source" / "run_log.txt"

# 台账列名 ←→ 代码里的短键；客户表格改表头，只改这一张映射
COL = {
    "code": "文件编号",
    "name": "文件名称",
    "ver": "版本",
    "status": "状态",
    "eff": "生效日期",
    "owner": "修订人",
    "note": "本次修订内容",
}
CURRENT = "现行"


def ensure_deps():
    need = []
    for mod, pkg in (("docx", "python-docx"), ("openpyxl", "openpyxl")):
        try:
            __import__(mod)
        except ImportError:
            need.append(pkg)
    if not need:
        return
    for extra in ([], ["--user"]):
        cmd = [sys.executable, "-m", "pip", "install",
               "-i", "https://pypi.tuna.tsinghua.edu.cn/simple"] + need + extra
        if subprocess.call(cmd) == 0:
            return
        raise SystemExit("依赖安装失败，请手动执行：pip install " + " ".join(need))


def host_label():
    name = socket.gethostname()
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sock.connect(("8.8.8.8", 80))
        ip = sock.getsockname()[0]
        sock.close()
    except OSError:
        try:
            ip = socket.gethostbyname(name)
        except OSError:
            ip = "未知"
    return f"{name}（{ip}）"


def log_head():
    LOG.parent.mkdir(parents=True, exist_ok=True)
    with LOG.open("a", encoding="utf-8") as fh:
        fh.write(f"运行主机：{host_label()}\n")
        fh.write(f"开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}\n")
        fh.write(f"Python：{sys.version.split()[0]}（{Path(sys.executable).name}）\n")


def read_ledger():
    from openpyxl import load_workbook
    wb = load_workbook(LEDGER, data_only=True)
    ws = wb.active
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, raw))
        if not record.get(COL["code"]):
            continue
        rows.append({key: "" if record.get(col) is None else str(record.get(col)).strip()
                     for key, col in COL.items()})
    return rows


def ver_key(ver):
    digits = re.sub(r"\D", "", ver)
    return int(digits) if digits else 0


def pick_current(rows):
    buckets = {}
    for row in rows:
        if row["status"] == CURRENT:
            buckets.setdefault(row["code"], []).append(row)
    current = {code: max(group, key=lambda r: (ver_key(r["ver"]), r["eff"]))
               for code, group in buckets.items()}
    skipped = sorted({row["code"] for row in rows} - set(current))
    return current, skipped


def build_ref_map(current):
    return {code: f"《{item['name']}》（编号 {code}，现行版本 {item['ver']}）"
            for code, item in current.items()}


def build_fields(item, ref_map):
    fields = {
        "DOC_NO": item["code"],
        "TITLE": item["name"],
        "VER": item["ver"],
        "EFF": item["eff"],
        "OWNER": item["owner"],
    }
    for code, text in ref_map.items():
        fields["SOP:" + code] = text
    return fields


def replace_in_paragraph(par, fields, pending):
    text = par.text
    if "{{" not in text:
        return
    new = text
    for key, value in fields.items():
        new = new.replace("{{" + key + "}}", value)
    for code in re.findall(r"\{\{SOP:([^}]+)\}\}", new):
        new = new.replace("{{SOP:" + code + "}}", f"[待确认:{code}]")
        if code not in pending:
            pending.append(code)
    if new != text:
        for run in par.runs[1:]:
            run.text = ""
        par.runs[0].text = new


def iter_paragraphs(doc):
    for par in doc.paragraphs:
        yield par
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    yield par
    for section in doc.sections:
        for part in (section.header, section.footer):
            for par in part.paragraphs:
                yield par


def append_revision(doc, item):
    table = doc.tables[0]
    row = table.add_row()
    for cell, value in zip(row.cells, [item["ver"], item["eff"], item["note"], item["owner"]]):
        cell.text = value
    return len(table.rows)


def render_one(item, fields, pending):
    from docx import Document
    doc = Document(str(TEMPLATE))
    for par in iter_paragraphs(doc):
        replace_in_paragraph(par, fields, pending)
    rows_before = len(doc.tables[0].rows)
    rows_after = append_revision(doc, item)
    out_path = OUT_DIR / f"{item['code']}_{item['ver']}_{item['name']}.docx"
    doc.save(str(out_path))
    return out_path, rows_before, rows_after


def verify(rows, made, current, skipped, pending):
    from docx import Document
    checks = []

    def check(name, ok):
        checks.append((name, bool(ok)))

    names = " ".join(path.name for path, _, _ in made)
    check("台账读到 10 行", len(rows) == 10)
    check("现行版 6 份（10 行里含 3 个历史版本）", len(current) == 6)
    check("生成了 6 个文件", len(made) == 6)
    check("全部作废的编号被跳过", skipped == ["SOP-QA-011"])
    check("V9 与 V10 并存时取 V10（版本按数字比）",
          "SOP-QC-007_V10" in names and "SOP-QC-007_V9_" not in names)
    check("台账里没有的引用标成待确认", pending == ["SOP-QA-030"])

    header_ok = footer_ok = clean_ok = ref_ok = table_ok = True
    for path, before, after in made:
        doc = Document(str(path))
        body = "\n".join(p.text for p in doc.paragraphs)
        head = "\n".join(p.text for s in doc.sections for p in s.header.paragraphs)
        foot = "\n".join(p.text for s in doc.sections for p in s.footer.paragraphs)
        code, ver = path.name.split("_")[0], path.name.split("_")[1]
        header_ok = header_ok and code in head and ver in head and "{{" not in head
        footer_ok = footer_ok and code in foot
        clean_ok = clean_ok and "{{" not in body and "{{" not in foot
        ref_ok = ref_ok and "《高效液相色谱仪操作标准操作规程》（编号 SOP-QC-003，现行版本 V03）" in body
        table_ok = table_ok and after == before + 1 and doc.tables[0].rows[-1].cells[0].text == ver

    check("每份页眉都换成本份的编号与版本", header_ok)
    check("每份页脚都换成本份的编号（页脚是 section 级）", footer_ok)
    check("正文与页脚无残留 {{ 占位符", clean_ok)
    check("正文交叉引用已展开成名称+编号+现行版本", ref_ok)
    check("修订历史表各追加 1 行且末行是本份版本", table_ok)

    for name, ok in checks:
        print(("OK   " if ok else "FAIL ") + name)
    return all(ok for _, ok in checks)


def main():
    ensure_deps()
    start = datetime.now()
    OUT_DIR.mkdir(exist_ok=True)
    log_head()

    rows = read_ledger()
    current, skipped = pick_current(rows)
    ref_map = build_ref_map(current)
    pending = []
    made = []
    for code in sorted(current):
        item = current[code]
        made.append(render_one(item, build_fields(item, ref_map), pending))

    passed = verify(rows, made, current, skipped, pending)
    cost = (datetime.now() - start).total_seconds()
    print(f"\n生成 {len(made)} 份，跳过编号 {len(skipped)} 个，待确认引用 {len(pending)} 处")
    print("结果目录：02_output")
    print("判定：" + ("PASS" if passed else "FAIL"))
    with LOG.open("a", encoding="utf-8") as fh:
        fh.write(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}  用时 {cost:.1f}s  "
                 f"判定 {'PASS' if passed else 'FAIL'}\n")
    return 0 if passed else 1


if __name__ == "__main__":
    code = main()
    if "--no-pause" not in sys.argv:
        input("\n按回车键退出……")
    sys.exit(code)

# -*- coding: utf-8 -*-
"""重造示例数据。

上半段：建受控文档模板（source/templates/SOP模板.docx）
下半段：造台账（01_raw_data/SOP台账.xlsx），顺手清掉上一次的输出与日志。
"""

import shutil
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
TPL_DIR = BASE / "source" / "templates"
TPL = TPL_DIR / "SOP模板.docx"
LEDGER = BASE / "01_raw_data" / "SOP台账.xlsx"
OUT_DIR = BASE / "02_output"
LOG = BASE / "source" / "run_log.txt"

HEADER = "文件编号：{{DOC_NO}}｜版本：{{VER}}｜生效日期：{{EFF}}"
FOOTER = "{{DOC_NO}}｜受控文件，打印后请核对页眉版本"
BODY = [
    "{{TITLE}}",
    "1. 目的",
    "本标准操作规程规定 {{TITLE}} 的执行要求，适用于本部门及相关岗位。",
    "2. 引用文件",
    "本规程引用 {{SOP:SOP-QC-003}} 与 {{SOP:SOP-QA-009}}。",
    "3. 程序",
    "按 {{SOP:SOP-QC-003}} 第 5.2 条执行；与 {{SOP:SOP-QA-030}} 有冲突时，以本规程为准。",
]
HISTORY_HEAD = ["版本", "生效日期", "修订内容", "修订人"]
HISTORY_LAST = ["V02", "2024-06-15", "按 2023 版药典修订术语", "李敏"]

# 故意造的边界：同编号多版本 / V9 与 V10 并存 / 全部作废 / 前导零 / 只有一条
LEDGER_ROWS = [
    ("SOP-QC-003", "高效液相色谱仪操作标准操作规程", "V01", "作废", "2023-04-01", "李敏", "首次发布"),
    ("SOP-QC-003", "高效液相色谱仪操作标准操作规程", "V02", "作废", "2024-06-15", "李敏", "按 2023 版药典修订术语"),
    ("SOP-QC-003", "高效液相色谱仪操作标准操作规程", "V03", "现行", "2026-03-01", "李敏", "增加系统适用性试验判定"),
    ("SOP-QC-007", "纯化水取样与检验规程", "V9", "作废", "2025-01-10", "王强", "取样点编号调整"),
    ("SOP-QC-007", "纯化水取样与检验规程", "V10", "现行", "2026-05-06", "王强", "增加 TOC 检测频次"),
    ("SOP-QA-009", "偏差处理规程", "V02", "现行", "2025-11-20", "张蕾", "补充分级时限"),
    ("SOP-QA-011", "供应商审计管理规程", "V04", "作废", "2024-09-02", "张蕾", "并入集团统一规程"),
    ("SOP-QA-012", "变更控制管理规程", "V05", "现行", "2026-01-15", "张蕾", "增加紧急变更通道"),
    ("SOP-QC-021", "稳定性考察样品管理规程", "V01", "现行", "2026-07-01", "陈晨", "首次发布"),
    ("SOP-MFG-004", "洁净区清洁消毒规程", "V06", "现行", "2026-08-18", "陈晨", "更换消毒剂种类"),
]


def build_template():
    from docx import Document
    from docx.shared import Pt

    TPL_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    head_par = section.header.paragraphs[0]
    head_par.text = HEADER
    foot_par = section.footer.paragraphs[0]
    foot_par.text = FOOTER

    for i, text in enumerate(BODY):
        par = doc.add_paragraph(text)
        if i == 0:
            par.runs[0].bold = True
            par.runs[0].font.size = Pt(16)
        if text.startswith("3. 程序"):
            doc.add_paragraph("4. 修订历史")

    table = doc.add_table(rows=2, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, HISTORY_HEAD):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, HISTORY_LAST):
        cell.text = text

    props = doc.core_properties
    props.author = "示例"
    props.last_modified_by = "示例"
    props.title = "SOP 受控文档模板"
    doc.save(str(TPL))
    print(f"模板已重建：{TPL.relative_to(BASE)}")


def build_ledger():
    from openpyxl import Workbook

    LEDGER.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "SOP台账"
    ws.append(["文件编号", "文件名称", "版本", "状态", "生效日期", "修订人", "本次修订内容"])
    for row in LEDGER_ROWS:
        ws.append(list(row))
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["G"].width = 26
    wb.save(str(LEDGER))
    print(f"台账已重建：{LEDGER.relative_to(BASE)}（{len(LEDGER_ROWS)} 行）")


def reset_outputs():
    if OUT_DIR.exists():
        shutil.rmtree(OUT_DIR)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if LOG.exists():
        LOG.unlink()
    print("已清空 02_output 与 source/run_log.txt")


if __name__ == "__main__":
    build_template()
    build_ledger()
    reset_outputs()
    print("重造完成，接着跑：python run.py --no-pause")

# -*- coding: utf-8 -*-
"""实验室每日检测日报汇总：一批格式各异的片段 -> 一份带格式的日报。

跑法：python run.py            （Windows 也可以直接双击）
      python run.py --no-pause
"""

import csv
import re
import socket
import subprocess
import sys
from copy import copy
from datetime import date, datetime
from pathlib import Path

BASE = Path(__file__).resolve().parent
TEMPLATE = BASE / "source" / "templates" / "日报模板.xlsx"
RAW_DIR = BASE / "01_raw_data"
OUT_PATH = BASE / "02_output" / "检测日报.xlsx"
LOG = BASE / "source" / "run_log.txt"

# 要汇总哪一天：写成常量，别从「今天」推 —— 补跑昨天的日报时才不会全错
REPORT_DAY = date(2026, 10, 20)

# 表头不一定在第一行，靠这几个关键字找它在第几行
HEADER_KEYS = ("样品编号", "样品号", "Sample ID")

# 列名别名表：各人导出的表头五花八门，先归一成短键，后面只认短键
ALIAS = {
    "sid": ("样品编号", "样品号", "样品ID", "Sample ID"),
    "item": ("检测项目", "项目", "检验项目", "Item"),
    "result": ("结果", "检测结果", "检验结果", "Result"),
    "unit": ("单位", "Unit"),
    "time": ("检测时间", "完成时间", "检验时间", "TestTime"),
}
COL_ORDER = ("sid", "item", "result", "unit", "time")
EMPTY_RESULT = ("", "—", "-", "/", "未出", "待定")


def norm(value):
    """单元格归一：转字符串、去空格、转小写，专给比列名用。"""
    return re.sub(r"\s+", "", str(value if value is not None else "")).lower()


def ensure_deps():
    try:
        import openpyxl  # noqa: F401
        return
    except ImportError:
        pass
    for extra in ([], ["--user"]):
        cmd = [sys.executable, "-m", "pip", "install",
               "-i", "https://pypi.tuna.tsinghua.edu.cn/simple", "openpyxl"] + extra
        if subprocess.call(cmd) == 0:
            return
    raise SystemExit("依赖安装失败，请手动执行：pip install openpyxl")


def host_label():
    name = socket.gethostname()
    try:
        sock = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sock.connect(("8.8.8.8", 80))
        ip = sock.getsockname()[0]
        sock.close()
    except OSError:
        try:
            ip = socket.gethostbyname(name)
        except OSError:
            ip = "未知"
    return f"{name}（{ip}）"


def log_head():
    LOG.parent.mkdir(parents=True, exist_ok=True)
    with LOG.open("a", encoding="utf-8") as fh:
        fh.write(f"运行主机：{host_label()}\n")
        fh.write(f"开始执行：{datetime.now():%Y-%m-%d %H:%M:%S}\n")
        fh.write(f"Python：{sys.version.split()[0]}（{Path(sys.executable).name}）\n")


def find_header(rows):
    """表头在哪一行？返回 0 基行号，找不到返回 -1。"""
    keys = {norm(k) for k in HEADER_KEYS}
    for i, row in enumerate(rows):
        if keys & {norm(c) for c in row}:
            return i
    return -1


def map_columns(header):
    """把这一行的列名翻译成短键 -> 列下标。"""
    cells = [norm(c) for c in header]
    pos = {}
    for key, names in ALIAS.items():
        for name in names:
            if norm(name) in cells:
                pos[key] = cells.index(norm(name))
                break
    return pos


def read_fragment(path):
    """读一个片段，返回 (表头行号从 1 数, 记录列表)。"""
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            rows = [list(r) for r in csv.reader(fh)]
    else:
        from openpyxl import load_workbook
        sheet = load_workbook(path, data_only=True).active
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    head = find_header(rows)
    if head < 0:
        raise SystemExit(f"{path.name} 里找不到表头，估计列名被改过")
    pos = map_columns(rows[head])
    missing = [k for k in COL_ORDER if k not in pos]
    if missing:
        raise SystemExit(f"{path.name} 缺列：{missing}")
    records = []
    for raw in rows[head + 1:]:
        if not any(norm(c) for c in raw):
            continue
        rec = {"组别": path.stem.split("_")[0], "来源文件": path.name}
        for key in COL_ORDER:
            rec[key] = raw[pos[key]] if pos[key] < len(raw) else ""
        records.append(rec)
    return head + 1, records


def parse_time(value):
    """是 datetime 就直接用；是字符串就试几种常见写法；都不中返回 None。"""
    if isinstance(value, datetime):
        return value
    text = str(value).strip().replace("/", "-")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def pick_latest(records):
    """同一样品同一项目出过多次结果，只留检测时间最晚的那条（复测取末次）。"""
    keep = {}
    for rec in records:
        key = (norm(rec["sid"]), norm(rec["item"]))
        stamp = parse_time(rec["time"]) or datetime.min
        if key not in keep or stamp >= keep[key][0]:
            keep[key] = (stamp, rec)
    return [rec for _, rec in keep.values()]


def split_by_day(records, day):
    """按「检测时间」的日期分堆 —— 不按文件名，补交昨天的数据不能算成今天的。"""
    today, other = [], []
    for rec in records:
        stamp = parse_time(rec["time"])
        (today if stamp and stamp.date() == day else other).append(rec)
    return today, other


def is_pending(rec):
    """结果空着，或者写着破折号、待定，都算当天没出结果。"""
    return norm(rec["result"]) in {norm(x) for x in EMPTY_RESULT}


def write_rows(sheet, rows, start_row, style_row):
    """从 start_row 起写数据，格式一律照抄 style_row 那行 —— 模板的框线和数字格式就不会丢。"""
    width = sheet.max_column
    source = [sheet.cell(row=style_row, column=c) for c in range(1, width + 1)]
    for i, values in enumerate(rows):
        for c, value in enumerate(values, start=1):
            cell = sheet.cell(row=start_row + i, column=c)
            cell.value = value
            cell.font = copy(source[c - 1].font)
            cell.border = copy(source[c - 1].border)
            cell.alignment = copy(source[c - 1].alignment)
            cell.number_format = source[c - 1].number_format
    return start_row + len(rows)


def build_report(done, pending, other):
    from openpyxl import load_workbook
    book = load_workbook(TEMPLATE)

    summary = book["汇总"]
    summary["B2"] = REPORT_DAY
    summary["B2"].number_format = "yyyy-mm-dd"
    rows = [[rec["sid"], rec["item"], rec["result"], rec["unit"],
             parse_time(rec["time"]), rec["组别"]] for rec in done]
    write_rows(summary, rows, 4, 4)

    late = book["待出结果"]
    rows = [[rec["sid"], rec["item"], rec["组别"],
             parse_time(rec["time"]), "当天没出结果，次日接着追"] for rec in pending]
    write_rows(late, rows, 3, 3)

    stats = book["统计"]
    buckets = {}
    for index, group in enumerate((done, pending, other)):
        for rec in group:
            buckets.setdefault(rec["组别"], [0, 0, 0])[index] += 1
    rows = [[name] + counts + [sum(counts)] for name, counts in sorted(buckets.items())]
    write_rows(stats, rows, 3, 3)

    OUT_PATH.parent.mkdir(exist_ok=True)
    book.save(OUT_PATH)


def verify(raw, merged, done, pending, other, heads):
    from openpyxl import load_workbook
    checks = []

    def check(name, ok):
        checks.append((name, bool(ok)))

    book = load_workbook(OUT_PATH)
    summary = book["汇总"]
    late = book["待出结果"]
    stats = book["统计"]

    names = [summary.cell(row=r, column=1).value for r in range(4, 4 + len(done))]
    result = {summary.cell(row=r, column=1).value: summary.cell(row=r, column=3).value
              for r in range(4, 4 + len(done))}
    rows = [r for r in range(3, 3 + len({rec["组别"] for rec in merged}))]
    total = sum(stats.cell(row=r, column=5).value or 0 for r in rows)

    check("5 个片段全部读到（16 条原始记录）", len(raw) == 16)
    check("表头压在第 2 行、第 3 行的都认出来了",
          dict(heads)["理化组_20261020.xlsx"] == 2
          and dict(heads)["前处理组_20261020.xlsx"] == 3
          and dict(heads)["微生物组_20261020.xlsx"] == 1)
    check("英文字段名的 csv 也读得动（仪器组 3 条）",
          len([r for r in raw if r["组别"] == "仪器组"]) == 3)
    check("两组复测去重后剩 14 条", len(merged) == 14)
    check("复测取末次：S-0001 的 pH 是 6.85 不是 6.82", result.get("S-0001") == 6.85)
    check("跨文件复测也取末次：S-0002 水分 3.17 不是 3.15", result.get("S-0002") == 3.17)
    check("昨天测的 2 条没混进今日日报", len(other) == 2)
    check("当天已出结果 9 条", len(done) == 9)
    check("结果空着或写破折号的 3 条进待出页签", len(pending) == 3)
    check("待出页签写了 3 行就收住", late.cell(row=5, column=1).value is not None
          and late.cell(row=6, column=1).value is None)
    check("“<0.05” 原样保留，没被当数字吃掉", result.get("S-0006") == "<0.05")
    check("样品编号前导零没丢（S-0006 不是 6）", "S-0006" in names)
    check("统计表合计 14 = 已出 9 + 待出 3 + 补交 2", total == 14)
    check("模板格式还在：标题加粗 14 号、冻结窗格、原列宽",
          summary["A1"].font.bold and summary["A1"].font.sz == 14
          and summary.freeze_panes == "A4"
          and summary.column_dimensions["A"].width == 14)

    for name, ok in checks:
        print(("OK   " if ok else "FAIL ") + name)
    return all(ok for _, ok in checks)


def main():
    ensure_deps()
    start = datetime.now()
    OUT_PATH.parent.mkdir(exist_ok=True)
    log_head()

    files = sorted(p for p in RAW_DIR.iterdir()
                   if p.suffix.lower() in (".xlsx", ".csv") and not p.name.startswith("~$"))
    raw, heads = [], []
    for path in files:
        head, records = read_fragment(path)
        heads.append((path.name, head))
        raw.extend(records)

    merged = pick_latest(raw)
    today, other = split_by_day(merged, REPORT_DAY)
    pending = [rec for rec in today if is_pending(rec)]
    done = [rec for rec in today if not is_pending(rec)]

    build_report(done, pending, other)
    passed = verify(raw, merged, done, pending, other, heads)

    cost = (datetime.now() - start).total_seconds()
    print(f"\n读进 {len(raw)} 条，去掉复测 {len(raw) - len(merged)} 条，"
          f"当日已出 {len(done)} 条 / 待出 {len(pending)} 条 / 补交非当日 {len(other)} 条")
    print("结果文件：02_output/检测日报.xlsx")
    print("判定：" + ("PASS" if passed else "FAIL"))
    with LOG.open("a", encoding="utf-8") as fh:
        fh.write(f"执行完成：{datetime.now():%Y-%m-%d %H:%M:%S}  用时 {cost:.1f}s  "
                 f"判定 {'PASS' if passed else 'FAIL'}\n")
    return 0 if passed else 1


if __name__ == "__main__":
    code = main()
    if "--no-pause" not in sys.argv:
        input("\n按回车键退出……")
    sys.exit(code)

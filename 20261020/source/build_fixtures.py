# -*- coding: utf-8 -*-
"""重造模板与片段数据。

上半个文件建模板（三个页签的标题、表头、框线、列宽、冻结窗格、以及一行空着的「样式行」），
下半个文件造五个格式各不相同的片段文件（表头行数不同、列名中英混写、xlsx 与 csv 混着来）。
"""

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent.parent
TPL = BASE / "source" / "templates" / "日报模板.xlsx"
RAW = BASE / "01_raw_data"
OUT = BASE / "02_output"

THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")


def head(cell, text):
    cell.value = text
    cell.font = Font(bold=True)
    cell.fill = HEAD_FILL
    cell.border = BOX
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_row(sheet, row, cols, widths, time_col=None):
    """留一行空白的「样式行」：run.py 写数据时照抄它的字体、框线、对齐和数字格式。"""
    for c in range(1, cols + 1):
        cell = sheet.cell(row=row, column=c)
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center")
        cell.number_format = "yyyy-mm-dd hh:mm" if c == time_col else "@"
    for c, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(c)].width = width


def build_template():
    book = Workbook()

    s = book.active
    s.title = "汇总"
    s["A1"] = "实验室检测日报"
    s["A1"].font = Font(bold=True, size=14)
    s.merge_cells("A1:F1")
    s["A2"] = "汇总日期："
    s["C2"] = "制表："
    for text, c in (("样品编号", 1), ("检测项目", 2), ("结果", 3),
                    ("单位", 4), ("检测时间", 5), ("组别", 6)):
        head(s.cell(row=3, column=c), text)
    style_row(s, 4, 6, (14, 22, 12, 10, 20, 14), time_col=5)
    s.freeze_panes = "A4"

    p = book.create_sheet("待出结果")
    p["A1"] = "待出结果（当天没出的，次日接着追）"
    p["A1"].font = Font(bold=True, size=13)
    p.merge_cells("A1:E1")
    for text, c in (("样品编号", 1), ("检测项目", 2), ("组别", 3),
                    ("检测时间", 4), ("备注", 5)):
        head(p.cell(row=2, column=c), text)
    style_row(p, 3, 5, (14, 22, 14, 20, 26), time_col=4)
    p.freeze_panes = "A3"

    t = book.create_sheet("统计")
    t["A1"] = "按组别统计"
    t["A1"].font = Font(bold=True, size=13)
    t.merge_cells("A1:E1")
    for text, c in (("组别", 1), ("已出结果", 2), ("待出结果", 3),
                    ("补交非当日", 4), ("合计", 5)):
        head(t.cell(row=2, column=c), text)
    style_row(t, 3, 5, (16, 12, 12, 14, 10))
    t.freeze_panes = "A3"

    TPL.parent.mkdir(parents=True, exist_ok=True)
    book.save(TPL)


def write_xlsx(path, info_rows, header, rows):
    """info_rows 是表头上方那几行说明——正是它们让表头不在第一行。"""
    book = Workbook()
    s = book.active
    for i, text in enumerate(info_rows, start=1):
        s.cell(row=i, column=1, value=text).font = Font(italic=True, color="808080")
    hr = len(info_rows) + 1
    for c, text in enumerate(header, start=1):
        head(s.cell(row=hr, column=c), text)
    for r, values in enumerate(rows, start=hr + 1):
        for c, v in enumerate(values, start=1):
            cell = s.cell(row=r, column=c, value=v)
            if isinstance(v, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
    for c in range(1, len(header) + 1):
        s.column_dimensions[get_column_letter(c)].width = 18
    book.save(path)


def write_csv(path, header, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(header)
        writer.writerows(rows)


def build_fragments():
    RAW.mkdir(parents=True, exist_ok=True)
    for old in list(RAW.glob("*.xlsx")) + list(RAW.glob("*.csv")):
        old.unlink()

    # 1. 表头在第 2 行（上面压了一行标题）；pH 没有单位，S-0020 结果还空着
    write_xlsx(
        RAW / "理化组_20261020.xlsx",
        ["理化组 2026-10-20 检测结果"],
        ["样品编号", "检测项目", "结果", "单位", "检测时间"],
        [
            ("S-0001", "pH", 6.82, "", datetime(2026, 10, 20, 9, 10)),
            ("S-0002", "水分", 3.15, "%", datetime(2026, 10, 20, 9, 20)),
            ("S-0012", "含量", 99.1, "%", datetime(2026, 10, 20, 9, 45)),
            ("S-0020", "含量", "", "%", datetime(2026, 10, 20, 9, 50)),
        ],
    )

    # 2. 表头就在第 1 行，但列名换了一套说法
    write_xlsx(
        RAW / "微生物组_20261020.xlsx",
        [],
        ["样品号", "项目", "检测结果", "单位", "完成时间"],
        [
            ("S-0003", "微生物限度", 120, "CFU/g", datetime(2026, 10, 20, 10, 5)),
            ("S-0004", "微生物限度", "<10", "CFU/g", datetime(2026, 10, 20, 10, 8)),
            ("S-0005", "微生物限度", "", "CFU/g", datetime(2026, 10, 20, 10, 30)),
        ],
    )

    # 3. 仪器软件直接导出的 csv：英文表头，时间是字符串，里面有一条是昨天测的
    write_csv(
        RAW / "仪器组_20261020导出.csv",
        ["Sample ID", "Item", "Result", "Unit", "TestTime"],
        [
            ("S-0006", "重金属(以Pb计)", "<0.05", "mg/kg", "2026-10-20 11:12"),
            ("S-0007", "重金属(以Pb计)", "0.12", "mg/kg", "2026-10-20 11:15"),
            ("S-0008", "重金属(以Pb计)", "0.08", "mg/kg", "2026-10-19 16:40"),
        ],
    )

    # 4. 表头压在第 3 行（上面两行是填写须知和制表人），结果用破折号占位
    write_xlsx(
        RAW / "前处理组_20261020.xlsx",
        ["填写须知：结果未出请留空，不要写“待定”", "制表：前处理组"],
        ["样品编号", "检测项目", "结果", "单位", "检测时间"],
        [
            ("S-0009", "干燥失重", 0.42, "%", datetime(2026, 10, 20, 13, 0)),
            ("S-0010", "干燥失重", 0.38, "%", datetime(2026, 10, 20, 13, 5)),
            ("S-0011", "干燥失重", "—", "%", datetime(2026, 10, 20, 13, 10)),
        ],
    )

    # 5. 补录：两条是今天的复测（比原结果晚），一条是昨天漏交的
    write_xlsx(
        RAW / "补录_20261019遗留.xlsx",
        [],
        ["样品编号", "检测项目", "结果", "单位", "检测时间"],
        [
            ("S-0002", "水分", 3.17, "%", datetime(2026, 10, 20, 15, 40)),
            ("S-0013", "含量", 98.5, "%", datetime(2026, 10, 19, 17, 20)),
            ("S-0001", "pH", 6.85, "", datetime(2026, 10, 20, 15, 50)),
        ],
    )


def main():
    build_template()
    build_fragments()
    OUT.mkdir(exist_ok=True)
    for old in OUT.glob("*.xlsx"):
        old.unlink()
    print("模板：" + str(TPL.relative_to(BASE)))
    print("片段：" + "、".join(sorted(p.name for p in RAW.iterdir() if p.suffix in (".xlsx", ".csv"))))


if __name__ == "__main__":
    main()

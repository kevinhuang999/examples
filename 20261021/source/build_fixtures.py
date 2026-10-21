# -*- coding: utf-8 -*-
"""造 01_raw_data：模拟质量部共享盘里散落的一堆检验记录 Excel 模板。

上半个文件建表格骨架（标题行 / 表头 / 列宽 / 表头样式）；
下半个文件按不同版本造出十一个文件——正常版、副本、改版、旧版、隐藏说明页、
无版本号版，外加锁文件、零字节文件、旧格式 .xls、以及不该被盘点的台账。

跑法：python source/build_fixtures.py
"""

import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE = Path(__file__).resolve().parent.parent
RAW = BASE / "01_raw_data"

# A 版式：12 列，最完整的一版
COLS_A = ["样品编号", "样品名称", "检测项目", "检测结果（mg/L）", "单位", "检出限",
          "检测方法", "仪器编号", "检测人", "检测日期", "复核人", "结论"]
# B 版式：在 A 的基础上多了复检结果和备注
COLS_B = ["样品编号", "样品名称", "检测项目", "检测结果（mg/L）", "单位", "检出限",
          "检测方法", "仪器编号", "检测人", "检测日期", "复核人",
          "复检结果", "备注", "结论"]
# C 版式：老记录本照抄的 10 列
COLS_C = ["序号", "样品编号", "检测项目", "检测结果（mg/L）", "单位",
          "检测方法", "检测人", "检测日期", "复核人", "结论"]
# 结果汇总页的列：两个可见页签都不同，但只在「检验记录」页上算指纹
COLS_SUM = ["检测项目", "检测数量", "合格数", "不合格数", "合格率"]

HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
HEAD_FONT = Font(bold=True, size=10)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_head(ws, cols, header_row, title=None, note=None):
    """按指定行号写标题 / 说明 / 表头，表头加粗加底色，列宽按字数给。"""
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=12)
    if note:
        ws.cell(row=2, column=1, value=note).font = Font(size=9, italic=True)
    for idx, name in enumerate(cols, 1):
        cell = ws.cell(row=header_row, column=idx, value=name)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = max(10, len(name) * 2)


def add_sample_row(ws, cols, row):
    """模板里留一行上次的示例数据——真实文件基本都这样，顺手把格式带出来。"""
    values = {c: "" for c in cols}
    values.update({"样品编号": "S20260301-001", "样品名称": "注射用水",
                   "检测项目": "pH 值", "检测结果（mg/L）": "6.85", "单位": "-",
                   "检测人": "示例", "结论": "符合规定"})
    for idx, name in enumerate(cols, 1):
        cell = ws.cell(row=row, column=idx, value=values.get(name, ""))
        cell.border = BOX


def new_book(ver=None, sheets=("检验记录", "结果汇总"), cols=COLS_A,
             header_row=2, note=None, hidden_note=False, sum_sheet=True):
    """建一个模板文件：标题行带版本号，表头在 header_row。"""
    wb = Workbook()
    first = wb.active
    first.title = sheets[0]
    title = "检验记录表" if ver is None else f"检验记录表（版本号：{ver}）"
    write_head(first, cols, header_row, title=title, note=note)
    add_sample_row(first, cols, header_row + 1)

    if sum_sheet and "结果汇总" in sheets:
        ws2 = wb.create_sheet("结果汇总")
        write_head(ws2, COLS_SUM, 2, title="结果汇总表")
    if hidden_note and "填写说明" in sheets:
        ws3 = wb.create_sheet("填写说明")
        ws3["A1"] = "1. 检测结果一律保留两位小数；2. 复核人需签字。"
        ws3.sheet_state = "hidden"      # 隐藏页，盘点时不算结构
    return wb


def save(wb, name, stamp, sub="."):
    """存盘并把修改时间改成指定值——不然后造的文件修改时间全一样，看不出先后。"""
    path = RAW / sub / name
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    ts = datetime(*stamp).timestamp()
    os.utime(path, (ts, ts))


def main():
    if RAW.exists():
        for old in RAW.rglob("*"):
            if old.is_file():
                old.unlink()
    RAW.mkdir(parents=True, exist_ok=True)
    (RAW / "02_旧格式").mkdir(exist_ok=True)

    # ① 基准版 V2.1：两个可见页签，表头在第 2 行
    save(new_book(ver="V2.1"), "检验记录模板.xlsx", (2026, 3, 2, 9, 10))

    # ② 副本：结构一模一样，但有人昨天刚保存过，修改时间最新
    save(new_book(ver="V2.1"), "检验记录模板 - 副本.xlsx", (2026, 9, 18, 16, 42))

    # ③ 名字最像终版，其实只有 14 列、版本号 V2.0
    save(new_book(ver="V2.0", cols=COLS_B), "检验记录模板(最终版).xlsx", (2026, 5, 11, 9, 0))

    # ④ 改版：标题占了 2 行，表头被压到第 3 行；列序打乱，列名带空格和半角括号
    cols_a_shuffled = [" 样品编号 ", "检测项目", "检测结果(mg/L)", "单位", "样品名称",
                       "检出限", "检测方法", "仪器编号", "检测人", "检测日期",
                       "复核人", "结论"]
    save(new_book(ver="V2.0", cols=cols_a_shuffled, header_row=3,
                  note="填表说明：结果保留两位小数；复核人需签字。"),
         "检验记录模板_2024改.xlsx", (2026, 4, 20, 11, 30))

    # ⑤ 旧版：只有 10 列，而且只有一个页签
    save(new_book(ver="V1.4", cols=COLS_C, sheets=("检验记录",), sum_sheet=False),
         "检验记录模板（旧）.xlsx", (2025, 8, 6, 15, 20))

    # ⑥ 台账：名字里没有「模板」，不该进盘点
    wb = Workbook()
    ws = wb.active
    ws.title = "台账"
    write_head(ws, ["序号", "样品编号", "领用日期", "领用人"], 1)
    save(wb, "台账_2024.xlsx", (2026, 9, 1, 8, 0))

    # ⑦ 最新一版 V2.2：多挂了一个隐藏的「填写说明」页
    save(new_book(ver="V2.2", sheets=("检验记录", "结果汇总", "填写说明"), hidden_note=True),
         "检验记录模板_2025新.xlsx", (2026, 8, 28, 10, 5))

    # ⑧ 无版本号：标题行被人删了，表头直接在第 1 行
    save(new_book(ver=None, header_row=1), "检验记录模板_无版本号.xlsx", (2026, 6, 12, 14, 0))

    # ⑨ Excel 打开文件时留下的锁文件
    (RAW / "~$检验记录模板.xlsx").write_bytes(b"\x00" * 256)

    # ⑩ 零字节：从邮件附件里拖出来，其实是坏的
    (RAW / "检验记录模板_损坏.xlsx").write_bytes(b"")

    # ⑪ 旧格式 .xls：openpyxl 打不开
    (RAW / "02_旧格式" / "检验记录模板_V1.xls").write_bytes(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

    print(f"已重建 {RAW}：11 个文件（含 5 份 A 版式、1 份 B 版式、1 份旧版、4 类应跳过）")


if __name__ == "__main__":
    main()
